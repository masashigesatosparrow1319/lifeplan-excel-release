from dataclasses import replace
import pytest
from src.lifeplan import demo, calculate as p1_calculate, annualize
from src.p2a import CashEvent, P2Plan, from_p1, calculate, corrected_price


def event(**changes):
    args=dict(id='E001',name='旅行',kind='expense',amount=1000000,account_id='cash',occurrence_month=(2026,7),additional_confirmed=True)
    args.update(changes)
    return CashEvent(**args)


def bare(start=(2026,1),end=(2026,12),cash=1000000):
    p=from_p1(demo())
    return replace(p,start=start,end=end,rules=[],accounts=[replace(p.accounts[0],balance=cash)])


def test_no_events_p1_regression():
    p=demo(); rows=calculate(from_p1(p))
    for a,b in zip(rows,p1_calculate(p)):
        assert {k:v for k,v in a.items() if k!='ledger'}==b


@pytest.mark.parametrize('annual,expected',[(4800000,2800000),(1800000,1300000)])
def test_midyear_income(annual,expected):
    p=from_p1(demo())
    p.rules=[replace(p.rules[0],end=(2026,6)),replace(p.rules[0],id='new',annual=annual,start=(2026,7)),p.rules[1]]
    rows=calculate(p)
    assert rows[5]['income']==300000 and rows[6]['income']==annual//12
    assert rows[-1]['cash']==expected


def test_retire_rehire():
    p=from_p1(demo()); salary,living=p.rules
    p.rules=[replace(salary,end=(2026,6)),replace(salary,id='retired',annual=0,start=(2026,7),end=(2026,9)),replace(salary,id='rehire',annual=2400000,start=(2026,10)),living]
    rows=calculate(p)
    assert [r['income'] for r in rows]==[300000]*6+[0]*3+[200000]*3
    assert rows[-1]['cash']==1000000


def test_once_and_income():
    p=from_p1(demo(),[event()]); rows=calculate(p)
    assert rows[5]['cash']==1600000 and rows[6]['cash']==700000 and rows[-1]['cash']==1200000
    assert [sum(i['amount'] for i in r['ledger'] if i['source_id']=='E001') for r in rows]==[0]*6+[1000000]+[0]*5
    p.events=[event(kind='income',amount=600000,occurrence_month=(2026,11))]
    assert calculate(p)[10]['income']==900000 and calculate(p)[-1]['cash']==2800000


def test_once_not_annual():
    p=bare(); p.events=[event(amount=100,occurrence_month=(2026,3))]
    rows=calculate(p)
    assert [r['expense'] for r in rows]==[0,0,100]+[0]*9 and rows[-1]['cash']==999900


@pytest.mark.parametrize('last,expected_years,expected_cash',[(2042,[2030,2034,2038,2042],8000000),(2041,[2030,2034,2038],11000000)])
def test_cycle(last,expected_years,expected_cash):
    p=bare((2029,1),(2043,12),20000000)
    p.events=[event(event_type='periodic',occurrence_month=None,first_month=(2030,5),interval_months=48,last_month=(last,5),amount=3000000)]
    rows=calculate(p)
    assert [r['month'] for r in rows if r['expense']]==[(y,5) for y in expected_years]
    assert rows[-1]['cash']==expected_cash
    p.start=(2032,1)
    assert calculate(p)[-1]['cash']==expected_cash+3000000


def test_interval_crosses_years():
    p=bare((2026,1),(2030,12)); p.events=[event(event_type='periodic',occurrence_month=None,first_month=(2026,11),interval_months=14,last_month=(2029,12),amount=100)]
    rows=calculate(p)
    assert [r['month'] for r in rows if r['expense']]==[(2026,11),(2028,1),(2029,3)]
    assert rows[-1]['cash']==999700


def test_outside_horizon():
    p=from_p1(demo(),[event(occurrence_month=(2026,3)),event(id='E002',kind='income',amount=600000,occurrence_month=(2026,10))])
    p.start=(2026,4); p.end=(2026,9)
    assert calculate(p)[-1]['cash']==1600000


def test_restricted_cash_and_shortage():
    p=from_p1(demo(),[event(occurrence_month=(2026,3))]); p.rules=[p.rules[1]]
    p.accounts[1]=replace(p.accounts[1],balance=1000000)
    rows=calculate(p); first=next(r for r in rows if r['shortage'])
    assert first['month']==(2026,3) and first['shortage']==600000 and first['total']==400000
    assert rows[-1]['cash']==-2400000 and rows[-1]['accounts']['reserve']==1000000


def test_zero_cash_no_warning():
    p=bare(); p.events=[event(occurrence_month=(2026,3))]
    rows=calculate(p); assert rows[-1]['cash']==0 and not any(r['shortage'] for r in rows)


def test_duplicate_and_distinct_keys():
    p=bare(); p.events=[event(economic_key='TRIP'),event(id='E002',economic_key='trip')]
    with pytest.raises(ValueError,match='二重計上'): calculate(p)
    p.events[1]=replace(p.events[1],economic_key='APPLIANCE')
    assert calculate(p)[-1]['cash']==-1000000
    p.events[1]=replace(p.events[1],economic_key='TRIP',occurrence_month=(2026,8))
    assert calculate(p)[-1]['cash']==-1000000


def test_rule_event_collision():
    p=from_p1(demo(),[event(kind='income',economic_key='salary')])
    with pytest.raises(ValueError,match='二重計上'): calculate(p)


def test_disabled_and_zero():
    p=bare(); p.events=[event(),event(id='E002',economic_key='E001',enabled=False,amount=-1)]
    assert calculate(p)[-1]['cash']==0
    p.events[1]=replace(p.events[1],enabled=True,amount=0)
    assert calculate(p)[-1]['cash']==0


@pytest.mark.parametrize('change',[{'amount':-1},{'amount':0.1},{'amount':1000000001},{'id':'salary'},{'account_id':'reserve'},{'account_id':'missing'},{'person_id':'missing'},{'additional_confirmed':False},{'event_type':'unknown'},{'economic_key':'a*'},{'economic_key':0},{'occurrence_month':(2026,13)}])
def test_invalid_inputs(change):
    p=from_p1(demo(),[event(**change)])
    with pytest.raises(ValueError): calculate(p)


@pytest.mark.parametrize('interval',[0,-1,1.5,None])
def test_invalid_period(interval):
    p=bare(); p.events=[event(event_type='periodic',occurrence_month=None,first_month=(2026,1),last_month=(2027,1),interval_months=interval)]
    with pytest.raises(ValueError): calculate(p)


def test_capacity_and_yearly():
    p=bare((2026,1),(2125,12)); p.events=[event(id=f'E{i:03}',amount=100,event_type='periodic',occurrence_month=None,first_month=(2026,1),last_month=(2125,1),interval_months=12) for i in range(200)]
    rows=calculate(p)
    assert rows[-1]['cash']==-1000000
    assert next(r['month'] for r in rows if r['shortage'])==(2076,1)
    assert all(v['expense']==20000 for v in annualize(rows).values())
    p.events.append(event(id='extra'))
    with pytest.raises(ValueError,match='容量'): calculate(p)


@pytest.mark.parametrize('amount,bps,year,expected',[(1000000,0,2028,1000000),(1000000,200,2027,1020000),(1000000,200,2028,1040400),(1000000,300,2028,1060900),(6,1000,2028,8),(1,500,2027,1),(10,500,2027,11),(1000000,-1000,2028,810000)])
def test_inflation_expected(amount,bps,year,expected):
    p=bare((2026,1),(2028,12)); p.events=[event(amount=amount,occurrence_month=(year,1),inflation_enabled=True,base_month=(2026,1),rate_bps=bps)]
    rows=calculate(p)
    assert sum(r['expense'] for r in rows)==expected
    assert rows[-1]['cash']==1000000-expected


def test_anniversary_and_rate_change():
    p=bare((2026,1),(2028,12)); e=event(occurrence_month=(2026,7),inflation_enabled=True,base_month=(2026,1),rate_bps=200)
    p.events=[e]
    assert sum(r['expense'] for r in calculate(p))==1000000
    e.occurrence_month=(2028,1)
    assert sum(r['expense'] for r in calculate(p))==1040400
    e.rate_bps=300
    assert sum(r['expense'] for r in calculate(p))==1060900
    e.base_month=(2027,1); e.rate_bps=200
    assert sum(r['expense'] for r in calculate(p))==1020000


def test_periodic_inflation():
    p=bare((2030,1),(2034,12)); p.events=[event(event_type='periodic',occurrence_month=None,first_month=(2030,5),last_month=(2034,5),interval_months=48,inflation_enabled=True,base_month=(2030,5),rate_bps=200)]
    assert [r['expense'] for r in calculate(p) if r['expense']]==[1000000,1082432]
    assert sum(r['expense'] for r in calculate(p))==2082432


@pytest.mark.parametrize('change',[{'rate_bps':1001},{'rate_bps':-1001},{'rate_bps':0.5},{'base_month':(2029,1)},{'kind':'income'},{'amount':1000000000,'rate_bps':1}])
def test_invalid_inflation(change):
    e=event(occurrence_month=(2028,1),inflation_enabled=True,base_month=(2026,1),rate_bps=200)
    p=bare((2026,1),(2028,12)); p.events=[replace(e,**change)]
    with pytest.raises(ValueError): calculate(p)


@pytest.fixture(scope='module')
def workbook_template(tmp_path_factory):
    from src.p2a_excel import generate
    path=tmp_path_factory.mktemp('p2a')/'base.xlsx'
    generate(from_p1(demo(),[event()]),path)
    return path


def test_workbook_roundtrip_and_input_frames(workbook_template):
    from src.p2a_excel import read_plan
    from openpyxl import load_workbook
    p=read_plan(workbook_template)
    assert calculate(p)[-1]['cash']==1200000
    w=load_workbook(workbook_template)
    assert w['イベント入力'].max_row==201 and w['収支'].max_row==51
    assert w['イベント入力'].protection.sheet and w['イベント入力']['B201'].protection.locked is False
    assert w['イベント入力']['A201'].protection.locked is True
    assert w['月次計上'].max_column==251 and w['重複確認'].max_column==253
    assert w['重複確認']['C3'].value.startswith('=IF(C$2>0,INDEX($B3:B3')
    assert not any(c.value and 'MOD(' in str(c.value) for row in w['年価格'] for c in row)


def test_add_change_disable_in_excel(workbook_template,tmp_path):
    from src.p2a_excel import read_plan
    from openpyxl import load_workbook
    w=load_workbook(workbook_template); s=w['イベント入力']
    # A new row uses its preassigned ID, default once/expense/person/purpose/key.
    for cell,value in {'B3':'家電','E3':__import__('datetime').date(2026,8,1),'H3':100000,'I3':'cash','M3':'はい'}.items(): s[cell]=value
    path=tmp_path/'edited.xlsx'; w.save(path)
    assert calculate(read_plan(path))[-1]['cash']==1100000
    s['H3']=200000; w.save(path)
    assert calculate(read_plan(path))[-1]['cash']==1000000
    s['S3']=0; w.save(path)
    assert calculate(read_plan(path))[-1]['cash']==1200000
    # Unused slots are ignored; clearing a filled row returns it to an unused slot.
    for c in range(2,20): s.cell(2,c,None).value=None
    w.save(path)
    assert calculate(read_plan(path))[-1]['cash']==2200000


def test_excel_input_collision_and_bounds(workbook_template,tmp_path):
    from src.p2a_excel import read_plan
    from openpyxl import load_workbook
    from datetime import date
    w=load_workbook(workbook_template); s=w['イベント入力']
    for cell,value in {'B3':'別の旅行','E3':date(2026,7,1),'H3':1,'I3':'cash','M3':'はい','L3':'E001'}.items(): s[cell]=value
    path=tmp_path/'bad.xlsx'; w.save(path)
    with pytest.raises(ValueError,match='二重計上'): calculate(read_plan(path))
    s['L3']='OTHER'; s['G3']=0; w.save(path)
    with pytest.raises(ValueError,match='競合'): calculate(read_plan(path))
    s['G3']=None; s['S3']=2; w.save(path)
    with pytest.raises(ValueError,match='有効'): read_plan(path)


def test_empty_template_and_partial_row(tmp_path):
    from src.p2a_excel import generate,read_plan
    from openpyxl import load_workbook
    path=tmp_path/'empty.xlsx'; generate(bare(),path)
    assert calculate(read_plan(path))[-1]['cash']==1000000
    w=load_workbook(path); w['イベント入力']['H2']=100; w.save(path)
    with pytest.raises(ValueError): calculate(read_plan(path))


@pytest.mark.parametrize('cell',['F2','J2','K2','L2'])
def test_zero_is_not_an_empty_optional_field(workbook_template,tmp_path,cell):
    from src.p2a_excel import read_plan
    from openpyxl import load_workbook
    w=load_workbook(workbook_template)
    w['イベント入力'][cell]=0
    path=tmp_path/'zero-input.xlsx'; w.save(path)
    with pytest.raises(ValueError): calculate(read_plan(path))


def test_precision_price():
    assert 10**9*11000+5000 < 2**53
    for amount in (0,1,6,10,999999999,1000000000):
        for bps in (-1000,-1,0,1,500,1000):
            numerator=amount*(10000+bps)+5000
            assert int(float(numerator)/10000)==numerator//10000


def test_generated_ids_avoid_existing_rule(tmp_path):
    from src.p2a_excel import generate,read_plan
    from openpyxl import load_workbook
    from datetime import date
    p=from_p1(demo()); p.rules[0]=replace(p.rules[0],id='E001')
    path=tmp_path/'ids.xlsx'; generate(p,path)
    w=load_workbook(path); s=w['イベント入力']
    assert s['A2'].value=='E002'
    for cell,value in {'B2':'旅行','E2':date(2026,7,1),'H2':100000,'I2':'cash','M2':'はい'}.items(): s[cell]=value
    w.save(path)
    assert calculate(read_plan(path))[-1]['cash']==2100000


def test_code_case_matches_excel_lookup():
    p=from_p1(demo(),[event(account_id='CASH',person_id='HOUSEHOLD')])
    p.rules[0]=replace(p.rules[0],account='CASH')
    assert calculate(p)[-1]['cash']==1200000


def test_full_case_fixed_expectations():
    from tests.p2a_cases import CASES
    for case in CASES:
        if not case.edits and not case.invalid:
            assert calculate(case.plan)[-1]['cash']==case.expected_cash,case.name


def test_p1_japanese_reference_ids_are_preserved(tmp_path):
    from src.p2a_excel import generate,read_plan
    p=from_p1(demo())
    p.persons[0]=replace(p.persons[0],id='本人')
    p.accounts=[replace(a,owner='本人' if a.owner=='adult1' else a.owner) for a in p.accounts]
    p.accounts[0]=replace(p.accounts[0],id='銀行*口座')
    p.rules=[replace(r,id=f'収支{idx}',target=f'項目{idx}',account='銀行*口座') for idx,r in enumerate(p.rules)]
    p.events=[event(account_id='銀行*口座',person_id='本人')]
    assert calculate(p)[-1]['cash']==1200000
    path=tmp_path/'legacy.xlsx'; generate(p,path)
    assert calculate(read_plan(path))[-1]['cash']==1200000


def test_p2a_powershell_bom():
    from pathlib import Path
    raw=Path('tools/recalculate_p2a.ps1').read_bytes()
    assert raw.startswith(b'\xef\xbb\xbf')
    assert 'Join-Path' not in raw.decode('utf-8-sig').split('$ErrorActionPreference',1)[0]
