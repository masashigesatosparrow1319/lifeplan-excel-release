"""Artificial XLSX/XML fixtures ONLY, not Excel recalculation evidence.

Unlike the earlier A COM mocks, these exercise the real cases.verify_saved().
Cached values below are explicit test doubles, never product/runtime outputs.
"""
from datetime import datetime
import copy
import json
from pathlib import Path
import shutil
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
import pytest

from tools.excel2024_validation import cases, common

M = cases.NS['m']


def rewrite(path, mutate):
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    mutate(parts)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(n, data)


def artificial_caches(path, probes):
    values = {c: v for s, c, v in probes if s == '確認'}
    def patch(parts):
        root = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        for cell in root.findall('.//m:c', cases.NS):
            if cell.find('m:f', cases.NS) is None or cell.get('r') not in values:
                continue
            value = values[cell.get('r')]
            cell.set('t', 'str' if isinstance(value, str) else 'n')
            v = cell.find('m:v', cases.NS)
            if v is None: v = ET.SubElement(cell, '{' + M + '}v')
            v.text = str(value)
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(root)
    rewrite(path, patch)


@pytest.fixture
def packet(tmp_path):
    directory = tmp_path / 'case'
    (directory / 'inputs').mkdir(parents=True)
    (directory / 'results').mkdir()
    file = '01_function_gate.xlsx'
    source, target = (directory / p / file for p in ('inputs', 'results'))
    w = Workbook(); s = w.active; s.title = '確認'
    s['B2'] = '人工fixture'; s['B8'] = 100
    s['C3'] = '文字列'; s['D4'] = 0; s['E4'] = False; s['F4'] = datetime(2030, 4, 2)
    w.calculation.calcMode = 'auto'
    probes = [['確認', c, v] for c, v in cases.EXPECTED24.items()]
    for _, c, _ in probes: s[c] = '=1'  # artificial expression, not S1 model
    w.save(source); w.close(); shutil.copyfile(source, target)
    artificial_caches(target, probes)
    manifest = dict(case_id='functions24', file=file, sha256=common.sha(source), formula_count=24, probes=probes)
    return directory, manifest, source, target


def edit(target, probes, sheet='確認', cell='A1', value=None, style=False):
    w = load_workbook(target)
    if sheet not in w.sheetnames: w.create_sheet(sheet)
    s = w[sheet]
    if style: s[cell].fill = PatternFill('solid', fgColor='FFFF00')
    else: s[cell] = value
    w.save(target); w.close()
    artificial_caches(target, probes)


def verify(packet):
    return cases.verify_saved(packet[0], packet[1])


def test_legacy_emptycell_coordinate_exception_and_real_fixed_branch(packet):
    directory, manifest, source, target = packet
    w = load_workbook(source, read_only=True)
    try:
        with pytest.raises(AttributeError, match='coordinate'):
            for row in w['確認']:
                for c in row:
                    if c.data_type != 'f': _ = c.coordinate
    finally: w.close()
    assert verify(packet)['numeric_comparison'] == 'matched'


@pytest.mark.parametrize('cell,value', [('A1', 0), ('A1', False), ('A1', 'x'), ('A1', datetime(2030, 1, 1)),
                                       ('C3', None), ('C3', '変化'), ('D4', False), ('E4', 0), ('F4', 40000)])
def test_inside_blank_value_and_type_changes_rejected(packet, cell, value):
    edit(packet[3], packet[1]['probes'], cell=cell, value=value)
    with pytest.raises(ValueError, match='非数式'): verify(packet)


@pytest.mark.parametrize('cell,value', [('N2', 1), ('B14', 'x'), ('A1', '前'), ('CV500', False), ('CV500', 0)])
def test_outside_original_values_rejected(packet, cell, value):
    edit(packet[3], packet[1]['probes'], cell=cell, value=value)
    with pytest.raises(ValueError, match='非数式'): verify(packet)


@pytest.mark.parametrize('cell', ['N2', 'B14', 'A1', 'CV500'])
def test_style_only_union_expansion_and_blank_cells_allowed(packet, cell):
    edit(packet[3], packet[1]['probes'], cell=cell, style=True)
    assert verify(packet)['numeric_comparison'] == 'matched'


def test_numeric_int_float_representation_not_rejected(packet):
    edit(packet[3], packet[1]['probes'], cell='D4', value=0.0)
    assert verify(packet)['numeric_comparison'] == 'matched'


@pytest.mark.parametrize('value', [0, False, 'new'])
def test_new_sheet_nonempty_values_rejected(packet, value):
    edit(packet[3], packet[1]['probes'], sheet='追加', cell='C7', value=value)
    with pytest.raises(ValueError, match='非数式'): verify(packet)


@pytest.mark.parametrize('cell,value', [('I8', '=2'), ('I8', 0), ('Z20', '=1')])
def test_formula_change_loss_addition_keeps_existing_rejection(packet, cell, value):
    edit(packet[3], packet[1]['probes'], cell=cell, value=value)
    with pytest.raises(ValueError, match='数式の消失/変更'): verify(packet)


def test_new_empty_formatted_sheet_has_no_added_values(packet):
    edit(packet[3], packet[1]['probes'], sheet='空白', cell='C7', style=True)
    assert verify(packet)['numeric_comparison'] == 'matched'


@pytest.mark.parametrize('problem', ['cache_missing', 'cache_wrong', 'manual', 'error', 'input_sha', 'missing_result',
                                    'broken_zip', 'hidden_cell', 'huge_dimension', 'missing_dimension', 'time_limit'])
def test_existing_checks_and_defensive_limits(packet, monkeypatch, problem):
    directory, manifest, source, target = packet
    if problem == 'input_sha': manifest['sha256'] = '0' * 64
    elif problem == 'missing_result': target.unlink()
    elif problem == 'broken_zip': target.write_bytes(b'not a zip')
    elif problem == 'error': edit(target, manifest['probes'], cell='G5', value='#DIV/0!')
    elif problem == 'time_limit': monkeypatch.setattr(cases, 'F24_SECONDS', -1)
    else:
        def patch(parts):
            key = 'xl/workbook.xml' if problem == 'manual' else 'xl/worksheets/sheet1.xml'
            root = ET.fromstring(parts[key])
            if problem == 'manual': root.find('m:calcPr', cases.NS).set('calcMode', 'manual')
            elif problem in ('huge_dimension', 'hidden_cell', 'missing_dimension'):
                dim = root.find('m:dimension', cases.NS)
                if problem == 'missing_dimension': root.remove(dim)
                else: dim.set('ref', 'A1:XFD1048576' if problem == 'huge_dimension' else 'A1:B2')
            else:
                cell = next(c for c in root.findall('.//m:c', cases.NS) if c.get('r') == 'I8')
                cell.find('m:v', cases.NS).text = None if problem == 'cache_missing' else '999'
            parts[key] = ET.tostring(root)
        rewrite(target, patch)
    with pytest.raises((ValueError, AssertionError, common.GateError, FileNotFoundError, zipfile.BadZipFile)):
        verify(packet)


@pytest.mark.parametrize('limit', ['F24_ARCHIVE', 'F24_EXPANDED', 'F24_CELLS', 'F24_ROWS', 'F24_COLUMNS'])
def test_small_resource_limits_fail_closed(packet, monkeypatch, limit):
    monkeypatch.setattr(cases, limit, 1)
    with pytest.raises(common.GateError) as error: verify(packet)
    assert error.value.status == 'BLOCKED'


def test_budget_boundary_inclusive_and_one_over_rejected(packet, monkeypatch):
    # Pin the actual archive byte boundary; no huge allocation is required.
    size = max(packet[2].stat().st_size, packet[3].stat().st_size)
    monkeypatch.setattr(cases, 'F24_ARCHIVE', size)
    assert verify(packet)['numeric_comparison'] == 'matched'
    monkeypatch.setattr(cases, 'F24_ARCHIVE', size - 1)
    with pytest.raises(common.GateError): verify(packet)


# Offline fixtures below use a mock Excel worker and ARTIFICIAL cached XML.
# They never claim to be the owner's G2-010 files or prove Excel execution.
@pytest.fixture
def offline_packet(tmp_path, monkeypatch):
    from tools.excel2024_validation import offline_saved as offline, preopen_input as builder
    from tools.excel2024_validation import windows_preflight as pre, A_DIAGNOSTIC, A_PROTOCOL
    from tests.test_excel2024_functions24 import profile
    from tests.test_excel2024_preopen_diagnostic import runtime as fake_runtime
    import hashlib
    root = tmp_path / 'historical'; root.mkdir()
    directory = root / 'functions24'; directory.mkdir()
    files = pre.source_fingerprint()['files'].copy()
    files.pop('tools/excel2024_validation/offline_saved.py')
    files['tools/excel2024_validation/cases.py'] = offline.OLD_CASES
    old_pin = hashlib.sha256(json.dumps(files, sort_keys=True).encode()).hexdigest()
    manifest = builder.prepare(directory, 'ARTIFICIAL-RUN', 'G2-20000101-001-WIN-01', tmp_path, old_pin)
    request = dict(run_id=manifest['run_id'], case_id='functions24', diagnostic=A_DIAGNOSTIC,
                   protocol=A_PROTOCOL, control=manifest['control'], profile=profile(),
                   source_sha256=old_pin, manifest_sha256=common.sha(directory / 'case.json'))
    result = fake_runtime(directory, request)
    target = directory / 'results' / manifest['file']
    shutil.copyfile(directory / 'inputs' / manifest['file'], target)
    artificial_caches(target, manifest['probes'])
    result['saved_sha256'] = common.sha(target)
    common.write_json(directory / 'worker.json', result)
    result['residual_observation'] = dict(outcome='PROCESS_EXITED')
    parent = dict(mode='functions24-preopen-force0', diagnostic=A_DIAGNOSTIC, protocol=A_PROTOCOL,
                  status='FAIL', exit_code=1, run_id=manifest['run_id'],
                  source=dict(files=files, source_sha256=old_pin, commit=None), stages=[
                      dict(case_id='tiny3', status='NOT_RUN'),
                      dict(case_id='functions24', status='FAIL', runtime=result,
                           control_verification=result['preopen_control'],
                           residual=dict(same_owned_process_remains=False),
                           first_error=dict(type='AttributeError', message="'EmptyCell' has no 'coordinate'")),
                      dict(case_id='annual1', status='NOT_RUN')])
    common.write_json(root / 'run.json', parent)
    # Test-only target pins; production CLI never accepts substitute G2-010 XLSX pins.
    monkeypatch.setattr(offline, 'INPUT_SHA', manifest['sha256'])
    monkeypatch.setattr(offline, 'SAVED_SHA', result['saved_sha256'])
    pin = pre.source_fingerprint()['release_source_sha256']
    return root, tmp_path / 'offline-audit', pin, common.sha(root / 'run.json')


def audit(packet):
    from tools.excel2024_validation import offline_saved
    root, output, pin, old_run_sha = packet
    return offline_saved.execute(root, output, pin, old_run_sha, confirmed=True)


def test_offline_real_readonly_checkers_with_artificial_cache_fixture(offline_packet):
    from tools.excel2024_validation import offline_saved
    root, output, pin, run_sha = offline_packet
    before = offline_saved.inventory(root)
    result = audit(offline_packet)
    assert result['status'] == 'OFFLINE_REVERIFICATION_PASS', result
    assert result['original_status'] == 'FAIL' and result['gate2_acceptance'] is False
    assert result['excel_started'] is result['com_connected'] is result['approval_id_used'] is False
    assert offline_saved.inventory(root) == before
    assert common.sha(root / 'run.json') == run_sha
    assert common.sha(output / 'offline-report.json') == (output / 'offline-report.sha256').read_text(encoding='ascii').strip()
    assert audit(offline_packet)['status'] == 'BLOCKED'  # never overwrite an audit


@pytest.mark.parametrize('problem', ['saved_sha', 'input_sha', 'run_sha', 'source_sha', 'parent_pass',
                                    'worker_cleanup', 'process_remains', 'transform', 'request', 'marker'])
def test_offline_rejects_evidence_changes_and_preserves_original_fail(offline_packet, problem):
    root, output, pin, run_sha = offline_packet
    case = root / 'functions24'
    if problem == 'saved_sha':
        with (case / 'results/01_function_gate.xlsx').open('ab') as stream: stream.write(b'X')
    elif problem == 'input_sha':
        with (case / 'inputs/01_function_gate.xlsx').open('ab') as stream: stream.write(b'X')
    elif problem == 'source_sha': offline_packet = (root, output, '0'*64, run_sha)
    elif problem == 'run_sha': offline_packet = (root, output, pin, '0'*64)
    else:
        path = (root / 'run.json' if problem in ('parent_pass', 'process_remains') else
                case / ('worker.json' if problem == 'worker_cleanup' else problem + '.json'))
        if problem == 'marker':
            path = Path(common.read_json(case / 'case.json')['control']['approval_path'])
        value = common.read_json(path)
        if problem == 'parent_pass': value['status'] = 'PASS'
        elif problem == 'process_remains': value['stages'][1]['runtime']['residual_observation']['outcome'] = 'OWNED_PROCESS_REMAINS'
        elif problem == 'worker_cleanup': value['cleanup_errors'] = ['ARTIFICIAL']
        else: value['ARTIFICIAL_CHANGE'] = True
        common.write_json(path, value)
        if path == root / 'run.json': offline_packet = (root, output, pin, common.sha(path))
    result = audit(offline_packet)
    assert result['status'] != 'OFFLINE_REVERIFICATION_PASS'
    assert 'ARTIFICIAL-RUN' not in json.dumps(result)
    assert str(root) not in json.dumps(result)


def test_offline_rejects_missing_output_or_evidence_without_exposing_private_paths(offline_packet):
    root, output, pin, run_sha = offline_packet
    (root / 'functions24/worker.json').unlink()
    assert audit(offline_packet)['status'] == 'OFFLINE_REVERIFICATION_FAIL'


def test_offline_does_not_consume_id_prepare_connect_or_launch(offline_packet, monkeypatch):
    from tools.excel2024_validation import offline_saved, preopen_input, run, functions24, session
    def forbidden(*args, **kwargs): raise AssertionError('MUST_NOT_CALL')
    for module, name in [(preopen_input, 'prepare'), (run, 'execute'), (run, 'supervise'),
                         (functions24, 'execute'), (session.Session, 'connect')]:
        monkeypatch.setattr(module, name, forbidden)
    assert audit(offline_packet)['status'] == 'OFFLINE_REVERIFICATION_PASS'


def test_offline_pins_and_initial_consent_required(offline_packet):
    from tools.excel2024_validation import offline_saved
    root, output, pin, run_sha = offline_packet
    result = offline_saved.execute(root, output, pin, run_sha, confirmed=False)
    assert result['status'] == 'BLOCKED' and not output.exists()


@pytest.mark.parametrize('problem', ['duplicate', 'dtd', 'row_order', 'cell_order', 'oversized_actual_cell'])
def test_bounded_xml_and_zip_structure_rejected(packet, problem):
    target = packet[3]
    if problem == 'duplicate':
        with zipfile.ZipFile(target, 'a') as z:
            with pytest.warns(UserWarning, match='Duplicate'):
                z.writestr('xl/workbook.xml', b'<invalid/>')
    else:
        def patch(parts):
            key = 'xl/worksheets/sheet1.xml'
            if problem == 'dtd': parts[key] = b'<!DOCTYPE a []>' + parts[key]; return
            root = ET.fromstring(parts[key])
            data = root.find('m:sheetData', cases.NS)
            if problem == 'row_order': data.append(copy.deepcopy(data[0]))
            elif problem == 'cell_order': data[0].append(copy.deepcopy(data[0][0]))
            else: data[0][0].set('r', 'XFD1048576')
            parts[key] = ET.tostring(root)
        rewrite(target, patch)
    with pytest.raises(common.GateError) as error: verify(packet)
    assert error.value.status == 'BLOCKED'


def test_offline_snapshot_change_is_blocked_not_a_pass(offline_packet, monkeypatch):
    from tools.excel2024_validation import offline_saved
    real = offline_saved.inventory
    count = [0]
    def changed(root):
        value = real(root); count[0] += 1
        if count[0] == 2: value['run.json'] = 'SIMULATED_RACE'
        return value
    monkeypatch.setattr(offline_saved, 'inventory', changed)
    assert audit(offline_packet)['status'] == 'BLOCKED'
    assert not offline_packet[1].exists()


def test_offline_cli_returns_redacted_separate_status_only(monkeypatch, capsys):
    from tools.excel2024_validation import offline_saved
    import sys
    monkeypatch.setattr(sys, 'argv', ['offline_saved', '--source-run', 'PRIVATE_PATH', '--output', 'PRIVATE_OUTPUT',
                                      '--expected-source-sha', '0'*64, '--expected-run-sha', '0'*64])
    monkeypatch.setattr(sys, 'addaudithook', lambda hook: None)  # real guard is exercised by the full suite
    assert offline_saved.main() == 2
    text = capsys.readouterr().out
    assert 'PRIVATE' not in text and json.loads(text)['status'] == 'BLOCKED'


@pytest.mark.parametrize('value', [0, False, '追加', datetime(2030, 4, 2)])
def test_internal_emptycell_hole_cannot_be_filled(packet, value):
    edit(packet[3], packet[1]['probes'], cell='G5', value=value)
    with pytest.raises(ValueError, match='非数式'): verify(packet)


@pytest.mark.parametrize('cell,allowed', [('DX512', True), ('DY512', False), ('DX513', False)])
def test_dimension_boundary_and_style_only_blocked_over_budget(packet, cell, allowed):
    edit(packet[3], packet[1]['probes'], cell=cell, style=True)
    if allowed: assert verify(packet)['numeric_comparison'] == 'matched'
    else:
        with pytest.raises(common.GateError) as error: verify(packet)
        assert error.value.status == 'BLOCKED'


def test_offline_resource_blocked_stays_blocked(offline_packet, monkeypatch):
    monkeypatch.setattr(cases, 'F24_ROWS', 1)
    result = audit(offline_packet)
    assert result['status'] == 'BLOCKED'
    assert result['reason'] == 'SAVED_VERIFIER_LIMIT_OR_FAILURE'
