import os
from dataclasses import replace
from datetime import date
from pathlib import Path
import json
import hashlib
import math
import random
from fractions import Fraction
import pytest
from openpyxl import load_workbook
from src.lifeplan import demo, calculate, annualize, allocate, generate, Person, Rule, verification_identity, MAX_AMOUNT, MAX_WEIGHT


ALLOCATION_CASES = [
    ('equal_100',100,[1]*12,[9]*4+[8]*8),
    ('first_six_100',100,[1]*6+[0]*6,[17]*4+[16]*2+[0]*6),
    ('december_zero',100,[1]*11+[0],[10]+[9]*10+[0]),
    ('january_only',100,[1]+[0]*11,[100]+[0]*11),
    ('zero_amount',0,[1]*12,[0]*12),
    ('skewed',100,[1000]+[1]*11,[99,1]+[0]*10),
    ('integer_weights',100,[1,2,3]+[0]*9,[17,33,50]+[0]*9),
    ('unequal_equal_remainders',100,[1,4,7]+[0]*9,[9,33,58]+[0]*9),
    ('integer_scaled_remainders',100,[10,40,70]+[0]*9,[9,33,58]+[0]*9),
    ('upper_equal',10**9,[1]*12,[83333334]*4+[83333333]*8),
    ('upper_weights',10**9,[1000]*12,[83333334]*4+[83333333]*8),
    ('upper_integer_weights',10**9,[1,2,3]+[0]*9,[166666667,333333333,500000000]+[0]*9),
    ('upper_skewed',10**9,[1000]+[1]*11,[989119683]+[989120]*8+[989119]*3),
    ('near_upper',999999999,[1000,999,1]+[0]*9,[499999999,499500000,500000]+[0]*9),
    ('upper_only',10**9,[1000]+[0]*11,[10**9]+[0]*11),
]

# Intentionally invalid Excel inputs bypassing data validation (e.g. paste).
INVALID_WORKBOOK_CASES = [
    ('invalid_decimal_weights','収支',{'H2':0.1,'I2':0.2,'J2':0.3},3),
    ('invalid_weight_upper','収支',{'H2':1001},3),
    ('invalid_weight_negative','収支',{'H2':-1},3),
    ('invalid_zero_weights','収支',{f'{chr(c)}2':0 for c in range(ord('H'),ord('S')+1)},3),
    ('invalid_annual_upper','収支',{'D2':10**9+1},3),
    ('invalid_annual_fraction','収支',{'D2':100.1},3),
    ('invalid_balance_upper','口座',{'E2':10**9+1},5),
    ('invalid_balance_fraction','口座',{'E2':100.1},5),
]


def allocation_plan(amount, weights, end=(2026,12)):
    p=demo()
    return replace(p,rules=[replace(p.rules[0],annual=amount,weights=weights,end=end),replace(p.rules[1],annual=0)])


def cases():
    p=demo()
    yield 'basic',p,2200000
    yield 'retire',replace(p,rules=[replace(p.rules[0],end=(2026,6)),p.rules[1]]),400000
    yield 'partial',replace(p,start=(2026,4),end=(2026,9)),1600000
    yield 'second_half',replace(p,rules=[replace(p.rules[0],start=(2026,7)),p.rules[1]]),400000
    yield 'bonus',replace(p,rules=[replace(p.rules[0],weights=[0]*5+[1]+[0]*5+[1]),p.rules[1]]),2200000
    yield 'rounding',replace(p,rules=[replace(p.rules[0],annual=100),replace(p.rules[1],annual=0)]),1000100
    yield 'couple',replace(p,persons=p.persons+[Person('adult2','adult',date(1992,2,1))]),2200000
    yield 'shortage',replace(p,rules=[replace(p.rules[0],annual=0),p.rules[1]]),-1400000
    yield 'change',replace(p,rules=[replace(p.rules[0],end=(2026,6)),replace(p.rules[0],id='new',start=(2026,7),annual=4800000),p.rules[1]]),2800000
    yield 'weight_change',replace(p,rules=[replace(p.rules[0],end=(2026,6)),replace(p.rules[0],id='new',start=(2026,7),weights=[0]*11+[1]),p.rules[1]]),4000000
    yield 'cross_year',replace(p,start=(2026,11),end=(2027,2),rules=[replace(r,end=(2027,12)) for r in p.rules]),1400000
    for name,amount,weights,expected in ALLOCATION_CASES:
        yield name,allocation_plan(amount,weights),1000000+amount
    # Allocate first, then mask inactive months: do not redistribute July-Dec.
    yield 'rounding_retire',allocation_plan(100,[1]*12,(2026,6)),1000052


@pytest.mark.parametrize('name,plan,expected',list(cases()))
def test_reference(name,plan,expected):
    rows=calculate(plan)
    assert rows[-1]['cash']==expected
    assert rows[-1]['total']==expected+500000
    assert all(r['accounts']['reserve']==500000 for r in rows)
    yearly=annualize(rows)
    assert sum(y['income'] for y in yearly.values())==sum(r['income'] for r in rows)
    assert sum(y['expense'] for y in yearly.values())==sum(r['expense'] for r in rows)


def test_first_shortage():
    p=dict((n,p) for n,p,e in cases())['shortage']
    rows=calculate(p)
    first=next(r for r in rows if r['shortage'])
    assert first['month']==(2026,6) and first['shortage']==200000
    assert rows[4]['cash']==0 and rows[4]['shortage']==0


def test_allocation():
    assert allocate(100,[1]*12)==[9]*4+[8]*8
    assert sum(allocate(100,[1]*12))==100
    for weights in ([0]*12,[-1]+[1]*11,[1]*11):
        with pytest.raises(ValueError): allocate(100,weights)


@pytest.mark.parametrize('name,amount,weights,expected',ALLOCATION_CASES)
def test_allocation_explicit_expected(name,amount,weights,expected):
    assert allocate(amount,weights)==expected
    assert sum(expected)==amount
    assert all(type(v) is int and v>=0 for v in expected)
    assert all(v==0 for v,w in zip(expected,weights) if w==0)
    total=sum(Fraction(str(w)) for w in weights)
    assert all(abs(Fraction(v)-amount*Fraction(str(w))/total)<1 for v,w in zip(expected,weights))


def test_retirement_does_not_redistribute():
    rows=calculate(allocation_plan(100,[1]*12,(2026,6)))
    assert [r['income'] for r in rows]==[9,9,9,9,8,8,0,0,0,0,0,0]
    assert annualize(rows)[2026]['income']==52
    # Positive weights outside effective periods are not transferred to active months.
    rows=calculate(allocation_plan(100,[1]*6+[0]*6,(2026,3)))
    assert [r['income'] for r in rows]==[17,17,17]+[0]*9
    assert annualize(rows)[2026]['income']==51


def test_all_zero_weights_rejected_even_for_zero_amount():
    for amount in (0,100):
        with pytest.raises(ValueError,match='合計が正'):
            allocate(amount,[0]*12)
        with pytest.raises(ValueError,match='合計が正'):
            calculate(allocation_plan(amount,[0]*12))


def test_scale_and_integer_ties():
    assert allocate(100,[1]*6+[0]*6)==[17]*4+[16]*2+[0]*6
    assert allocate(100,[1000]*6+[0]*6)==[17]*4+[16]*2+[0]*6
    assert allocate(100,[1.0,2.0,3.0]+[0]*9)==[17,33,50]+[0]*9


@pytest.mark.parametrize('value',[-1,MAX_WEIGHT+1,0.1,0.2,0.3,1.5,float('nan'),float('inf'),True,'1',None])
def test_invalid_ratio(value):
    with pytest.raises(ValueError,match='月別比率'):
        allocate(100,[value]+[1]*11)


@pytest.mark.parametrize('value',[-1,MAX_AMOUNT+1,0.1,float('nan'),float('inf'),True,'100',None,10**1000])
def test_invalid_annual(value):
    with pytest.raises(ValueError,match='年額'):
        allocate(value,[1]*12)


@pytest.mark.parametrize('value',[-1,MAX_AMOUNT+1,0.1,float('nan'),float('inf'),True,'100',None,10**1000])
def test_invalid_initial_balance(value):
    p=demo()
    with pytest.raises(ValueError,match='初期現預金残高'):
        calculate(replace(p,accounts=[replace(p.accounts[0],balance=value),p.accounts[1]]))


def test_initial_balance_limit():
    p=demo()
    rows=calculate(replace(p,accounts=[replace(p.accounts[0],balance=MAX_AMOUNT),p.accounts[1]]))
    assert rows[0]['opening']==MAX_AMOUNT
    assert rows[-1]['cash']==MAX_AMOUNT+1200000  # limit applies to initial balance, not computed closing cash


def test_double_integer_precision_boundaries():
    assert MAX_AMOUNT*MAX_WEIGHT==10**12 < 2**53
    assert 12*MAX_WEIGHT==12000
    # Closest possible noninteger quotient is >=1/12000 from an integer.
    assert math.ulp(float(MAX_AMOUNT)) < 1/12000
    rng=random.Random(2010)
    samples=[(rng.randint(0,MAX_AMOUNT),rng.randint(1,MAX_WEIGHT),rng.randint(0,11*MAX_WEIGHT)) for _ in range(10000)]
    # Construct remainders of 1 and W-1 near the annual upper bound.
    for total in range(2,12001):
        w=min(MAX_WEIGHT,total)
        if math.gcd(w,total)==1:
            for remainder in (1,total-1):
                a0=(remainder*pow(w,-1,total))%total
                amount=a0+(MAX_AMOUNT-a0)//total*total
                samples.append((amount,w,total-w))
    for amount,weight,others in samples:
        total=weight+others
        product=amount*weight
        quotient,remainder=divmod(product,total)
        assert float(amount)*float(weight)==product
        double_q=math.floor(float(product)/float(total))
        assert double_q==quotient
        assert float(product)-double_q*float(total)==remainder


def test_conflict_and_independent_income():
    p=demo()
    with pytest.raises(ValueError,match='期間重複'):
        calculate(replace(p,rules=p.rules+[replace(p.rules[0],id='duplicate')]))
    extra=replace(p.rules[0],id='sidejob',target='sidejob',annual=120000)
    assert calculate(replace(p,rules=p.rules+[extra]))[-1]['cash']==2320000
    with pytest.raises(ValueError,match='期間重複'):
        calculate(replace(p,rules=p.rules+[replace(p.rules[1],id='expense_duplicate')]))


def test_restricted_and_child():
    p=demo()
    with pytest.raises(ValueError,match='制限口座'):
        calculate(replace(p,rules=[replace(p.rules[1],account='reserve')]))
    child=Person('child1','child',date(2030,1,1),True)
    calculate(replace(p,persons=p.persons+[child]))
    calculate(replace(p,persons=p.persons+[replace(child,id=f'child{i}') for i in range(6)]))
    with pytest.raises(ValueError): calculate(replace(p,persons=[]))
    with pytest.raises(ValueError): calculate(replace(p,persons=p.persons+[replace(child,id=f'child{i}') for i in range(7)]))


def test_cross_year():
    p=demo(); p=replace(p,start=(2026,11),end=(2027,2),rules=[replace(r,end=(2027,12)) for r in p.rules])
    ys=annualize(calculate(p))
    assert ys[2026]['income']==600000 and ys[2027]['income']==600000
    assert ys[2026]['cash']==1200000 and ys[2027]['cash']==1400000


@pytest.mark.parametrize('name,plan,expected',list(cases()))
def test_workbook_structure(tmp_path,name,plan,expected):
    path=tmp_path/f'{name}.xlsx'; generate(plan,path)
    w=load_workbook(path)
    assert w.calculation.fullCalcOnLoad
    assert w['月次CF'].max_row==len(calculate(plan))+1
    assert w['月次CF']['E2'].value=='=B2+C2-D2'
    assert w['収支']['D2'].value==plan.rules[0].annual
    assert w['入力チェック'].max_row==len(plan.rules)+len(plan.accounts)+2
    assert w['不足サマリ']['B4'].value=='未判定（P1対象外）'
    assert w['配分根拠'].max_column==38
    for sheet,cell,max_value in [('収支','D2',MAX_AMOUNT),('収支','H2',MAX_WEIGHT),('口座','E2',MAX_AMOUNT)]:
        validations=[v for v in w[sheet].data_validations.dataValidation if cell in v.sqref]
        assert len(validations)==1 and validations[0].type=='whole'
        assert validations[0].formula1=='0' and validations[0].formula2==str(max_value)


@pytest.mark.parametrize('name,plan,expected',list(cases()))
def test_recalculated_excel(name,plan,expected):
    directory=os.environ.get('LIFEPLAN_RECALCULATED_DIR')
    if not directory:
        pytest.skip('未検証: LIFEPLAN_RECALCULATED_DIR に再計算済みケースを指定してください')
    w=load_workbook(Path(directory)/f'{name}.xlsx',data_only=True)
    formulas=load_workbook(Path(directory)/f'{name}.xlsx',data_only=False)
    assert (w['_検証メタ']['B1'].value,w['_検証メタ']['B2'].value)==verification_identity(plan),'旧版または別入力のExcelです。最新ケースを再計算してください'
    # Reject old-input workbooks and frozen value-only copies.
    for i,r in enumerate(plan.rules,2):
        assert w['収支'].cell(i,4).value==r.annual
        assert w['収支'].cell(i,5).value.date()==date(*r.start,1)
        assert w['収支'].cell(i,6).value.date()==date(*r.end,1)
        assert [w['収支'].cell(i,j).value for j in range(8,20)]==r.weights
        assert formulas['月別配分'].cell(i,2).data_type=='f'
        expected_allocation=allocate(r.annual,r.weights)
        assert [w['月別配分'].cell(i,j).value for j in range(2,14)]==expected_allocation
        products=[int(r.annual)*int(v) for v in r.weights]
        total=sum(r.weights)
        assert [w['配分根拠'].cell(i,j).value for j in range(2,14)]==[p//total for p in products]
        assert [w['配分根拠'].cell(i,j).value for j in range(14,26)]==[p%total for p in products]
    for s in w:
        assert not [c.coordinate for row in s for c in row if c.data_type=='e'],s.title
    assert all(row[1].value=='OK' for row in w['入力チェック'].iter_rows(min_row=2))
    for allocation_name,amount,weights,expected_allocation in ALLOCATION_CASES:
        if name==allocation_name:
            assert [w['月別配分'].cell(2,j).value for j in range(2,14)]==expected_allocation
    rows=calculate(plan)
    for i,r in enumerate(rows,2):
        for col,key in [('B','opening'),('C','income'),('D','expense'),('E','cash'),('F','total'),('G','shortage')]:
            assert w['月次CF'][f'{col}{i}'].value==r[key],(name,i,col)
        for j,a in enumerate(plan.accounts,9):
            assert w['月次CF'].cell(i,j).value==r['accounts'][a.id]
    for i,(year,r) in enumerate(annualize(rows).items(),2):
        for col,key in [('C','income'),('D','expense'),('E','cash'),('F','total')]:
            assert w['年次CF'][f'{col}{i}'].value==r[key]
    shortages=[r for r in rows if r['shortage']]
    assert w['不足サマリ']['B3'].value==max(r['shortage'] for r in rows)
    if shortages:
        assert w['不足サマリ']['B2'].value.date()==date(*shortages[0]['month'],1)
    else:
        assert w['不足サマリ']['B2'].value=='なし'


def test_unrecalculated_file_is_not_a_pass(tmp_path,monkeypatch):
    generate(demo(),tmp_path/'basic.xlsx')
    monkeypatch.setenv('LIFEPLAN_RECALCULATED_DIR',str(tmp_path))
    with pytest.raises(AssertionError):
        test_recalculated_excel('basic',demo(),2200000)


def test_old_generation_is_rejected(tmp_path,monkeypatch):
    path=tmp_path/'basic.xlsx'
    w=generate(demo(),path)
    w['_検証メタ']['B2']='old-generator'
    w.save(path)
    monkeypatch.setenv('LIFEPLAN_RECALCULATED_DIR',str(tmp_path))
    with pytest.raises(AssertionError,match='旧版'):
        test_recalculated_excel('basic',demo(),2200000)


def invalid_workbook(path,sheet,changes):
    w=generate(demo(),path)
    for cell,value in changes.items(): w[sheet][cell]=value
    w.save(path)


@pytest.mark.parametrize('name,sheet,changes,check_row',INVALID_WORKBOOK_CASES)
def test_excel_invalid_input_warning(name,sheet,changes,check_row):
    directory=os.environ.get('LIFEPLAN_RECALCULATED_DIR')
    if not directory:
        pytest.skip('未検証: Excel再計算済みの入力エラーケースが必要です')
    w=load_workbook(Path(directory)/f'{name}.xlsx',data_only=True)
    for cell,value in changes.items(): assert w[sheet][cell].value==value
    assert (w['_検証メタ']['B1'].value,w['_検証メタ']['B2'].value)==verification_identity(demo())
    assert str(w['入力チェック'].cell(check_row,2).value).startswith('入力エラー')
    assert str(w['不足サマリ']['B5'].value).startswith('入力エラー')


def test_powershell_51_file_contract():
    raw=Path('tools/recalculate_excel.ps1').read_bytes()
    assert raw.startswith(b'\xef\xbb\xbf'),'Windows PowerShell 5.1向けUTF-8 BOMが必要です'
    text=raw.decode('utf-8-sig')
    param_block=text.split('$ErrorActionPreference',1)[0]
    assert 'Join-Path' not in param_block
    assert '$PSScriptRoot' not in param_block


if __name__=='__main__':
    output=Path('output/recalc_input')
    all_cases=list(cases())
    desired={f'{name}.xlsx' for name,plan,expected in all_cases}|{f'{name}.xlsx' for name,sheet,changes,row in INVALID_WORKBOOK_CASES}
    # Remove obsolete generated cases only if their previous manifest hash still matches.
    previous=output/'manifest.json'
    if previous.exists():
        for entry in json.loads(previous.read_text(encoding='utf-8')):
            file=output/entry['file']
            if file.name not in desired and file.parent==output and file.exists():
                if hashlib.sha256(file.read_bytes()).hexdigest()!=entry['sha256']:
                    raise ValueError(f'変更された旧ケースは削除しません: {file}')
                file.unlink()
    manifest=[]
    for name,plan,expected in all_cases:
        path=output/f'{name}.xlsx'
        generate(plan,path)
        manifest.append(dict(file=path.name,expected_final_cash=expected,
            expected_salary_allocation=allocate(plan.rules[0].annual,plan.rules[0].weights),
            expected_monthly_income=[r['income'] for r in calculate(plan)],
            sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    for name,sheet,changes,row in INVALID_WORKBOOK_CASES:
        path=output/f'{name}.xlsx'
        invalid_workbook(path,sheet,changes)
        manifest.append(dict(file=path.name,expected_input_error=True,
            sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    (output/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print(f'{len(manifest)} cases generated in {output}')
