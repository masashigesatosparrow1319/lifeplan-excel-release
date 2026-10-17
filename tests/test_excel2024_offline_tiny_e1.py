"""SYNTHETIC_TEST_ONLY: no actual E1/Windows files, processes, or Excel runs.

The real auditor, real read-only XLSX checker and guard predicates are exercised.
Only temporary artificial XML receives literal test caches. No runtime/cache
verifier is replaced by a passing mock.
"""
import copy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
import sys
import zipfile
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from tools.excel2024_validation import SCHEMA, E1_DIAGNOSTIC, E1_PROTOCOL
from tools.excel2024_validation import offline_tiny_e1 as audit, windows_preflight as pre, cases

UTC = '2030-01-01T00:00:00+00:00'
RID = 'SYNTHETIC_TEST_ONLY'
WPID = 700001  # synthetic Python PID; never queried
EPID = 700002  # synthetic Excel PID; never queried
OLD = {'診断!B1': 5, '診断!B2': 6, '診断!B3': 11}
NEW = {'診断!B1': 10, '診断!B2': 21, '診断!B3': 31}


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def get(path): return json.loads(path.read_text(encoding='utf-8'))
def sha(path): return hashlib.sha256(path.read_bytes()).hexdigest()


def rewrite(path, mutate):
    with zipfile.ZipFile(path) as z: parts = {n: z.read(n) for n in z.namelist()}
    mutate(parts)
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, raw in parts.items(): z.writestr(n, raw)


def synthetic_cache(path):
    def patch(parts):
        name = 'xl/worksheets/sheet1.xml'; root = ET.fromstring(parts[name])
        for c in root.findall('.//m:c', cases.NS):
            value = {'A1': 7, 'B1': 10, 'B2': 21, 'B3': 31}.get(c.get('r'))
            if value is not None:
                v = c.find('m:v', cases.NS)
                if v is None: v = ET.SubElement(c, '{' + cases.NS['m'] + '}v')
                c.attrib.pop('t', None); v.text = str(value)
        parts[name] = ET.tostring(root)
    rewrite(path, patch)


def make_packet(tmp_path):
    root = tmp_path / 'SYNTHETIC_TEST_ONLY'; root.mkdir()
    directory = root / 'tiny3'; directory.mkdir()
    for part in ('inputs', 'work', 'results'):
        (directory / part).mkdir()
        shutil.copyfile(pre.REPO / 'output/s1_tiny_diagnostic/tiny_calc.xlsx', directory / part / audit.FILE)
    synthetic_cache(directory / 'results' / audit.FILE)
    manifest = dict(run_id=RID, case_id='tiny3', file=audit.FILE, sha256=audit.ORIGINAL_SHA, formula_count=3,
                    probes=[['診断', c, v] for c, v in cases.SPEC['expected_after'].items()],
                    edits=[['診断', 'A1', 7], ['診断', 'A2', 3]], challenges=[])
    save(directory / 'case.json', manifest)
    files = audit.TRUSTED_FILES.copy()
    source = dict(files=files, source_sha256=audit.digest(json.dumps(files, sort_keys=True).encode()), commit=None,
                  dependencies={'openpyxl': '3.1.5', 'pytest': '8.4.2', 'psutil': '7.0.0', 'pywin32': '311'})
    facts = dict(pid=EPID, created=datetime(2030, 1, 1, tzinfo=timezone.utc).timestamp(),
                 path=r'C:\SYNTHETIC_TEST_ONLY\EXCEL.EXE', version='16.0',
                 exe_sha256='1'*64, file_version='16.0.20430.20182', build='20430.0',
                 installation=dict(product_release_ids=['HomeBusiness2024Retail']))
    owner = {k: facts[k] for k in ('pid', 'created', 'path')}
    profile = dict(schema=SCHEMA, product='Microsoft Office Home and Business 2024', licensed=True,
                   ui_confirmed_on_owned_pid=True, facts=facts)
    request = dict(run_id=RID, case_id='tiny3', diagnostic=E1_DIAGNOSTIC, profile=profile,
                   manifest_sha256=sha(directory / 'case.json'), source_sha256=source['source_sha256'])
    save(directory / 'request.json', request)
    records = []
    def emit(name, **fields):
        records.append(dict(schema=SCHEMA, run_id=RID, case_id='tiny3', worker_pid=WPID,
                            utc=UTC, sequence=len(records)+1, event=name, **fields))
    def call(name): emit(name+'.before_call', timeout_seconds=180 if name=='final.full_rebuild' else 60); emit(name+'.after_call', elapsed_seconds=0.0)
    def snapshot_emit(item):
        for key in ('state_before','values','mode','state_after'):call(item['phase']+'.'+key)
        emit('snapshot',sample=item)
    def wait():
        emit('calculation.wait_started',timeout_seconds=180)
        call('calculation_state')
        emit('calculation.observed',state=0,observation=1,elapsed_seconds=0.0)
        emit('calculation.wait_done',observations=1,state_counts={'0':1},elapsed_seconds=0.0)
    def sample(phase, mode=-4105, old=False, state=0):
        values = OLD.copy() if old else NEW.copy()
        return dict(phase=phase, mode=mode, state_before=state, state=state, values=values,
                    differences=cases.compare_probes(manifest, lambda s, c: values[s+'!'+c]))
    contexts = []
    def context(phase, mode=-4105, old=False, state=0):
        item = dict(phase=phase, utc=UTC, force_full=False, sample=sample('e1.'+phase, mode, old, state),
                    workbooks=[dict(name=audit.FILE, path=str(directory / ('results' if phase=='reopened' else 'work') / audit.FILE),
                                    read_only=phase=='reopened')])
        for field in ('force_full','workbooks'):call('e1.'+phase+'.'+field)
        contexts.append(item); snapshot_emit(item['sample']); emit('e1.context', evidence=item)
    for name in ('connect', 'ownership'): call(name)
    emit('ownership.verified', identity=owner)
    call('identity'); emit('identity.verified', identity=facts)
    for name in ('settings', 'open.identity', 'open', 'read_only'): call(name)
    samples = [sample('opened', old=True)]; snapshot_emit(samples[0]); context('opened', old=True)
    for name in ('manual', 'manual.read', 'edit'): call(name)
    context('before_flag', -4135, old=True, state=2)
    for name in ('final.force_full', 'final.force_full_read'): call(name)
    emit('calculation.protocol', label='final', force_full_observed=False, protocol=E1_PROTOCOL)
    context('after_flag', -4135, old=True, state=2)
    call('final.full_rebuild'); context('after_full_rebuild', -4135)
    call('final.automatic_restore'); call('final.automatic_read')
    emit('calculation.automatic_observed', label='final', mode=-4105)
    snapshot_emit(sample('final.before_wait')); context('after_automatic'); wait()
    samples.append(sample('before_save')); snapshot_emit(samples[-1])
    for name in ('save.identity', 'save_as', 'close', 'reopen.identity', 'reopen'): call(name)
    context('reopened'); wait(); samples.append(sample('reopened')); snapshot_emit(samples[-1])
    for name in ('reopened.close', 'cleanup.diagnostics.identity','cleanup.observe.Ready','cleanup.observe.Interactive',
                 'cleanup.observe.DisplayAlerts','cleanup.quit', 'cleanup.release'): call(name)
    worker = dict(schema=SCHEMA, run_id=RID, case_id='tiny3', status='SAVED_REOPENED',
                  request_sha256=sha(directory/'request.json'), manifest_sha256=sha(directory/'case.json'),
                  first_error=None, cleanup_errors=[], progress_log_errors=[], identity=facts,
                  owned_pid=EPID, observed_process=owner, observed_identity=facts, completed_utc=UTC,
                  full_rebuild_calls=1, samples=samples,
                  saved_sha256=sha(directory/'results'/audit.FILE),
                  e1=dict(diagnostic=E1_DIAGNOSTIC, protocol=E1_PROTOCOL, evidence=contexts))
    save(directory / 'worker.json', worker)
    runtime = copy.deepcopy(worker)
    runtime.update(worker_pid=WPID, worker_still_running=False, force_termination=False, last_progress=records[-1],
                   residual_observation=dict(outcome='PROCESS_EXITED', samples=[owner.copy()]))
    parent = dict(schema=SCHEMA, mode='e1', diagnostic=E1_DIAGNOSTIC, protocol=E1_PROTOCOL,
                  run_id=RID, status='DIAGNOSTIC_PASS', exit_code=0,
                  diagnostic_only=True, diagnostic_recalculation_verified=True, gate2_acceptance=False,
                  product_acceptance=False, excel_recalculation_verified=False, force_termination=False, fixture_kind='SYNTHETIC_TEST_ONLY',
                  python='3.12.10 SYNTHETIC_TEST_ONLY', python_bits=64, os='Windows SYNTHETIC_TEST_ONLY', source=source,
                  stages=[dict(case_id='tiny3', status='PASS', runtime=runtime,
                               residual=dict(same_owned_process_remains=False),
                               verification=dict(saved_sha256=worker['saved_sha256'])),
                          dict(case_id='functions24', status='NOT_RUN'), dict(case_id='annual1', status='NOT_RUN')],
                  unit_tests=dict(returncode=0, junit='unit-tests.xml', log='unit-tests.txt', passed=1, skipped=35, failed=0, errors=0))
    save(root / 'run.json', parent)
    (directory / 'progress.jsonl').write_text(''.join(json.dumps(r,ensure_ascii=False)+'\n' for r in records),encoding='utf-8')
    (directory / 'worker-console.txt').write_text('',encoding='utf-8')
    (root / 'unit-tests.xml').write_text('<testsuites><testsuite tests="36" failures="0" errors="0" skipped="35"/></testsuites>',encoding='utf-8')
    (root / 'unit-tests.txt').write_text('SYNTHETIC_TEST_ONLY: 1 passed, 35 skipped',encoding='utf-8')
    return root


@pytest.fixture
def packet(tmp_path): return make_packet(tmp_path)


def execute(root, output=None, **kwargs):
    return audit.execute(root, output or root.parent / 'NEW_AUDIT', sha(root/'run.json'),
                         pre.source_fingerprint()['release_source_sha256'], owner_confirmed=True, acl_confirmed=True, **kwargs)


def change(root, which, mutate, bind_worker=True):
    path = root / ('run.json' if which=='parent' else 'tiny3/'+which+'.json')
    value = get(path); mutate(value); save(path,value)
    if which=='worker' and bind_worker:
        p=get(root/'run.json'); old=p['stages'][0]['runtime']; old.update(copy.deepcopy(value)); save(root/'run.json',p)


def progress(root, mutate):
    path=root/'tiny3/progress.jsonl'; values=[json.loads(s) for s in path.read_text(encoding='utf-8').splitlines()]
    mutate(values)
    path.write_text(''.join(json.dumps(s,ensure_ascii=False)+'\n' for s in values),encoding='utf-8')


def test_real_readonly_auditor_normal_synthetic_only(packet):
    before=audit.inventory(packet,time_start())
    result=execute(packet)
    assert result['status']=='OFFLINE_E1_REVERIFICATION_PASS',result
    assert result['input_label']=='SYNTHETIC_TEST_ONLY'
    assert result['gate2_acceptance'] is False and result['current_process_state']=='NOT_CHECKED'
    assert result['checks']['python_relation']=='VERIFIED'
    assert result['python_evidence_basis']=='LEGACY_DIRECT_SOURCE_AND_PROGRESS'
    assert audit.inventory(packet,time_start())==before
    report=packet.parent/'NEW_AUDIT/offline-e1-report.json'
    assert sha(report)==report.with_suffix('.sha256').read_text(encoding='ascii').strip()
    assert RID not in json.dumps(result) or result['input_label']==RID
    assert execute(packet)['status']=='BLOCKED'


def time_start():
    import time
    return time.monotonic()


@pytest.mark.parametrize('field,value', [('mode','smoke'),('diagnostic','WRONG'),('protocol','WRONG'),
                                       ('status','FAIL'),('exit_code',1),('exit_code',False),
                                       ('gate2_acceptance',True),('diagnostic_only',False),('python_bits',32)])
def test_parent_identity_changes_rejected(packet,field,value):
    change(packet,'parent',lambda p:p.update({field:value}))
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('kind', ['map','map_self_signed','commit','dependencies','digest','foreign_stage'])
def test_source_or_stage_cannot_self_certify(packet,kind):
    def mutate(p):
        if kind=='foreign_stage': p['stages'][1]['status']='PASS';return
        s=p['source']
        if kind.startswith('map'):
            s['files']['src/lifeplan.py']='0'*64
            if kind=='map_self_signed': s['source_sha256']=audit.digest(json.dumps(s['files'],sort_keys=True).encode())
        else: s['commit' if kind=='commit' else 'dependencies' if kind=='dependencies' else 'source_sha256']='WRONG'
    change(packet,'parent',mutate)
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('field,value', [('full_rebuild_calls',2),('full_rebuild_calls',True),('status','FAIL'),
                                       ('cleanup_errors',['ERROR']),('progress_log_errors',['ERROR']),
                                       ('owned_pid',99),('request_sha256','0'*64),('manifest_sha256','0'*64)])
def test_worker_fields_bound_to_real_validator(packet,field,value):
    change(packet,'worker',lambda p:p.update({field:value}))
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('phase',range(6))
@pytest.mark.parametrize('kind',['flag','mode','values','missing'])
def test_all_six_observations_checked(packet,phase,kind):
    def mutate(p):
        e=p['e1']['evidence']
        if kind=='missing':e.pop(phase)
        elif kind=='flag':e[phase]['force_full']=0
        elif kind=='mode':e[phase]['sample']['mode']=123
        else:e[phase]['sample']['values']['診断!B1']=999
    change(packet,'worker',mutate)
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('part',['run.json','tiny3/case.json','tiny3/request.json','tiny3/worker.json',
                                 'tiny3/progress.jsonl','tiny3/inputs/tiny_calc.xlsx','tiny3/work/tiny_calc.xlsx',
                                 'tiny3/results/tiny_calc.xlsx','tiny3/worker-console.txt','unit-tests.xml','unit-tests.txt'])
def test_missing_evidence_blocks(packet,part):
    (packet/part).unlink()
    # The missing root run prevents an external baseline in the helper too.
    result=(audit.execute(packet,packet.parent/'NEW_AUDIT','0'*64,pre.source_fingerprint()['release_source_sha256'],
                          owner_confirmed=True,acl_confirmed=True) if part=='run.json' else execute(packet))
    assert result['status']=='BLOCKED',result


@pytest.mark.parametrize('kind',['sequence','truncate','context','snapshot','runid','rebuild','ownership','bad_json'])
def test_progress_full_binding(packet,kind):
    path=packet/'tiny3/progress.jsonl'
    if kind=='bad_json':path.write_text('{broken',encoding='utf-8')
    else:
        def mutate(v):
            if kind=='sequence':v[0]['sequence']=999
            elif kind=='truncate':v.pop()
            elif kind=='runid':v[0]['run_id']='FOREIGN'
            elif kind=='rebuild':v[0]['event']='final.full_rebuild.after_call'
            elif kind=='ownership':v[next(i for i,e in enumerate(v) if e['event']=='ownership.verified')]['identity']['created']=0
            elif kind=='snapshot':v[next(i for i,e in enumerate(v) if e['event']=='snapshot')]['sample']['values']['診断!B1']=99
            else:v[next(i for i,e in enumerate(v) if e['event']=='e1.context')]['evidence']['force_full']=True
        progress(packet,mutate)
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('kind',['no_exit','remains','wrong_fingerprint','child_unproven','unfinished'])
def test_exit_and_legacy_python_equivalence_not_assumed(packet,kind):
    def mutate(p):
        r=p['stages'][0]['runtime']
        if kind=='no_exit':r.pop('residual_observation')
        elif kind=='remains':r['residual_observation']['outcome']='OWNED_PROCESS_REMAINS'
        elif kind=='wrong_fingerprint':r['residual_observation']['samples'][0]['created']=0
        elif kind=='unfinished':r['worker_still_running']=True
        else:r['worker_pid']=WPID+100
    change(packet,'parent',mutate)
    result=execute(packet)
    assert result['status']!='OFFLINE_E1_REVERIFICATION_PASS'
    if kind=='child_unproven':
        assert result['status']=='BLOCKED' and result['checks']['python_relation']=='UNSUPPORTED'


def rebind_saved_hash(root):
    value=sha(root/'tiny3/results'/audit.FILE)
    change(root,'worker',lambda w:w.update(saved_sha256=value))
    change(root,'parent',lambda p:p['stages'][0]['verification'].update(saved_sha256=value))


def saved_edit(root,cell=None,value=None,*,style=False,sheet='診断'):
    path=root/'tiny3/results'/audit.FILE
    w=load_workbook(path)
    if sheet not in w.sheetnames:w.create_sheet(sheet)
    s=w[sheet]
    if style:s[cell].fill=PatternFill('solid',fgColor='FFFF00')
    else:s[cell]=value
    w.save(path);w.close()
    synthetic_cache(path)
    if cell=='A1' and sheet=='診断':
        def patch(parts):
            key='xl/worksheets/sheet1.xml';r=ET.fromstring(parts[key])
            c=next(c for c in r.findall('.//m:c',cases.NS) if c.get('r')=='A1')
            c.find('m:v',cases.NS).text=str(value)
            parts[key]=ET.tostring(r)
        rewrite(path,patch)
    rebind_saved_hash(root)


@pytest.mark.parametrize('cell,value',[('A1',2),('A2',None),('A2',False),('A2',0),('A2','3'),
                                       ('C1',False),('C2',0),('C3',datetime(2030,1,2)),('AF64','outside')])
def test_all_nonformula_positions_and_approved_edit_checked(packet,cell,value):
    saved_edit(packet,cell,value)
    result=execute(packet)
    assert result['status']=='OFFLINE_E1_REVERIFICATION_FAIL',result


@pytest.mark.parametrize('cell',['C1','AF64'])
def test_style_only_union_growth_is_allowed_within_budget(packet,cell):
    saved_edit(packet,cell,style=True)
    assert execute(packet)['status']=='OFFLINE_E1_REVERIFICATION_PASS'


def test_new_sheet_values_are_not_ignored(packet):
    saved_edit(packet,'A1',False,sheet='追加')
    assert execute(packet)['status']=='OFFLINE_E1_REVERIFICATION_FAIL'


@pytest.mark.parametrize('kind',['wrong_cache','missing_cache','formula_change','formula_loss','formula_add',
                                 'manual','error','huge_dimension','hidden_cell','broken_zip','duplicate_zip','dtd'])
def test_cache_formula_mode_xml_and_resource_checks_are_real(packet,kind):
    path=packet/'tiny3/results'/audit.FILE
    if kind=='broken_zip':path.write_bytes(b'NOT_AN_XLSX')
    elif kind=='duplicate_zip':
        with zipfile.ZipFile(path,'a') as z:
            with pytest.warns(UserWarning,match='Duplicate'):z.writestr('xl/workbook.xml',b'<duplicate/>')
    elif kind in ('formula_add','error'):
        saved_edit(packet,'C3','=1' if kind=='formula_add' else '#DIV/0!')
    else:
        def patch(parts):
            key='xl/workbook.xml' if kind=='manual' else 'xl/worksheets/sheet1.xml'
            r=ET.fromstring(parts[key])
            if kind=='manual':r.find('m:calcPr',cases.NS).set('calcMode','manual')
            elif kind in ('huge_dimension','hidden_cell'):
                r.find('m:dimension',cases.NS).set('ref','A1:XFD1048576' if kind=='huge_dimension' else 'A1:A1')
            elif kind=='dtd':parts[key]=b'<!DOCTYPE a []>'+parts[key];return
            else:
                c=next(c for c in r.findall('.//m:c',cases.NS) if c.get('r')=='B1')
                if kind=='formula_change':c.find('m:f',cases.NS).text='A1-A2'
                elif kind=='formula_loss':c.remove(c.find('m:f',cases.NS))
                else:c.find('m:v',cases.NS).text=None if kind=='missing_cache' else '999'
            parts[key]=ET.tostring(r)
        rewrite(path,patch)
    rebind_saved_hash(packet)
    result=execute(packet)
    assert result['status']!='OFFLINE_E1_REVERIFICATION_PASS',result


@pytest.mark.parametrize('which',['manifest','request','input','saved','worker','progress','unit_xml'])
def test_file_bytes_sha_truncation_and_malformed_json_rejected(packet,which):
    paths={'manifest':'tiny3/case.json','request':'tiny3/request.json','input':'tiny3/inputs/tiny_calc.xlsx',
           'saved':'tiny3/results/tiny_calc.xlsx','worker':'tiny3/worker.json',
           'progress':'tiny3/progress.jsonl','unit_xml':'unit-tests.xml'}
    path=packet/paths[which]
    path.write_bytes(path.read_bytes()[:20])
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('flag',['owner','acl'])
def test_owner_and_acl_not_assumed(packet,flag):
    value=audit.execute(packet,packet.parent/'NEW',sha(packet/'run.json'),pre.source_fingerprint()['release_source_sha256'],
                        owner_confirmed=flag!='owner',acl_confirmed=flag!='acl')
    assert value['status']=='BLOCKED' and not (packet.parent/'NEW').exists()


@pytest.mark.parametrize('which',['run','source'])
def test_independent_expected_pins_required(packet,which):
    r=audit.execute(packet,packet.parent/'NEW','0'*64 if which=='run' else sha(packet/'run.json'),
                    '0'*64 if which=='source' else pre.source_fingerprint()['release_source_sha256'],
                    owner_confirmed=True,acl_confirmed=True)
    assert r['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('limit',['MAX_FILE','MAX_TOTAL','MAX_ENTRIES','MAX_ARCHIVE','MAX_EXPANDED','MAX_ROWS','MAX_COLS','SECONDS'])
def test_resource_limits_block_without_office(packet,monkeypatch,limit):
    monkeypatch.setattr(audit,limit,0)
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


def test_output_inside_original_and_existing_audit_never_written(packet):
    assert execute(packet,packet/'new')['status']=='BLOCKED'
    existing=packet.parent/'OLD_AUDIT';existing.mkdir();(existing/'offline-e1-report.json').write_text('{}',encoding='utf-8')
    assert execute(packet,existing/'new')['status']=='BLOCKED'
    assert not (existing/'new').exists()


@pytest.mark.parametrize('kind',['symlink','hardlink','synthetic_symlink','synthetic_reparse'])
def test_links_rejected_and_targets_unchanged(packet,kind,monkeypatch):
    import stat
    from types import SimpleNamespace

    target=packet/'tiny3/worker-console.txt'
    original=target.read_bytes()
    replacement=packet.parent/'UNTOUCHED';replacement.write_bytes(original)
    synthetic=kind.startswith('synthetic_') or (kind=='symlink' and sys.platform=='win32')
    observed=[]
    if synthetic:
        # Keep a real directory entry: inventory must actually reach this file.
        real_lstat=Path.lstat
        def link_attributes(path,*args,**kwargs):
            value=real_lstat(path,*args,**kwargs)
            if path==target:
                observed.append(path)
                return SimpleNamespace(
                    st_mode=(stat.S_IFLNK | stat.S_IMODE(value.st_mode)
                             if kind!='synthetic_reparse' else value.st_mode),
                    st_nlink=value.st_nlink,
                    st_file_attributes=0x400 if kind=='synthetic_reparse' else 0)
            return value
        monkeypatch.setattr(Path,'lstat',link_attributes)
    elif kind=='symlink':
        target.unlink()
        target.symlink_to(replacement)
    else:
        import os
        target.unlink()
        os.link(replacement,target)
    result=execute(packet)
    assert result['status']=='BLOCKED'
    if kind!='hardlink':
        assert result['reason']=='LINK_OR_REPARSE'
    if synthetic:
        assert observed
    assert target.read_bytes()==original
    assert replacement.read_bytes()==original
    assert not (packet.parent/'NEW_AUDIT').exists()


def test_front_back_inventory_detects_race_without_writing_original(packet,monkeypatch):
    original=audit.inventory; n=[0]
    def changed(root,start):
        data=original(root,start);n[0]+=1
        if n[0]==2:data['run.json']='SIMULATED_CONCURRENT_CHANGE'
        return data
    monkeypatch.setattr(audit,'inventory',changed)
    assert execute(packet)['status']=='BLOCKED'
    assert not (packet.parent/'NEW_AUDIT').exists()


@pytest.mark.parametrize('event,args', [('subprocess.Popen',('git',['git','rev-parse','HEAD'],None,None)),
                                       ('subprocess.Popen',('EXCEL.EXE',['EXCEL.EXE'],None,None)),
                                       ('import',('pythoncom',None,None,None,None)),
                                       ('ctypes.dlopen',('ole32.dll',)),('os.system',('anything',))])
def test_dedicated_offline_guard_disallows_even_git_or_office(event,args):
    guard=audit.OfflineGuard()
    with pytest.raises(RuntimeError):guard(event,args)
    assert guard.violations==['OFFLINE_FORBIDDEN_OPERATION']


def test_cli_installs_guard_and_does_not_print_private_arguments(monkeypatch,capsys):
    monkeypatch.setattr(sys,'argv',['offline_tiny_e1','--source-run','PRIVATE_PATH','--output','PRIVATE_OUT',
                                  '--expected-run-sha','a'*64,'--expected-source-sha','b'*64])
    hooks=[];monkeypatch.setattr(sys,'addaudithook',hooks.append)
    assert audit.main()==2
    value=json.loads(capsys.readouterr().out)
    assert 'PRIVATE' not in json.dumps(value) and hooks and isinstance(hooks[0],audit.OfflineGuard)


def test_no_live_executors_are_called_by_real_inspection(packet,monkeypatch):
    from tools.excel2024_validation import run,e1,worker,session,preopen_input
    def forbidden(*args,**kwargs):raise AssertionError('LIVE_EXECUTOR_FORBIDDEN')
    for module,name in [(run,'execute'),(run,'supervise'),(e1,'execute'),(worker,'execute_case'),
                        (cases,'prepare'),(session.Session,'connect'),(preopen_input,'prepare')]:
        monkeypatch.setattr(module,name,forbidden)
    assert execute(packet)['status']=='OFFLINE_E1_REVERIFICATION_PASS'


def test_recorded_launcher_child_proof_supported_without_live_lookup(packet):
    def mutate(p):
        r=p['stages'][0]['runtime']
        supervisor=dict(pid=700000,ppid=0,created=1,path='SYNTHETIC_SUPERVISOR')
        launcher=dict(pid=WPID+100,ppid=700000,created=2,path='SYNTHETIC_LAUNCHER')
        child=dict(pid=WPID,ppid=WPID+100,created=3,path='SYNTHETIC_CHILD',sys_executable='SYNTHETIC_CHILD')
        r.update(worker_pid=launcher['pid'],python_process=child.copy(),
                 python_processes=dict(supervisor=supervisor,launcher=launcher,worker=child,
                                       worker_exit=dict(outcome='PROCESS_EXITED')))
    change(packet,'parent',mutate)
    result=execute(packet)
    assert result['status']=='OFFLINE_E1_REVERIFICATION_PASS',result
    assert result['checks']['python_relation']=='VERIFIED'


def test_recorded_launcher_child_proof_must_match_progress(packet):
    def mutate(p):
        r=p['stages'][0]['runtime']
        r['python_processes']={}
    change(packet,'parent',mutate)
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


@pytest.mark.parametrize('kind',['unit_counts','console','parent_error','flag','duplicate_json','xml_utf16_dtd'])
def test_additional_contract_boundaries(packet,kind):
    if kind=='unit_counts':change(packet,'parent',lambda p:p['unit_tests'].update(failed=1))
    elif kind=='console':(packet/'tiny3/worker-console.txt').write_text('SYNTHETIC_WARNING',encoding='utf-8')
    elif kind=='parent_error':change(packet,'parent',lambda p:p.update(first_error={'type':'ERROR'}))
    elif kind=='flag':change(packet,'parent',lambda p:p.update(force_termination=True))
    elif kind=='duplicate_json':
        path=packet/'tiny3/request.json';raw=path.read_text(encoding='utf-8');path.write_text(raw[:-1]+',"run_id":"OTHER"}',encoding='utf-8')
    else:
        def patch(parts):parts['xl/worksheets/sheet1.xml']=('<!DOCTYPE a []><a/>').encode('utf-16')
        rewrite(packet/'tiny3/results'/audit.FILE,patch);rebind_saved_hash(packet)
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'


def test_auditor_has_no_side_effectful_or_live_calls():
    import ast
    source=Path(audit.__file__).read_text(encoding='utf-8')
    tree=ast.parse(source)
    forbidden={'prepare','supervise','connect','execute_case','CalculateFullRebuild','SaveAs','source_identity','Process','pids'}
    calls={n.func.attr for n in ast.walk(tree) if isinstance(n,ast.Call) and isinstance(n.func,ast.Attribute)}
    assert not calls.intersection(forbidden)
    assert 'load_workbook(original, read_only=True' in source and 'load_workbook(saved, read_only=True' in source


def test_atomic_output_collision_cannot_overwrite_existing_data(packet,monkeypatch):
    original=Path.mkdir
    output=packet.parent/'NEW_AUDIT'
    def race(path,*args,**kwargs):
        if path==output:
            original(path)
            (path/'DO_NOT_TOUCH').write_text('PRESERVE',encoding='utf-8')
        return original(path,*args,**kwargs)
    monkeypatch.setattr(Path,'mkdir',race)
    result=execute(packet)
    assert result['status']=='BLOCKED' and (output/'DO_NOT_TOUCH').read_text(encoding='utf-8')=='PRESERVE'


def test_report_stdout_does_not_include_private_hashes_or_identifiers(packet,monkeypatch,capsys):
    # Execute the real checker through main; only audit-hook registration is
    # observed here to avoid a permanent inner hook during pytest. The outer
    # no-Excel integration run applies its real hook to every test operation.
    output=packet.parent/'CLI_AUDIT'
    monkeypatch.setattr(sys,'argv',['offline_tiny_e1','--source-run',str(packet),'--output',str(output),
                                  '--expected-run-sha',sha(packet/'run.json'),
                                  '--expected-source-sha',pre.source_fingerprint()['release_source_sha256'],
                                  '--owner-confirmed','--acl-confirmed'])
    hooks=[];monkeypatch.setattr(sys,'addaudithook',hooks.append)
    assert audit.main()==0
    text=capsys.readouterr().out
    assert str(packet) not in text and str(WPID) not in text and str(EPID) not in text
    assert sha(packet/'run.json') not in text and sha(packet/'tiny3/results'/audit.FILE) not in text
    assert hooks and json.loads(text)['status']=='OFFLINE_E1_REVERIFICATION_PASS'


def test_legacy_parent_progress_prefix_is_bound_without_discarding_tail(packet):
    rows=[json.loads(line) for line in (packet/'tiny3/progress.jsonl').read_text(encoding='utf-8').splitlines()]
    change(packet,'parent',lambda p:p['stages'][0]['runtime'].update(last_progress=rows[-2]))
    assert execute(packet)['status']=='OFFLINE_E1_REVERIFICATION_PASS'


def test_parent_progress_prefix_cannot_be_fabricated(packet):
    change(packet,'parent',lambda p:p['stages'][0]['runtime']['last_progress'].update(event='WRONG'))
    assert execute(packet)['status']=='OFFLINE_E1_REVERIFICATION_FAIL'


@pytest.mark.parametrize('kind',['missing','state','deadline','extra_calculate'])
def test_wait_protocol_and_closed_call_schema(packet,kind):
    def mutate(rows):
        if kind=='extra_calculate':
            rows[0]['event']='final.additional_calculate.before_call'
        else:
            entry=next(r for r in rows if r['event']=='calculation.observed')
            if kind=='missing':entry['event']='MISSING_WAIT_OBSERVATION'
            elif kind=='state':entry['state']=2
            else:entry['elapsed_seconds']=180
    progress(packet,mutate)
    assert execute(packet)['status']!='OFFLINE_E1_REVERIFICATION_PASS'
