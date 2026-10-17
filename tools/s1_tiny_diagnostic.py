"""Three ordinary formulas for observation only. Never writes formula caches."""
import argparse
import hashlib
import json
import ntpath
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.properties import CalcProperties

SCHEMA = 'S1-tiny-com-diagnostic-v1'
SPEC = dict(file='tiny_calc.xlsx', sheet='診断', formula_count=3,
            inputs_before={'A1': 2, 'A2': 3}, inputs_after={'A1': 7, 'A2': 3},
            formulas={'B1': '=A1+A2', 'B2': '=A1*A2', 'B3': '=SUM(B1:B2)'},
            expected_before={'B1': 5, 'B2': 6, 'B3': 11},
            expected_after={'B1': 10, 'B2': 21, 'B3': 31})


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def create(directory):
    directory = Path(directory).resolve()
    if directory.is_relative_to(Path(__file__).resolve().parents[1] / 'references'):
        raise ValueError('参照原本へ出力しません')
    if directory.exists() and any(directory.iterdir()):
        raise FileExistsError('新規空フォルダーを指定してください')
    directory.mkdir(parents=True, exist_ok=True)
    w = Workbook()
    s = w.active
    s.title = SPEC['sheet']
    for address, value in {**SPEC['inputs_before'], **SPEC['formulas']}.items():
        s[address] = value
    w.calculation = CalcProperties(calcMode='auto', fullCalcOnLoad=True)
    path = directory / SPEC['file']
    w.save(path)
    w.close()
    packet = dict(schema=SCHEMA, spec=SPEC, sha256=sha(path))
    (directory / 'manifest.json').write_text(json.dumps(packet, ensure_ascii=False, indent=2), encoding='utf-8')
    return packet


def classify(sample):
    values = sample['values']
    marker = all(values.get(k) == v for k, v in SPEC['inputs_after'].items())
    updated = all(values.get(k) == v for k, v in SPEC['expected_after'].items())
    if not marker:
        if all(values.get(k) == v for k, v in {**SPEC['inputs_before'], **SPEC['expected_before']}.items()):
            return 'baseline_values_matched'
        return 'input_marker_not_applied'
    if updated:
        return 'updated_values_state_done' if sample['state'] == 0 else 'updated_values_state_not_done'
    if all(values.get(k) == v for k, v in SPEC['expected_before'].items()):
        return 'old_values_observed'
    return 'unverified_values'


def check_cache(path):
    w = load_workbook(path, read_only=True, data_only=False)
    try:
        assert w.sheetnames == [SPEC['sheet']]
        formulas = {c.coordinate: c.value for row in w.active for c in row if c.data_type == 'f'}
        assert formulas == SPEC['formulas'], '診断数式が変更されています'
        assert all(w.active[k].value == v for k, v in SPEC['inputs_after'].items()), '診断入力が一致しません'
        assert w.calculation is None or w.calculation.calcMode in (None, 'auto'), '保存設定が自動ではありません'
    finally:
        w.close()
    w = load_workbook(path, read_only=True, data_only=True)
    try:
        actual = {k: w.active[k].value for k in SPEC['expected_after']}
        assert actual == SPEC['expected_after'], '未再計算または固定期待値不一致: ' + str(actual)
        assert not any(c.data_type == 'e' for row in w.active for c in row), 'Excelエラー値'
    finally:
        w.close()


def inspect(inputs, results):
    inputs, results = Path(inputs), Path(results)
    assert inputs.resolve() != results.resolve(), '原本と結果を分けてください'
    packet = json.loads((inputs / 'manifest.json').read_text(encoding='utf-8'))
    assert packet['schema'] == SCHEMA and packet['spec'] == SPEC
    assert sha(inputs / SPEC['file']) == packet['sha256'], '原本の変更を検出しました'
    log = json.loads((results / 'tiny-session.json').read_text(encoding='utf-8-sig'))
    assert log['schema'] == SCHEMA and log['input_sha256'] == packet['sha256']
    trace = log['trace']
    assert trace['identity_verified'] is True
    facts = trace['connection']
    assert facts['excel_version'] == '14.0' and facts['file_major_version'] == 14
    assert facts['process_id'] == trace['launched_process_id'] > 0
    assert ntpath.normcase(facts['executable_path']) == ntpath.normcase(trace['expected_path'])
    observations = log['observations']
    report = dict(s1_acceptance=False, tiny_gate_ready=False, manual_warning_check_required=True, cache='not_saved',
                  observations=[dict(phase=s['phase'], state=s['state'], calculation=s['calculation'], finding=classify(s)) for s in observations])
    evidence = results / SPEC['file']
    if evidence.exists():
        check_cache(evidence)
        report['cache'] = 'fixed_values_matched'
    if observations:
        last = observations[-1]
        report['tiny_gate_ready'] = bool(log['observation_completed'] and not log['cleanup_errors']
            and not log['progress_log_errors'] and log.get('first_error') is None
            and last['state'] == 0 and last['calculation'] == -4105
            and classify(last) == 'updated_values_state_done' and report['cache'] == 'fixed_values_matched')
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default='output/s1_tiny_diagnostic')
    parser.add_argument('--results', help='指定時は保存済み診断記録/キャッシュを読み取るだけ')
    args = parser.parse_args()
    print(json.dumps(inspect(args.output, args.results) if args.results else create(args.output), ensure_ascii=False, indent=2))
