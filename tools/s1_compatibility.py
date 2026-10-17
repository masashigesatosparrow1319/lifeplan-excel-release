"""Isolated Excel 2010 S1 gate. Does not call or change the P2-A generator.

Python models verify the algorithms, never evaluate Excel or populate caches.
"""
import argparse
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

SLOTS = 250
MONTHS = ('2026-01', '2026-02')
FILES = ('01_function_gate.xlsx', '02_monthly_cases.xlsx')
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


@dataclass(frozen=True)
class Entry:
    kind: str
    slot: int
    key: str
    amounts: tuple
    enabled: bool = True

    @property
    def index(self):
        return self.slot - 1 + (50 if self.kind == 'Event' else 0)


@dataclass(frozen=True)
class Case:
    id: str
    name: str
    entries: tuple
    expected: tuple

    def vectors(self):
        keys = [''] * SLOTS
        amounts = [[0] * SLOTS for _ in MONTHS]
        for e in self.entries:
            if e.enabled:
                keys[e.index] = e.key.upper()
                for m in range(2):
                    amounts[m][e.index] = e.amounts[m]
        return keys, amounts


E = Entry
CASES = (
    Case('N01', 'すべて別キー', (E('Rule', 1, 'A', (100, 100)), E('Event', 1, 'B', (200, 200)), E('Event', 2, 'C', (300, 300))), (0, 0)),
    Case('N02', '同キーでも別月', (E('Event', 1, 'K', (100, 0)), E('Event', 2, 'k', (0, 200))), (0, 0)),
    Case('N03', '同キーの中間0円・月ごと正額1件', (E('Event', 1, 'K', (100, 0)), E('Event', 2, 'K', (0, 0)), E('Event', 3, 'k', (0, 200))), (0, 0)),
    Case('N04', '同キーも全額0円', (E('Event', 1, 'K', (0, 0)), E('Event', 2, 'K', (0, 0))), (0, 0)),
    Case('N05', '正額の同キーを無効化', (E('Event', 1, 'K', (100, 100)), E('Event', 2, 'k', (100, 200), False)), (0, 0)),
    Case('N06', 'Ruleとイベントの独立支払い', (E('Rule', 1, 'RULE_COST', (100, 100)), E('Event', 1, 'EVENT_COST', (200, 200))), (0, 0)),
    Case('E01', '同月同キー3件正額', (E('Event', 1, 'K', (100, 100)), E('Event', 2, 'K', (200, 0)), E('Event', 3, 'K', (300, 0))), (1, 0)),
    Case('E02', '別キーを跨ぐ重複・中間0円も保持', (E('Event', 1, 'K', (100, 0)), E('Event', 2, 'K', (0, 0)), E('Event', 3, 'X', (50, 0)), E('Event', 4, 'K', (200, 0))), (1, 0)),
    Case('E03', 'Ruleとイベントの同月重複', (E('Rule', 1, 'K', (100, 100)), E('Event', 1, 'k', (200, 0))), (1, 0)),
    Case('E04', 'イベント200番目で同月重複', (E('Event', 1, 'K', (100, 0)), E('Event', 200, 'k', (200, 200))), (1, 0)),
)
GATE_IDS = (1, 2, 1, 0, 250, 250)
# name, amounts, fixed duplicate flag, fixed positive-group count
GATES = (
    ('別グループ', (100, 200, 0, 0, 0, 0), 0, 2),
    ('非連続同グループ', (100, 0, 200, 0, 0, 0), 1, 2),
    ('全0円・0ビン除外', (0, 0, 0, 0, 0, 0), 0, 0),
    ('最終ビン250の重複', (0, 0, 0, 0, 100, 200), 1, 2),
    ('最終ビン250に1件', (0, 0, 0, 0, 0, 100), 0, 1),
    ('2グループが同時重複', (100, 0, 200, 0, 300, 400), 1, 4),
)


def group_ids(keys):
    first = {}
    result = []
    for i, key in enumerate(keys, 1):
        key = key.upper()
        if key:
            first.setdefault(key, i)
        result.append(first[key] if key else 0)
    return result


def previous_indices(keys):
    previous, result = {}, []
    for i, key in enumerate(keys, 1):
        key = key.upper()
        result.append(previous.get(key, 0) if key else 0)
        if key:
            previous[key] = i
    return result


def legacy_counts(keys, amounts):
    """Same recurrence as current 重複確認; zero amounts retain prior counts."""
    result = []
    for previous, amount in zip(previous_indices(keys), amounts):
        result.append((result[previous - 1] if previous else 0) + int(amount > 0))
    return result


def frequency_model(ids, amounts):
    """Mathematical FREQUENCY bins 0..250; not an Excel evaluator."""
    bins = [0] * 252
    for identity, amount in zip(ids, amounts):
        bins[identity if amount > 0 else 0] += 1
    return bins


def s1_flag(keys, amounts):
    return int(any(n > 1 for n in frequency_model(group_ids(keys), amounts)[1:]))


def source_sha():
    return hashlib.sha256(Path(__file__).read_bytes()).hexdigest()


def book(kind):
    w = Workbook()
    w.active.title = '確認'
    w.calculation = CalcProperties(calcId=0, calcMode='auto', fullCalcOnLoad=True, forceFullCalc=True)
    bins = w.create_sheet('ビン')
    bins.append(['ビン上限', '0ビンを除くマスク'])
    for i in range(251):
        bins.append([i, int(i != 0)])
    bins.cell(253, 2, 1)  # FREQUENCY overflow bin
    bins.sheet_state = 'hidden'
    meta = w.create_sheet('_meta')
    meta.append(['version', 's1-compatibility-v1'])
    meta.append(['source_sha256', source_sha()])
    meta.append(['kind', kind])
    meta.sheet_state = 'hidden'
    return w


def histogram(data):
    return f"FREQUENCY({data},'ビン'!$A$2:$A$252)"


def duplicate_formula(data):
    return f'=IF(SUMPRODUCT(--({histogram(data)}>1),\'ビン\'!$B$2:$B$253)>0,1,0)'


def decorate(w):
    for s in w:
        s.freeze_panes = 'B2'
        for c in s[1]:
            c.font = Font(bold=True)
        s.column_dimensions['A'].width = 46
    w.active = 0


def gate_book():
    w = book('gate')
    s = w['確認']
    s['A1'] = 'S1関数単体ゲート：通常のF9のみ。CSE入力不要かを実機確認する'
    s['A2'] = '固定期待値と実数式を比較。生成時の数式キャッシュは空。0=重複なし、1=重複あり'
    s['A4'] = '数値グループID（MATCHを使わない対照）'
    for c, value in enumerate(GATE_IDS, 2):
        s.cell(4, c, value)
    for c, text in enumerate(['ケース', '金額1', '金額2', '金額3', '金額4', '金額5', '金額6', '固定期待', 'S1実測', '判定', '正額件数期待', '正額件数実測', '件数判定'], 1):
        s.cell(7, c, text)
    for row, (name, amounts, expected, count) in enumerate(GATES, 8):
        s.cell(row, 1, name)
        for c, amount in enumerate(amounts, 2):
            s.cell(row, c, amount)
        data = f'$B$4:$G$4*--(B{row}:G{row}>0)'
        for c, value in {8: expected, 9: duplicate_formula(data), 10: f'=IF(I{row}=H{row},"一致","不一致")', 11: count, 12: f'=SUMPRODUCT({histogram(data)},\'ビン\'!$B$2:$B$253)', 13: f'=IF(L{row}=K{row},"一致","不一致")'}.items():
            s.cell(row, c, value)
    decorate(w)
    return w


def case_book():
    w = book('cases')
    s = w['確認']
    s['A1'] = 'S1月次ゲート：現行の累積方式とS1方式を固定期待値で比較'
    s['A2'] = '0=重複なし、1=重複あり。異常系で1を検出した場合も、固定期待に合えば「一致」'
    for c, text in enumerate(['ケース', '説明', '年月', '固定期待', '現行実測', 'S1実測', '照合'], 1):
        s.cell(3, c, text)
    inputs = w.create_sheet('入力')
    inputs.append(['ケース', '種別', '入力枠番号', '費用キー', '有効0/1', '2026-01円', '2026-02円'])
    matrix = w.create_sheet('月次金額')
    legacy = w.create_sheet('現行判定')
    trial = w.create_sheet('S1判定')
    for ci, case in enumerate(CASES):
        start, base = 2 + 250 * ci, 1 + 6 * ci
        entries = {e.index: e for e in case.entries}
        for slot in range(250):
            row, c = start + slot, col(slot + 2)
            entry = entries.get(slot)
            kind = 'Rule' if slot < 50 else 'Event'
            inputs.append([case.id, kind, slot + 1 if slot < 50 else slot - 49, entry.key if entry else None, int(entry.enabled) if entry else 1, entry.amounts[0] if entry else 0, entry.amounts[1] if entry else 0])
            matrix.cell(base, slot + 2, f'=IF(AND(\'入力\'!E{row}=1,\'入力\'!D{row}<>""),UPPER(\'入力\'!D{row}),"")')
            matrix.cell(base + 1, slot + 2, f'{kind}:{slot+1 if slot<50 else slot-49}')
            for target in (legacy, trial):
                target.cell(base, slot + 2, f"='月次金額'!{c}{base}")
            prev = col(slot + 1)
            formula = '=0' if slot == 0 else f'=IF({c}{base}="",0,IFERROR(LOOKUP(2,1/($B${base}:{prev}${base}={c}{base}),COLUMN($B${base}:{prev}${base})-1),0))'
            legacy.cell(base + 1, slot + 2, formula)
            trial.cell(base + 1, slot + 2, f'=IF({c}{base}="",0,MATCH({c}{base},$B${base}:$IQ${base},0))')
            for m in range(2):
                mr = base + 2 + m
                matrix.cell(mr, slot + 2, f'=IF(AND(\'入力\'!E{row}=1,\'入力\'!D{row}<>""),\'入力\'!{col(6+m)}{row},0)')
                indicator = f'IF(\'月次金額\'!{c}{mr}>0,1,0)'
                chain = indicator if slot == 0 else f'IF({c}${base+1}>0,INDEX($B{mr}:{prev}{mr},1,{c}${base+1})+{indicator},{indicator})'
                legacy.cell(mr, slot + 2, '=' + chain)
        for m, month in enumerate(MONTHS):
            mr, sr = base + 2 + m, 4 + 2 * ci + m
            matrix.cell(mr, 1, case.id + ':' + month)
            legacy.cell(mr, 1, case.id + ':' + month)
            trial.cell(mr, 1, case.id + ':' + month)
            legacy.cell(mr, 252, f'=IF(MAX(B{mr}:IQ{mr})>1,1,0)')
            trial.cell(mr, 252, duplicate_formula(f'$B${base+1}:$IQ${base+1}*--(\'月次金額\'!B{mr}:IQ{mr}>0)'))
            values = [case.id, case.name, month, case.expected[m], f"='現行判定'!IR{mr}", f"='S1判定'!IR{mr}", f'=IF(AND(E{sr}=D{sr},F{sr}=D{sr}),"一致","不一致")']
            for c, value in enumerate(values, 1):
                s.cell(sr, c, value)
        for target in (matrix, legacy, trial):
            target.cell(base, 1, case.id + ':キー')
            target.cell(base + 1, 1, '入力枠' if target is matrix else '前方同キー位置' if target is legacy else '数値グループID')
    for column, lo, hi in [('E', 0, 1), ('F', 0, 10**9), ('G', 0, 10**9)]:
        validation = DataValidation(type='whole', operator='between', formula1=lo, formula2=hi)
        validation.showErrorMessage = True
        inputs.add_data_validation(validation)
        validation.add(f'{column}2:{column}2501')
    for row in inputs.iter_rows(min_row=2, min_col=4, max_col=7):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='FFF2CC')
    for hidden in (matrix, legacy, trial):
        hidden.sheet_state = 'hidden'
    decorate(w)
    s.freeze_panes = 'D4'
    s.column_dimensions['B'].width = 48
    return w


def stats(path):
    formulas, shared, arrays, caches, longest = 0, 0, 0, 0, 0
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                root = ET.fromstring(z.read(name))
                for cell in root.findall('.//m:c', NS):
                    f, v = cell.find('m:f', NS), cell.find('m:v', NS)
                    if f is not None:
                        formulas += 1
                        shared += int(f.attrib.get('t') == 'shared')
                        arrays += int(f.attrib.get('t') == 'array')
                        caches += int(v is not None and bool(v.text))
                        longest = max(longest, 1 + len(f.text or ''))
    return dict(bytes=Path(path).stat().st_size, formula_cells=formulas, shared=shared, arrays=arrays, nonempty_caches=caches, max_formula_chars=longest)


def prepare(directory):
    directory = Path(directory)
    if directory.resolve().is_relative_to(Path(__file__).resolve().parents[1] / 'references'):
        raise ValueError('参照原本へ出力しません')
    if directory.exists() and any(directory.iterdir()):
        raise FileExistsError('新規空フォルダーを指定してください')
    directory.mkdir(parents=True, exist_ok=True)
    manifest = []
    for name, factory in zip(FILES, (gate_book, case_book)):
        path = directory / name
        w = factory()
        w.save(path)
        w.close()
        record = dict(file=name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(), **stats(path))
        assert record['shared'] == record['arrays'] == record['nonempty_caches'] == 0
        assert record['max_formula_chars'] <= 8192
        manifest.append(record)
    (directory / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    return manifest


def expected_cells(kind):
    """Expected caches for inspection only; never written into an xlsx."""
    if kind == 'gate':
        for row, (_, _, flag, count) in enumerate(GATES, 8):
            for c, value in [('I', flag), ('J', '一致'), ('L', count), ('M', '一致')]:
                yield '確認', f'{c}{row}', value
        return
    for ci, case in enumerate(CASES):
        keys, amounts = case.vectors()
        base = 1 + 6 * ci
        for slot, (key, previous, group) in enumerate(zip(keys, previous_indices(keys), group_ids(keys)), 2):
            c = col(slot)
            for sheet in ('月次金額', '現行判定', 'S1判定'):
                yield sheet, f'{c}{base}', key
            yield '現行判定', f'{c}{base+1}', previous
            yield 'S1判定', f'{c}{base+1}', group
        for m in range(2):
            row, sr = base + 2 + m, 4 + 2 * ci + m
            counts = legacy_counts(keys, amounts[m])
            assert int(max(counts) > 1) == s1_flag(keys, amounts[m]) == case.expected[m]
            for slot, (amount, count) in enumerate(zip(amounts[m], counts), 2):
                yield '月次金額', f'{col(slot)}{row}', amount
                yield '現行判定', f'{col(slot)}{row}', count
            yield '現行判定', f'IR{row}', case.expected[m]
            yield 'S1判定', f'IR{row}', case.expected[m]
            for c, value in [('E', case.expected[m]), ('F', case.expected[m]), ('G', '一致')]:
                yield '確認', f'{c}{sr}', value


def snapshot(w):
    return {(s.title, c.coordinate): (c.data_type, c.value) for s in w for row in s for c in row if c.value is not None}


def canonical_formula(value, names):
    # Excel may remove optional sheet-name quotes. Preserve operators/$/strings.
    for name in names:
        value = value.replace("'" + name + "'!", name + '!')
    return value


def verify_file(original, result, kind):
    if Path(original).resolve() == Path(result).resolve():
        raise ValueError('入力と再計算結果を別フォルダーにしてください')
    if not Path(result).is_file():
        raise FileNotFoundError(result)
    metrics = stats(result)
    if metrics['arrays']:
        raise AssertionError('通常式ゲートに配列数式タグを検出しました')
    source = load_workbook(original, read_only=True, data_only=False)
    output = load_workbook(result, read_only=True, data_only=False)
    try:
        assert source['_meta']['B2'].value == source_sha(), '旧診断コードの入力です'
        assert source.sheetnames == output.sheetnames, 'シート構造が異なります'
        a, b = snapshot(source), snapshot(output)
        assert a.keys() == b.keys(), 'セル/数式の欠落・追加を検出しました'
        for key, (typ, value) in a.items():
            actual_type, actual = b[key]
            if typ == 'f':
                assert actual_type == 'f', f'{key}: 数式が失われています'
                assert canonical_formula(value, source.sheetnames) == canonical_formula(actual, output.sheetnames), f'{key}: 数式の変更'
            else:
                assert actual == value, f'{key}: 固定入力/期待値の変更'
    finally:
        source.close()
        output.close()
    caches = load_workbook(result, read_only=False, data_only=True)
    checked = 0
    try:
        assert not [(s.title, c.coordinate, c.value) for s in caches for row in s for c in row if c.data_type == 'e'], 'Excel数式エラー'
        for sheet, address, expected in expected_cells(kind):
            assert b.get((sheet, address), (None,))[0] == 'f', f'{sheet}!{address}: 数式が必要です'
            value = caches[sheet][address].value
            if expected == '':
                assert value in ('', None), f'{sheet}!{address}: 空白キーの不一致'
            else:
                assert value is not None, f'{sheet}!{address}: 未再計算（キャッシュなし）'
                assert value == expected, f'{sheet}!{address}: {value!r} != {expected!r}'
            checked += 1
    finally:
        caches.close()
    return dict(file=Path(result).name, numeric_comparison='matched', checked_formula_cells=checked,
                excel_version='manual confirmation required', repair_warning='manual confirmation required', **metrics)


def verify(directory, inputs):
    manifest = json.loads((Path(inputs) / 'manifest.json').read_text(encoding='utf-8'))
    assert [r['file'] for r in manifest] == list(FILES), '入力manifestが異なります'
    for entry in manifest:
        assert hashlib.sha256((Path(inputs) / entry['file']).read_bytes()).hexdigest() == entry['sha256'], '入力原本の更新を検出しました'
    return [verify_file(Path(inputs) / name, Path(directory) / name, kind) for name, kind in zip(FILES, ('gate', 'cases'))]


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default='output/s1_compatibility')
    parser.add_argument('--verify')
    args = parser.parse_args()
    result = verify(args.verify, args.output) if args.verify else prepare(args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))
