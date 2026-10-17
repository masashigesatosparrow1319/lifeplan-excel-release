"""Inspect real Excel caches for separately identified baseline/S1 packets."""
import argparse
from collections import Counter
import hashlib
import json
import ntpath
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from src import p2a_s1 as trial
from src import p2a_excel as baseline
from src.lifeplan import annualize
from src.p2a import calculate
from tools.p2a_s1_cases import SCHEMA, input_fingerprint


def saved_calculation_mode(path):
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read('xl/workbook.xml'))
    settings = root.find('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}calcPr')
    # OOXML's default is automatic when calcMode (or calcPr) is omitted.
    mode = settings.attrib.get('calcMode', 'auto') if settings is not None else 'auto'
    assert mode == 'auto', '保存済みxlsxの計算方法が自動ではありません'
    return mode


def formula_digest(path):
    w = load_workbook(path, read_only=True, data_only=False)
    h = hashlib.sha256()
    n = 0
    try:
        for s in w:
            for row in s:
                for c in row:
                    if c.data_type != 'f':
                        continue
                    f = c.value
                    if not isinstance(f, str):
                        raise AssertionError('通常式以外を検出しました')
                    for name in w.sheetnames:
                        f = f.replace("'" + name + "'!", name + '!')
                    h.update(json.dumps([s.title, c.coordinate, f], ensure_ascii=False).encode() + b'\n')
                    n += 1
    finally:
        w.close()
    return n, h.hexdigest()


def verify_entry(entry, inputs, results):
    original, result = Path(inputs) / entry['file'], Path(results) / entry['file']
    if original.resolve() == result.resolve():
        raise ValueError('入力と結果を別フォルダーにしてください')
    assert trial.sha(original) == entry['sha256'], '入力原本の更新を検出しました'
    if not result.is_file():
        raise FileNotFoundError(result)
    metrics = trial.stats(result)
    assert not metrics['arrays'], '配列数式を検出しました'
    assert formula_digest(original) == formula_digest(result), '数式が消失/変更されています'
    assert input_fingerprint(result) == entry['expected_input_sha256'], '別入力、または指定編集が未実施です'
    saved_mode = saved_calculation_mode(result)
    w = load_workbook(result, read_only=True, data_only=True)
    try:
        assert w['_検証メタ']['B1'].value == entry['generation_input_id'], '別生成入力です'
        wanted = trial.source_id() if entry['variant'] == 's1' else trial.base_source_id()
        assert w['_検証メタ']['B2'].value == wanted, '旧コード/別方式の結果です'
        if entry['variant'] == 's1':
            assert w['_検証メタ']['B3'].value == trial.VARIANT
            assert w['_検証メタ']['B4'].value == trial.base_source_id()
        if entry['invalid']:
            try:
                calculate(baseline.read_plan(result))
            except ValueError:
                pass
            else:
                raise AssertionError('異常入力をPythonが拒否しませんでした')
            assert w['入力チェック']['B2'].value == '入力/重複エラー', '異常入力が未検出/未再計算です'
            assert w['不足サマリ']['B5'].value == '入力/重複エラー'
        else:
            plan = baseline.read_plan(result)
            assert baseline.identity(plan)[0] == entry['normalized_input_id']
            rows = calculate(plan)
            assert rows[-1]['cash'] == entry['expected_cash']
            assert rows[-1]['total'] == entry['expected_total']
            assert w['入力チェック']['B2'].value == 'OK', '未再計算/入力エラーです'
            cached = list(w['月次CF'].iter_rows(min_row=2, values_only=True))
            assert len(cached) == len(rows)
            for reference, values in zip(rows, cached):
                assert (values[0].year, values[0].month) == reference['month']
                assert list(values[1:7]) == [reference[k] for k in ('opening', 'income', 'expense', 'cash', 'total', 'shortage')]
                assert list(values[8:8+len(plan.accounts)]) == [reference['accounts'][a.id] for a in plan.accounts]
            yearly = list(w['年次CF'].iter_rows(min_row=2, values_only=True))
            assert len(yearly) == len(annualize(rows))
            month_counts = Counter(r['month'][0] for r in rows)
            for (year, ref), values in zip(annualize(rows).items(), yearly):
                assert values[0] == year
                assert values[1] == month_counts[year]
                assert list(values[2:6]) == [ref[k] for k in ('income', 'expense', 'cash', 'total')]
            first = next((r for r in rows if r['shortage']), None)
            assert w['不足サマリ']['B3'].value == max(r['shortage'] for r in rows)
            if first:
                cached_first = w['不足サマリ']['B2'].value
                assert (cached_first.year, cached_first.month) == first['month']
            else:
                assert w['不足サマリ']['B2'].value == 'なし'
            assert not [(s.title, c.coordinate, c.value) for s in w for row in s for c in row if c.data_type == 'e'], 'Excel数式エラー'
        if entry['duplicate_months'] is not None:
            expected = {tuple(m) for m in entry['duplicate_months']}
            for row in w['重複確認'].iter_rows(min_row=3, max_col=252, values_only=True):
                when = row[0]
                if when is None:  # constant-bin rows beyond a short horizon
                    continue
                assert row[251] is not None, '重複フラグ未再計算'
                assert row[251] == int((when.year, when.month) in expected), '重複フラグ不一致'
    finally:
        w.close()
    return dict(file=entry['file'], variant=entry['variant'], numeric_comparison='matched',
                saved_calculation_mode=saved_mode, excel_version='see Windows/manual log', warning_check='manual confirmation required', **metrics)


def verify_excel2010_log(path, entries, packet):
    """Validate recorded runtime identity; never infer version from caches."""
    log = json.loads(Path(path).read_text(encoding='utf-8-sig'))
    assert log['schema'] == 'P2A-S1-excel2010-diagnostic-v1'
    assert len(entries) == 1, '診断版ログは1ケース専用です'
    assert log['source_id'] == packet['source_id']
    assert log['baseline_source_id'] == packet['baseline_source_id']
    assert log['trace']['identity_verified'] is True
    assert log['trace']['mode'] == 'single_case' and log['trace']['stage'] == 'single_case_complete'
    assert isinstance(log.get('result'), dict), '完了結果がありません（中断/未完了を含む）'
    assert log.get('body_complete', True) is True, '処理未完了です'
    assert not log.get('progress_log_errors', []), '進行ログの保存失敗があります'
    assert log['trace'].get('recalculation_protocol') == 'full_rebuild_auto_wait_v1', '旧待機順序の結果は新試験へ流用しません'
    assert log['trace'].get('automatic_restore_verified') is True, '自動復帰が未確認です'
    assert log['trace'].get('calculation_state_done_observed') is True, '状態0到達が未確認です'
    for field in ('automatic_calculation_observed', 'pre_save_calculation_observed', 'reopened_calculation_observed'):
        assert log['trace'].get(field) == -4105, '自動計算の実測不一致: ' + field
    assert log['result']['status'] == 'saved_and_reopened_numeric_check_pending'
    assert log['result']['file'] == entries[0]['file']
    assert log['cleanup_status'] == 'no_errors_recorded' and not log['cleanup_errors']
    for facts in (log['trace']['connection'], log['result']['connection']):
        assert facts['excel_version'] == '14.0' and facts['file_major_version'] == 14, 'Excel 2010以外です'
        assert facts['process_id'] == log['trace']['launched_process_id'] > 0, 'PID不一致'
        assert ntpath.normcase(ntpath.normpath(facts['executable_path'])) == ntpath.normcase(ntpath.normpath(log['trace']['expected_path'])), '実行パス不一致'
    return 'Excel 2010 runtime identity recorded; warnings require manual confirmation'


def verify_packet(inputs, results, file=None, excel2010_log=None):
    packet = json.loads((Path(inputs) / 'manifest.json').read_text(encoding='utf-8'))
    assert packet['schema'] == SCHEMA
    assert packet['source_id'] == trial.source_id() and packet['baseline_source_id'] == trial.base_source_id(), '旧試作packetです'
    entries = packet['entries']
    if file:
        entries = [e for e in entries if e['file'] == file]
        if len(entries) != 1:
            raise ValueError('指定した1ファイルがmanifestにありません')
    if excel2010_log:
        verify_excel2010_log(excel2010_log, entries, packet)
    return [verify_entry(e, inputs, results) for e in entries]


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--inputs', required=True)
    parser.add_argument('--results', required=True)
    parser.add_argument('--file', help='1ケース試運転のファイル名。省略時は全件照合')
    parser.add_argument('--excel2010-log', help='今回のランナーが記録した1ケースの接続ログ')
    args = parser.parse_args()
    print(json.dumps(verify_packet(args.inputs, args.results, args.file, args.excel2010_log), ensure_ascii=False, indent=2))
