"""Separate S1 regression/performance packets. Edits run in Excel via manifest.

Reference xlsx files made temporarily in Python are uncached inputs, never results.
"""
import argparse
import gc
from dataclasses import replace
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import shutil
import tempfile
import time
import zipfile

from openpyxl import load_workbook
from src.lifeplan import Person, Account, Rule, allocate, annualize, months, demo
from src.p2a import P2Plan, CashEvent, calculate, corrected_price, occurs, from_p1
from src import p2a_excel as baseline
from src import p2a_s1 as trial
from tests.p2a_cases import CASES, Case, one
from tools.s1_compatibility import group_ids, frequency_model, legacy_counts

SCHEMA = 'P2A-S1-packet-v1'
ASSETS = Path(__file__).resolve().parents[1] / 'docs' / 'performance_baseline_ffe49ef'


def canonical(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%dT%H:%M:%S')
    if type(value) is float and value.is_integer():
        return int(value)
    return value


def input_fingerprint(path):
    w = load_workbook(path, read_only=True, data_only=False)
    values = []
    try:
        for name, end, column in [('設定', 2, 2), ('人物', w['人物'].max_row, 4),
                                 ('口座', w['口座'].max_row, 6), ('収支', 51, 20), ('イベント入力', 201, 19)]:
            for row in w[name].iter_rows(min_row=2, max_row=end, max_col=column):
                values.append([name, [canonical(c.value) for c in row]])
    finally:
        w.close()
    return hashlib.sha256(json.dumps(values, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()


def patches(case):
    result = []
    for sheet, edits in case.edits.items():
        for cell, value in edits.items():
            typ = 'clear' if value is None else 'date' if isinstance(value, (date, datetime)) else 'number' if isinstance(value, (float, int)) else 'text'
            result.append(dict(sheet=sheet, cell=cell, type=typ,
                               value=value.isoformat()[:10] if typ == 'date' else value))
    return result


def apply_patches(w, edits):
    for p in edits:
        value = date.fromisoformat(p['value']) if p['type'] == 'date' else p['value']
        w[p['sheet']][p['cell']] = value


def source_vectors(path):
    """Dedup-only mathematical projection using unchanged P1/P2-A primitives."""
    plan = baseline.read_plan(path)
    plan.validate()
    w = load_workbook(path, read_only=True)
    try:
        ids = [r[0] if any(v is not None for v in (*r[1:7], r[19])) else ''
               for r in w['収支'].iter_rows(min_row=2, max_row=51, max_col=20, values_only=True)]
        ids += [r[0] if r[18] == 1 and any(v is not None for v in r[1:18]) else ''
                for r in w['イベント入力'].iter_rows(min_row=2, max_row=201, max_col=19, values_only=True)]
    finally:
        w.close()
    index = {source.upper(): i for i, source in enumerate(ids) if source}
    keys = [''] * 250
    rkeys = plan.keys()
    for r in plan.rules:
        keys[index[r.id.upper()]] = rkeys[r.id]
    for e in plan.events:
        keys[index[e.id.upper()]] = e.key
    allocations = {r.id: allocate(r.annual, r.weights) for r in plan.rules}
    prices, rows = {}, []
    gids = group_ids(keys)
    for ym in months(plan.start, plan.end):
        amounts = [0] * 250
        for r in plan.rules:
            if r.start <= ym <= r.end:
                amounts[index[r.id.upper()]] = allocations[r.id][ym[1] - 1]
        for e in plan.events:
            if occurs(e, ym):
                amounts[index[e.id.upper()]] = corrected_price(e, ym, prices)
        old = int(max(legacy_counts(keys, amounts)) > 1)
        s1 = int(any(n > 1 for n in frequency_model(gids, amounts)[1:]))
        assert old == s1
        rows.append(dict(month=ym, duplicate=s1, amounts=amounts))
    return ids, rows


def extra_cases():
    p = from_p1(demo(), [one()])
    add = dict(B201='末尾追加', E201=date(2026, 8, 1), H201=100000, I201='cash', M201='はい')
    yield Case('slot200_add', p, 1100000, {'イベント入力': add})
    middle = {k.replace('201', '100'): v for k, v in add.items()}
    yield Case('middle_gap_add', p, 1100000, {'イベント入力': middle})
    yield Case('slot200_disable', p, 1200000, {'イベント入力': {**add, 'S201': 0}})
    yield Case('same_key_other_month', p, 1100000, {'イベント入力': {**add, 'L201': 'E001'}})
    yield Case('slot200_duplicate', p, edits={'イベント入力': {**add, 'E201': date(2026, 7, 1), 'L201': 'E001'}}, invalid=True)
    yield Case('noncontiguous_zero_duplicate', p, edits={'イベント入力': {
        **add, 'E201': date(2026, 7, 1), 'L201': 'E001',
        'B3': '同キー0円', 'E3': date(2026, 7, 1), 'H3': 0, 'I3': 'cash', 'M3': 'はい', 'L3': 'E001'}}, invalid=True)
    yield Case('rule_event_zero', p, 2200000, {'イベント入力': {'H2': 0, 'L2': 'salary'}})


ALL_CASES = CASES + list(extra_cases())


def empty_directory(path):
    p = Path(path).resolve()
    if p.is_relative_to(Path(__file__).resolve().parents[1] / 'references'):
        raise ValueError('参照原本へ出力しません')
    if p.exists() and any(p.iterdir()):
        raise FileExistsError('新規空フォルダーを指定してください')
    p.mkdir(parents=True, exist_ok=True)
    return p


def expected_record(path, case):
    record = dict(expected_input_sha256=input_fingerprint(path), expected_cash=case.expected_cash, invalid=case.invalid)
    try:
        actual = baseline.read_plan(path)
        record['normalized_input_id'] = baseline.identity(actual)[0]
        rows = calculate(actual)
    except ValueError:
        if not case.invalid:
            raise
        record['normalized_input_id'] = None
    else:
        assert not case.invalid and rows[-1]['cash'] == case.expected_cash
        record['expected_total'] = rows[-1]['total']
        record['first_shortage'] = next((r['month'] for r in rows if r['shortage']), None)
    try:
        _, vectors = source_vectors(path)
        record['duplicate_months'] = [r['month'] for r in vectors if r['duplicate']]
    except ValueError:
        record['duplicate_months'] = None
    return record


def prepare_regression(directory, cases=ALL_CASES):
    out = empty_directory(directory)
    entries = []
    with tempfile.TemporaryDirectory(prefix='s1-regression-') as temporary:
        for case in cases:
            original, expected = Path(temporary) / 'normal.xlsx', Path(temporary) / 'edited.xlsx'
            w = baseline.generate(case.plan, original)
            all_edits = patches(case)
            setup = [p for p in all_edits if p['sheet'] in ('イベント入力', '収支') and p['cell'].startswith('A')]
            # ID columns are protected in the product. Inject only these invalid
            # fixtures before export; ordinary user edits still run in Excel.
            apply_patches(w, setup)
            if setup:
                w.save(original)
            target = out / (case.name + '.xlsx')
            trial.transform(original, target)
            edits = [p for p in all_edits if p not in setup]
            apply_patches(w, all_edits)
            w.save(expected)
            w.close()
            del w
            gc.collect()
            entry = dict(file=target.name, variant='s1', sha256=trial.sha(target),
                         generation_input_id=baseline.identity(case.plan)[0], edits=edits, setup_edits=setup,
                         monitor_seconds=case.monitor_seconds, **expected_record(expected, case))
            entries.append(entry)
    packet = dict(schema=SCHEMA, source_id=trial.source_id(), baseline_source_id=trial.base_source_id(), entries=entries)
    (out / 'manifest.json').write_text(json.dumps(packet, ensure_ascii=False, indent=2), encoding='utf-8')
    return packet


def plan_from_record(record):
    p = record['plan']
    persons = [Person(**dict(x, birth=date.fromisoformat(x['birth']))) for x in p['persons']]
    rules = [Rule(**dict(x, start=tuple(x['start']), end=tuple(x['end']))) for x in p['rules']]
    fields = ('occurrence_month', 'first_month', 'last_month', 'base_month')
    events = [CashEvent(**dict(x, **{k: tuple(x[k]) if x[k] is not None else None for k in fields})) for x in p['events']]
    return P2Plan(tuple(p['start']), tuple(p['end']), persons, [Account(**x) for x in p['accounts']], rules, events, p['rule_keys'])


def unchanged_parts(before, after):
    with zipfile.ZipFile(before) as a, zipfile.ZipFile(after) as b:
        assert set(a.namelist()) == set(b.namelist())
        members = trial.sheet_members(a)
        changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
        assert set(changed) == {members['重複確認'], members['_検証メタ']}
        return changed


def prepare_performance(directory, baseline_directory=None, trial_outputs=None):
    out = empty_directory(directory)
    artifacts = empty_directory(trial_outputs) if trial_outputs else None
    entries, comparisons = [], []
    for name in ('real60', 'real70', 'maximum_busy'):
        record = json.loads((ASSETS / (name + '.json')).read_text(encoding='utf-8'))
        assert record['source_id'] == trial.base_source_id()
        plan = plan_from_record(record)
        original, target = out / (name + '-normal.xlsx'), out / (name + '-s1.xlsx')
        if baseline_directory:
            source = Path(baseline_directory) / record['xlsx']
            assert trial.sha(source) == record['sha256']
            shutil.copyfile(source, original)
            generation_seconds = None
        else:
            started = time.perf_counter()
            w = baseline.generate(plan, original)
            w.close()
            del w
            gc.collect()
            generation_seconds = time.perf_counter() - started
        started = time.perf_counter()
        comparison = trial.transform(original, target)
        comparison['python_transform_and_audit_seconds'] = time.perf_counter() - started
        comparison['baseline_python_generation_seconds'] = generation_seconds
        comparison['case'] = name
        comparison['changed_xml_parts'] = unchanged_parts(original, target)
        comparison['all_other_xml_bytes_identical'] = True
        old_plan, new_plan = baseline.read_plan(original), baseline.read_plan(target)
        old_rows, new_rows = calculate(old_plan), calculate(new_plan)
        assert old_rows == new_rows and annualize(old_rows) == annualize(new_rows)
        assert old_rows[-1]['cash'] == record['expected']['cash']
        ids, vectors = source_vectors(target)
        for vector, row in zip(vectors, old_rows):
            ledger = {x['source_id'].upper(): x['amount'] for x in row['ledger']}
            assert vector['amounts'] == [ledger.get(source.upper(), 0) if source else 0 for source in ids]
            assert vector['duplicate'] == 0
        comparison['python_reference'] = dict(months=len(old_rows), annual_rows=len(annualize(old_rows)),
            cash=old_rows[-1]['cash'], total=old_rows[-1]['total'], first_shortage=record['expected']['first_shortage'],
            monthly_annual_accounts='Python reference matched; Excel caches UNVERIFIED')
        comparison['excel_2010'] = 'UNVERIFIED: recalculation, caches, memory, save/reopen'
        comparisons.append(comparison)
        expected = expected_record(original, Case(name, plan, record['expected']['cash']))
        for variant, file in [('baseline', original), ('s1', target)]:
            entries.append(dict(file=file.name, variant=variant, sha256=trial.sha(file), edits=[],
                generation_input_id=record['generation_plan_id'], monitor_seconds=120 if name == 'maximum_busy' else 30, **expected))
        if artifacts:
            shutil.copyfile(target, artifacts / target.name)
    packet = dict(schema=SCHEMA, source_id=trial.source_id(), baseline_source_id=trial.base_source_id(), entries=entries)
    (out / 'manifest.json').write_text(json.dumps(packet, ensure_ascii=False, indent=2), encoding='utf-8')
    report = dict(notice='Actual generated sizes/formula counts; Python timings are NOT Excel timings.', cases=comparisons)
    for folder in (out, artifacts):
        if folder:
            (folder / 'comparison.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return packet


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--performance', action='store_true')
    parser.add_argument('--output')
    parser.add_argument('--baseline-directory')
    parser.add_argument('--trial-outputs')
    args = parser.parse_args()
    directory = args.output or ('output/p2a_s1_performance_input' if args.performance else 'output/p2a_s1_recalc_input')
    result = prepare_performance(directory, args.baseline_directory, args.trial_outputs) if args.performance else prepare_regression(directory)
    print('Prepared', len(result['entries']), 'S1-specific entries. Excel results remain unverified.')
