"""G2-014: bounded, read-only, tiny3 E1 evidence auditor. No live Office work.

Windows/real evidence execution needs separate PM/owner approval. Synthetic
fixtures exercise this code but are never evidence of Excel recalculation.
"""
from datetime import date, datetime
import copy
import hashlib
import json
import ntpath
import os
from pathlib import Path
import re
import stat
import sys
import time
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
from . import SCHEMA, E1_DIAGNOSTIC, E1_PROTOCOL
from . import cases, e1, windows_preflight as pre
from .common import GateError

CANDIDATE_COMMIT = 'ddc46d7ef89e2fa9add95e10df8a678f59f44329'
# Independently read from git objects of the candidate, not from audited logs.
TRUSTED_FILES = {
    'tools/excel2024_validation/__init__.py': 'f5b822efbf36f3aece28b2d8a530ea8c5aa687bc720d7a61a3f6263441f183d1',
    'tools/excel2024_validation/cases.py': 'e278a9ddfcf971cab4d8194392c80df79eb9fa1390239dab7f9a646210df26b4',
    'tools/excel2024_validation/common.py': 'ff7007686ca1b00191ee55323ce64850a9bafd3d2b694bd86dc1d000755db4c3',
    'tools/excel2024_validation/e1.py': '397072e391967cc39476628097609445ba4d607f958acc0e0b539bf739d77ff4',
    'tools/excel2024_validation/run.py': 'ff6f15ff537d317b5b78c4802c30d1401e07aaa9280bad26480293f7b61f67fb',
    'tools/excel2024_validation/session.py': 'b4bfdbe6df6117fef7f2d5acef3cbd6cfbd7228daf46ddf159bec979221c7aca',
    'tools/excel2024_validation/worker.py': 'b2767018d3c5154218863af939daa79b1f8f588c46ffd07920bf8e184784901c',
    'src/lifeplan.py': '0539dc4159f89705ac96ad98ddbd96e18e5d8e4f27d1e621739671db13090946',
    'src/p2a.py': '89bece20ed9c032f3c2de463221f3dc3c693fab1218f6b12661bfddaa455b57f',
    'src/p2a_excel.py': '6c11245444e07a2a74acaec531be261f38543020f9f8dc3b8e44c32008c358c2',
    'src/p2a_s1.py': 'c194354168923551a29708c91cb416e9b0ce462c899d141b722906b74cbe0f21',
    'tools/s1_tiny_diagnostic.py': '6715251e0daa6b7310f7130eea0f8010ebc350416a2faaf9abf3d5695f73e252',
    'tools/s1_compatibility.py': '7f333ca92f29b7254f83ffbe3571b77256d24f255827956aa9787d9709cd12da',
    'tools/p2a_s1_cases.py': 'f0c54ada79f0d75d131eef7e001b952ad40beb17de2d521b408030633ad10735',
    'tools/verify_p2a_s1.py': '7d48831c8507455c24c4d2cf5e27b0fe4e7c6e10e0443314d710950ccd94d5fe',
    'tests/p2a_cases.py': '93191a84df9f06fd60c14f0fdf54b8642334c2c381a15126be7431595074ee6a',
}
ORIGINAL_SHA = '38c28cbceebc6bf9645b6791819cf00dce790f6f3ed3bbeb1c404bfcfa6e3682'
FILE = 'tiny_calc.xlsx'
MAX_FILE, MAX_TOTAL, MAX_ENTRIES = 4 * 1024**2, 16 * 1024**2, 64
MAX_ARCHIVE, MAX_EXPANDED, MAX_ROWS, MAX_COLS, SECONDS = 1024**2, 4 * 1024**2, 64, 32, 30
NS = cases.NS
PHASES = ('opened', 'before_flag', 'after_flag', 'after_full_rebuild', 'after_automatic', 'reopened')


class EvidenceError(Exception):
    def __init__(self, code, blocked=False):
        super().__init__(code)
        self.blocked = blocked


def need(value, code, *, blocked=False):
    if not value:
        raise EvidenceError(code, blocked)


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def deadline(started):
    need(time.monotonic() - started <= SECONDS, 'AUDIT_TIME_LIMIT', blocked=True)


def safe_path(path, *, new=False):
    path = Path(path)
    need(path.is_absolute() and '..' not in path.parts and not str(path).startswith(('\\\\', '//')),
         'ABSOLUTE_LOCAL_PATH_REQUIRED', blocked=True)
    need(not path.is_relative_to(pre.REPO) and not any(
        p.casefold() in (pre.LEDGER, 'onedrive', 'dropbox') for p in path.parts), 'PROTECTED_PATH', blocked=True)
    for variable in ('OneDrive', 'OneDriveConsumer', 'OneDriveCommercial'):
        sync = os.environ.get(variable)
        need(not sync or not path.is_relative_to(Path(sync).absolute()), 'SYNC_PATH', blocked=True)
    for item in (path, *path.parents):
        if item == path and new and not os.path.lexists(item):
            continue
        s = item.lstat()
        need(not stat.S_ISLNK(s.st_mode) and not getattr(s, 'st_file_attributes', 0) & 0x400,
             'LINK_OR_REPARSE', blocked=True)
        need(stat.S_ISDIR(s.st_mode) or stat.S_ISREG(s.st_mode) and s.st_nlink == 1,
             'SPECIAL_OR_HARDLINK', blocked=True)
    return path


def read(path):
    path = safe_path(path)
    s = path.stat()
    need(stat.S_ISREG(s.st_mode) and s.st_size <= MAX_FILE and os.access(path, os.R_OK),
         'EVIDENCE_FILE_LIMIT_OR_ACCESS', blocked=True)
    with path.open('rb') as stream:
        raw = stream.read(MAX_FILE + 1)
    need(len(raw) <= MAX_FILE, 'EVIDENCE_FILE_LIMIT', blocked=True)
    return raw


def unique(pairs):
    value = {}
    for key, item in pairs:
        need(key not in value, 'DUPLICATE_JSON_KEY')
        value[key] = item
    return value


def parse(raw):
    def bad_constant(value): raise EvidenceError('NONFINITE_JSON')
    return json.loads(raw.decode('utf-8-sig'), object_pairs_hook=unique, parse_constant=bad_constant)


def inventory(root, started):
    hashes, pending, count, total = {}, [root], 0, 0
    while pending:
        with os.scandir(pending.pop()) as entries:
            for entry in entries:
                count += 1
                need(count <= MAX_ENTRIES, 'ENTRY_LIMIT', blocked=True)
                path = safe_path(Path(entry.path))
                if path.is_dir(): pending.append(path)
                else:
                    raw = read(path); total += len(raw)
                    need(total <= MAX_TOTAL, 'TOTAL_BYTES_LIMIT', blocked=True)
                    hashes[str(path.relative_to(root))] = digest(raw)
                deadline(started)
    return hashes


def xml(raw):
    try:
        text = raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        raise EvidenceError('XML_ENCODING_UNSUPPORTED', blocked=True) from None
    need('\x00' not in text and '<!DOCTYPE' not in text.upper() and '<!ENTITY' not in text.upper(), 'XML_DECLARATION')
    return ET.fromstring(text)


def bounded_book(path, started):
    raw = read(path)
    need(len(raw) <= MAX_ARCHIVE, 'XLSX_BYTES_LIMIT', blocked=True)
    with zipfile.ZipFile(path) as z:
        entries = z.infolist(); names = [p.filename for p in entries]
        need(len(entries) <= 32 and len(set(names)) == len(names)
             and sum(p.file_size for p in entries) <= MAX_EXPANDED, 'ZIP_LIMIT', blocked=True)
        need(not any(p.flag_bits & 1 or '..' in p.filename.split('/') or p.filename.startswith('/') for p in entries)
             and not any('externalLink' in n or 'vbaProject' in n or n == 'xl/connections.xml' for n in names), 'ZIP_STRUCTURE')
        count = 0
        for name in names:
            payload = z.read(name)  # CRC checked; all parts read, not extracted
            if name.endswith(('.xml', '.rels')):
                root = xml(payload)
                if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                    count += 1
                    dim = root.find('m:dimension', NS)
                    need(dim is not None, 'DIMENSION_MISSING', blocked=True)
                    box = range_boundaries(dim.get('ref', ''))
                    need(None not in box and 1 <= box[0] <= box[2] <= MAX_COLS
                         and 1 <= box[1] <= box[3] <= MAX_ROWS, 'DIMENSION_LIMIT', blocked=True)
                    last_row = 0
                    for row in root.findall('m:sheetData/m:row', NS):
                        r = int(row.get('r', '0')); col = 0
                        need(last_row < r <= MAX_ROWS, 'ROW_ORDER_OR_LIMIT', blocked=True); last_row = r
                        for cell in row.findall('m:c', NS):
                            cr, cc = coordinate_to_tuple(cell.get('r', ''))
                            need(cr == r and col < cc <= MAX_COLS and box[0] <= cc <= box[2]
                                 and box[1] <= cr <= box[3], 'HIDDEN_OR_DUPLICATE_CELL', blocked=True)
                            col = cc
                        deadline(started)
            deadline(started)
        need(1 <= count <= 4, 'SHEET_LIMIT', blocked=True)


def nonformula_union(original, saved, started):
    a = load_workbook(original, read_only=True, data_only=False)
    b = load_workbook(saved, read_only=True, data_only=False)
    try:
        for name in dict.fromkeys([*a.sheetnames, *b.sheetnames]):
            sheets = [w[name] if name in w.sheetnames else None for w in (a, b)]
            rows = max(s.max_row if s is not None else 1 for s in sheets)
            cols = max(s.max_column if s is not None else 1 for s in sheets)
            need(rows <= MAX_ROWS and cols <= MAX_COLS, 'UNION_LIMIT', blocked=True)
            streams = [s.iter_rows(min_row=1, max_row=rows, min_col=1, max_col=cols) if s is not None
                       else iter(()) for s in sheets]
            for r in range(1, rows + 1):
                for c, (left, right) in enumerate(zip(*(next(s, (None,) * cols) for s in streams), strict=True), 1):
                    if left is not None and left.data_type == 'f': continue
                    x, y = left.value if left is not None else None, right.value if right is not None else None
                    expected = 7 if (name, r, c) == ('診断', 1, 1) else x
                    need(expected == y and isinstance(expected, bool) == isinstance(y, bool)
                         and (not isinstance(expected, date) and not isinstance(y, date) or type(expected) is type(y))
                         and not (right is not None and right.data_type == 'f'), 'NONFORMULA_CHANGED')
                deadline(started)
    finally:
        a.close(); b.close()


def timestamp(value):
    t = datetime.fromisoformat(value)
    need(t.utcoffset() is not None, 'TIMEZONE_MISSING')
    return t.timestamp()


def historical_source(source):
    raw = source['files']
    need(type(raw) is dict, 'SOURCE_SCHEMA_UNSUPPORTED', blocked=True)
    normalized = {k.replace('\\', '/'): v for k, v in raw.items()}
    need(len(normalized) == len(raw) and normalized == TRUSTED_FILES, 'SOURCE_MAP_UNSUPPORTED', blocked=True)
    need(source.get('commit') in (None, CANDIDATE_COMMIT), 'SOURCE_COMMIT_UNSUPPORTED', blocked=True)
    need(digest(json.dumps(raw, sort_keys=True).encode()) == source['source_sha256'], 'SOURCE_DIGEST_CHANGED')
    deps = source['dependencies']
    need(deps == {'openpyxl': '3.1.5', 'pytest': '8.4.2', 'psutil': '7.0.0', 'pywin32': '311'},
         'HISTORICAL_DEPENDENCIES_UNSUPPORTED', blocked=True)


def evidence_progress(records, parent, worker, runtime, manifest, directory, checks):
    need(records, 'PROGRESS_MISSING', blocked=True)
    previous_time, pending, calls = 0, {}, []
    for index, event in enumerate(records, 1):
        need(event['schema'] == SCHEMA and type(event['sequence']) is int and event['sequence'] == index and event['run_id'] == parent['run_id']
             and event['case_id'] == 'tiny3', 'PROGRESS_ID_OR_SEQUENCE')
        now = timestamp(event['utc']); need(now >= previous_time, 'PROGRESS_TIME_ORDER'); previous_time = now
        name = event['event']
        need(not name.endswith('.failed'), 'RECORDED_CALL_FAILED')
        if name.endswith('.before_call'):
            label = name.removesuffix('.before_call')
            need(label not in pending, 'DUPLICATE_CALL'); pending[label] = index
        elif name.endswith('.after_call'):
            label = name.removesuffix('.after_call')
            need(label in pending, 'UNPAIRED_CALL'); pending.pop(label); calls.append(label)
    need(not pending, 'PROGRESS_TRUNCATED', blocked=True)
    # The legacy parent reads progress before polling the process, so a final
    # event can arrive between that read and process exit. Its recorded prefix
    # must still match exactly; the COMPLETE file is checked independently below.
    last = runtime.get('last_progress', {})
    n = last.get('sequence')
    need(type(n) is int and 1 <= n <= len(records) and last == records[n - 1], 'LAST_PROGRESS_BINDING')
    required = ['connect', 'ownership', 'identity', 'settings', 'open.identity', 'open', 'read_only',
                'manual', 'manual.read', 'edit', 'final.force_full', 'final.force_full_read',
                'final.full_rebuild', 'final.automatic_restore', 'final.automatic_read',
                'save.identity', 'save_as', 'close', 'reopen.identity', 'reopen', 'reopened.close',
                'cleanup.diagnostics.identity', 'cleanup.quit', 'cleanup.release']
    need([c for c in calls if c in required] == required, 'CALL_ORDER_OR_COUNT')
    snapshot_phases = ['opened', 'e1.opened', 'e1.before_flag', 'e1.after_flag', 'e1.after_full_rebuild',
                       'final.before_wait', 'e1.after_automatic', 'before_save', 'e1.reopened', 'reopened']
    reads = [p + '.' + k for p in snapshot_phases for k in ('state_before', 'values', 'mode', 'state_after')]
    reads += ['e1.' + p + '.' + k for p in PHASES for k in ('force_full', 'workbooks')]
    reads += ['cleanup.observe.' + p for p in ('Ready', 'Interactive', 'DisplayAlerts')]
    need(all(calls.count(label) == 1 for label in reads), 'READ_CALLS_MISSING_OR_DUPLICATED', blocked=True)
    need(set(calls) <= set(required + reads + ['calculation_state', 'com.pump_messages']),
         'CALL_SCHEMA_UNSUPPORTED', blocked=True)
    need([c for c in calls if c.endswith('.full_rebuild')] == ['final.full_rebuild']
         and not any(c in ('calculate', 'range.calculate') for c in calls), 'EXTRA_CALCULATION')
    contexts = [r['evidence'] for r in records if r['event'] == 'e1.context']
    need(contexts == worker['e1']['evidence'], 'CONTEXT_PROGRESS_CHANGED')
    boundaries = [('open.after_call', 'manual.before_call'), ('edit.after_call', 'final.force_full.before_call'),
                  ('final.force_full_read.after_call', 'final.full_rebuild.before_call'),
                  ('final.full_rebuild.after_call', 'final.automatic_restore.before_call'),
                  ('final.automatic_read.after_call', 'save_as.before_call'),
                  ('reopen.after_call', 'reopened.close.before_call')]
    event_names = [r['event'] for r in records]
    context_events = [r for r in records if r['event'] == 'e1.context']
    for event, (begin, end) in zip(context_events, boundaries, strict=True):
        need(event_names.index(begin) < event['sequence'] - 1 < event_names.index(end), 'CONTEXT_CALL_ORDER')
        need(timestamp(event['evidence']['utc']) <= timestamp(event['utc']), 'CONTEXT_TIME_ORDER')
    snapshots = [r['sample'] for r in records if r['event'] == 'snapshot']
    need([s['phase'] for s in snapshots] == snapshot_phases, 'SNAPSHOT_PHASE_ORDER')
    for sample in worker['samples'] + [e['sample'] for e in contexts]:
        need(sum(s == sample for s in snapshots) == 1, 'SNAPSHOT_PROGRESS_MISMATCH')
    protocol = [r for r in records if r['event'] == 'calculation.protocol']
    need(len(protocol) == 1 and protocol[0].get('force_full_observed') is False
         and protocol[0].get('protocol') == E1_PROTOCOL and protocol[0].get('label') == 'final', 'PROTOCOL_PROGRESS')
    automatic = [r for r in records if r['event'] == 'calculation.automatic_observed']
    need(len(automatic) == 1 and automatic[0].get('mode') == -4105, 'AUTO_PROGRESS')
    starts = [i for i, r in enumerate(records) if r['event'] == 'calculation.wait_started']
    ends = [i for i, r in enumerate(records) if r['event'] == 'calculation.wait_done']
    need(len(starts) == len(ends) == 2 and starts[0] < ends[0] < starts[1] < ends[1],
         'WAIT_EVIDENCE_MISSING_OR_ORDER', blocked=True)
    for begin, end in zip(starts, ends, strict=True):
        observed = [r for r in records[begin:end] if r['event'] == 'calculation.observed']
        need(records[begin].get('timeout_seconds') == 180 and observed
             and observed[-1].get('state') == 0 and records[end].get('observations') == len(observed), 'WAIT_NOT_DONE')
        counts = {}
        for number, item in enumerate(observed, 1):
            need(type(item.get('state')) is int and item['state'] in (0, 1, 2) and item.get('observation') == number
                 and 0 <= item.get('elapsed_seconds', -1) < 180, 'WAIT_STATE_OR_DEADLINE')
            counts[str(item['state'])] = counts.get(str(item['state']), 0) + 1
        need(records[end].get('state_counts') == counts and 0 <= records[end].get('elapsed_seconds', -1) < 180,
             'WAIT_COUNTS_OR_DEADLINE')
    need(not any(r['event'] == 'calculation.wait_timeout' for r in records), 'RECORDED_WAIT_TIMEOUT')
    # Source-pinned successful ownership verification before Open, not PID string equality alone.
    owned = [r for r in records if r['event'] == 'ownership.verified']
    identity = [r for r in records if r['event'] == 'identity.verified']
    need(len(owned) == len(identity) == 1, 'OWNERSHIP_RECORD_MISSING', blocked=True)
    facts = worker['identity']; owner = owned[0]['identity']
    need(type(facts['pid']) is int and facts['pid'] > 0 and worker['owned_pid'] == facts['pid']
         and owner == {k: facts[k] for k in ('pid', 'created', 'path')}
         and identity[0]['identity'] == facts and worker.get('observed_process') == owner,
         'EXCEL_OWNER_FINGERPRINT')
    names = [r['event'] for r in records]
    need(names.index('ownership.verified') < names.index('identity.verified') < names.index('open.before_call'), 'OWNERSHIP_ORDER')
    started = timestamp(next(r['utc'] for r in records if r['event'] == 'connect.before_call'))
    need(started - 1 <= facts['created'] <= timestamp(identity[0]['utc']), 'EXCEL_CREATION_TIME')
    checks['excel_ownership'] = 'VERIFIED'
    residual = runtime.get('residual_observation')
    need(residual is not None, 'EXCEL_EXIT_RECORD_MISSING', blocked=True)
    need(residual.get('outcome') == 'PROCESS_EXITED'
         and parent['stages'][0].get('residual', {}).get('same_owned_process_remains') is False, 'EXCEL_EXIT_NOT_PROVEN', blocked=True)
    for observed in residual.get('samples', []):
        need({k: observed[k] for k in ('pid', 'created', 'path')} == owner, 'RESIDUAL_FINGERPRINT_CHANGED')
    checks['excel_exit'] = 'VERIFIED'
    pids = {r['worker_pid'] for r in records}
    need(all(type(p) is int and p > 0 for p in pids) and len(pids) == 1, 'PROGRESS_WORKER_PID')
    if runtime.get('python_processes') is not None:
        checks['python_processes'] = 'PRESENT'
        from .functions24 import validate_processes
        validate_processes(runtime)
        need(runtime['python_process']['pid'] in pids, 'PYTHON_WORKER_PROGRESS_MISMATCH')
        checks['python_relation'] = 'VERIFIED'
        basis = 'RECORDED_LAUNCHER_RELATION'
    elif runtime.get('worker_pid') in pids:
        # Pinned legacy supervisor polls this SAME worker, checks its returncode,
        # reads its request-bound final JSON and only then returns unfinished=False.
        checks['python_relation'] = 'VERIFIED'
        basis = 'LEGACY_DIRECT_SOURCE_AND_PROGRESS'
    else:
        checks['python_relation'] = 'UNSUPPORTED'
        raise EvidenceError('LEGACY_LAUNCHER_RELATION_UNPROVEN', blocked=True)
    need(runtime.get('worker_still_running') is False and runtime.get('force_termination') is False
         and not runtime.get('cleanup_pending'), 'WORKER_NOT_FINISHED', blocked=True)
    need(timestamp(worker['completed_utc']) >= previous_time, 'WORKER_COMPLETION_TIME')
    checks['progress'] = 'VERIFIED'
    return basis


def inspect(root, expected_run_sha, checks, started):
    from .run import validate_runtime  # Pure checker, not execute/supervise/source_identity.
    directory = root / 'tiny3'
    paths = [root / 'run.json', directory / 'case.json', directory / 'request.json', directory / 'worker.json']
    need(digest(read(paths[0])) == expected_run_sha, 'RUN_PIN_MISMATCH')
    parent, manifest, request, worker = [parse(read(p)) for p in paths]
    need(parent.get('schema') == SCHEMA and parent.get('mode') == 'e1'
         and parent.get('diagnostic') == E1_DIAGNOSTIC and parent.get('protocol') == E1_PROTOCOL
         and parent.get('status') == 'DIAGNOSTIC_PASS' and type(parent.get('exit_code')) is int and parent.get('exit_code') == 0
         and parent.get('diagnostic_only') is True and parent.get('diagnostic_recalculation_verified') is True
         and parent.get('gate2_acceptance') is False and parent.get('product_acceptance') is False
         and parent.get('excel_recalculation_verified') is False and parent.get('force_termination') is False
         and not parent.get('first_error'), 'NOT_APPROVED_SINGLE_E1')
    need([s['case_id'] for s in parent['stages']] == ['tiny3', 'functions24', 'annual1']
         and [s['status'] for s in parent['stages']] == ['PASS', 'NOT_RUN', 'NOT_RUN'], 'STAGES_CHANGED')
    checks['source'] = 'PRESENT'; historical_source(parent['source']); checks['source'] = 'VERIFIED'
    need(parent.get('python_bits') == 64 and '3.12.10' in parent.get('python', '')
         and 'Windows' in parent.get('os', ''), 'LEGACY_ENVIRONMENT_UNSUPPORTED', blocked=True)
    need(manifest == dict(run_id=parent['run_id'], case_id='tiny3', file=FILE, sha256=ORIGINAL_SHA,
                         formula_count=3, probes=[['診断', c, v] for c, v in cases.SPEC['expected_after'].items()],
                         edits=[['診断', 'A1', 7], ['診断', 'A2', 3]], challenges=[]), 'MANIFEST_CHANGED')
    need(request.get('schema', SCHEMA) == SCHEMA and request['run_id'] == parent['run_id'] == worker['run_id']
         and request.get('case_id') == 'tiny3' and request.get('diagnostic') == E1_DIAGNOSTIC
         and request['source_sha256'] == parent['source']['source_sha256']
         and request['manifest_sha256'] == digest(read(paths[1])), 'REQUEST_BINDING')
    for part in ('inputs', 'work'):
        need(digest(read(directory / part / FILE)) == ORIGINAL_SHA, 'ORIGINAL_INPUT_OR_WORK_CHANGED')
    need(digest((pre.REPO / 'output/s1_tiny_diagnostic' / FILE).read_bytes()) == ORIGINAL_SHA, 'LOCAL_ORIGINAL_CHANGED')
    checks['input'] = 'VERIFIED'
    runtime = copy.deepcopy(parent['stages'][0]['runtime'])
    need(not parent['stages'][0].get('first_error') and not parent['stages'][0].get('residual_error'), 'PARENT_RECORDED_ERROR')
    need(all(runtime.get(k) == v for k, v in worker.items() if k != 'observed_identity'), 'PARENT_WORKER_MISMATCH')
    validate_runtime(runtime, directory, request, request['profile'])
    need(type(worker.get('full_rebuild_calls')) is int and worker['full_rebuild_calls'] == 1, 'REBUILD_COUNT_TYPE')
    e1.validate_evidence(runtime, manifest)
    need([s['phase'] for s in worker['samples']] == ['opened', 'before_save', 'reopened'], 'REQUIRED_SAMPLES_MISSING')
    for item in worker['e1']['evidence']:
        need(item['phase'] in PHASES and item['force_full'] is False, 'FLAG_NOT_STRICT_FALSE')
        sample = item['sample']
        need(all(type(sample[k]) is int for k in ('state_before', 'state', 'mode'))
             and sample['state_before'] in (0, 1, 2) and sample['state'] in (0, 1, 2), 'INVALID_STATE_TYPE')
        book = item['workbooks'][0]
        destination = directory / ('results' if item['phase'] == 'reopened' else 'work') / FILE
        need(ntpath.normcase(book['path']) == ntpath.normcase(str(destination))
             and book['name'] == FILE and book['read_only'] is (item['phase'] == 'reopened'), 'BOOK_PATH_OR_READONLY')
        timestamp(item['utc'])
        if item['phase'] == 'opened':
            need(sample['mode'] == -4105 and sample['values'] == {
                '診断!B1': 5, '診断!B2': 6, '診断!B3': 11}, 'OPENED_CONTEXT_BASELINE')
    initial = worker['samples'][0]
    need(initial['mode'] == -4105 and initial['values'] == {
        '診断!B1': 5, '診断!B2': 6, '診断!B3': 11}, 'INITIAL_VALUES_OR_MODE')
    for sample in worker['samples']:
        need(all(type(sample[k]) is int for k in ('state_before', 'state', 'mode'))
             and sample['state_before'] in (0, 1, 2) and sample['state'] in (0, 1, 2), 'INVALID_SAMPLE_TYPE')
    checks['runtime'] = checks['observations'] = 'VERIFIED'
    records = [parse(line) for line in read(directory / 'progress.jsonl').splitlines() if line.strip()]
    python_basis = evidence_progress(records, parent, worker, runtime, manifest, directory, checks)
    need(not read(directory / 'worker-console.txt').strip(), 'WORKER_CONSOLE_REQUIRES_REVIEW', blocked=True)
    # Never follow paths from request/profile/unit command. Only fixed local basenames.
    units = parent.get('unit_tests', {})
    checks['units'] = 'PRESENT' if units else 'MISSING'
    need(units.get('returncode') == 0 and units.get('junit') == 'unit-tests.xml'
         and units.get('log') == 'unit-tests.txt', 'UNIT_RECORD_UNSUPPORTED', blocked=True)
    need(all((root / f).is_file() for f in ('unit-tests.xml', 'unit-tests.txt')), 'UNIT_ARTIFACTS_MISSING', blocked=True)
    suite = xml(read(root / 'unit-tests.xml'))
    sums = {k: sum(int(s.get(k, 0)) for s in suite.iter('testsuite')) for k in ('tests', 'failures', 'errors', 'skipped')}
    need(sums['tests'] > 0 and sums['failures'] == sums['errors'] == 0
         and units.get('failed') == units.get('errors') == 0
         and all(type(units.get(k)) is int and units[k] >= 0 for k in ('passed', 'failed', 'errors', 'skipped'))
         and units['passed'] > 0
         and units.get('passed') == sums['tests'] - sums['skipped'] and units.get('skipped') == sums['skipped'], 'UNIT_COUNTS_CHANGED')
    read(root / 'unit-tests.txt'); checks['units'] = checks['unit_artifacts'] = 'VERIFIED'
    original, saved = directory / 'inputs' / FILE, directory / 'results' / FILE
    for path in (original, saved): bounded_book(path, started)
    need(digest(read(saved)) == worker['saved_sha256'] == parent['stages'][0]['verification']['saved_sha256'], 'SAVED_SHA_CHANGED')
    verified = cases.verify_saved(directory, manifest)
    nonformula_union(original, saved, started)
    need(verified['saved_sha256'] == worker['saved_sha256'], 'SAVED_SHA_CHANGED')
    checks['saved'] = 'VERIFIED'
    deadline(started)
    return dict(original_status='DIAGNOSTIC_PASS', source_candidate=CANDIDATE_COMMIT,
                input_label='SYNTHETIC_TEST_ONLY' if parent.get('fixture_kind') == 'SYNTHETIC_TEST_ONLY'
                else 'UNLABELLED_PAST_EVIDENCE',
                legacy_source_resolution='FIXED_MAP_MATCH_NOT_SIGNED_HISTORY', saved_sha256=verified['saved_sha256'],
                formula_count=3, numeric_comparison='matched', nonformula_union='matched',
                recorded_excel_exit='PROCESS_EXITED', current_process_state='NOT_CHECKED',
                python_evidence_basis=python_basis)


class OfflineGuard:
    def __init__(self):
        from .no_excel_tests import Guard
        self.base = Guard(); self.violations = []

    def __call__(self, event, args):
        try:
            if event == 'subprocess.Popen': raise RuntimeError('PROCESS_LAUNCH')
            self.base(event, args)
        except RuntimeError:
            if len(self.violations) < 20: self.violations.append('OFFLINE_FORBIDDEN_OPERATION')
            raise


def execute(root, output, expected_run_sha, expected_source_sha, *, owner_confirmed=False, acl_confirmed=False):
    checks = {k: 'MISSING' for k in ('source', 'input', 'runtime', 'observations', 'progress',
                                    'excel_ownership', 'excel_exit', 'python_relation', 'python_processes',
                                    'units', 'unit_artifacts', 'saved', 'evidence_unchanged')}
    result = dict(status='BLOCKED', scope='E1_OFFLINE_EVIDENCE_ONLY', checks=checks,
                  excel_started=False, com_connected=False, approval_id_used=False, gate2_acceptance=False)
    started = time.monotonic(); before = None; safe_output = None
    try:
        need(owner_confirmed and acl_confirmed and __debug__, 'OWNER_AND_ACL_CONFIRMATION_REQUIRED', blocked=True)
        need(re.fullmatch('[0-9a-f]{64}', expected_run_sha or ''), 'EXPECTED_RUN_PIN_REQUIRED', blocked=True)
        source = pre.check_source(expected_source_sha)
        root, safe_output = safe_path(root), safe_path(output, new=True)
        need(root.is_dir() and not os.path.lexists(safe_output) and not safe_output.is_relative_to(root)
             and not root.is_relative_to(safe_output), 'NEW_SEPARATE_OUTPUT_REQUIRED', blocked=True)
        for ancestor in (safe_output.parent, *safe_output.parent.parents):
            need(not any((ancestor / p).exists() for p in ('tools/excel2024_validation', 'run.json',
                     'case.json', 'offline-report.json', 'offline-e1-report.json')), 'OUTPUT_IN_PROTECTED_TREE', blocked=True)
        need(os.access(root, os.R_OK | os.X_OK) and os.access(safe_output.parent, os.W_OK | os.X_OK),
             'ACCESS_NOT_CONFIRMED', blocked=True)
        before = inventory(root, started)
        try:
            result.update(inspect(root, expected_run_sha, checks, started), status='OFFLINE_E1_REVERIFICATION_PASS')
        except EvidenceError as error:
            result.update(status='BLOCKED' if error.blocked else 'OFFLINE_E1_REVERIFICATION_FAIL', reason=str(error))
        except (FileNotFoundError, KeyError):
            result.update(status='BLOCKED', reason='EVIDENCE_OR_SCHEMA_MISSING')
        except (Exception,):
            result.update(status='OFFLINE_E1_REVERIFICATION_FAIL', reason='EVIDENCE_CHECK_FAILED')
        need(inventory(root, started) == before, 'EVIDENCE_CHANGED_DURING_AUDIT', blocked=True)
        checks['evidence_unchanged'] = 'VERIFIED'
        result.update(run_json_sha256=before.get('run.json'), evidence_files=len(before),
                      evidence_inventory_sha256=digest(json.dumps(before, sort_keys=True).encode()),
                      new_portable_source_sha256=source['release_source_sha256'])
        safe_output.mkdir(exist_ok=False)
        with (safe_output / 'offline-e1-report.json').open('x', encoding='utf-8') as stream:
            json.dump(result, stream, ensure_ascii=False, indent=2); stream.flush(); os.fsync(stream.fileno())
        with (safe_output / 'offline-e1-report.sha256').open('x', encoding='ascii') as stream:
            stream.write(digest(read(safe_output / 'offline-e1-report.json')) + '\n')
    except EvidenceError as error:
        result.update(status='BLOCKED', reason=str(error))
    except BaseException:
        result.update(status='BLOCKED', reason='PRECONDITION_PRESERVATION_OR_PUBLICATION_FAILED')
    return result


def main():
    parser = pre.SafeParser(description=__doc__)
    parser.add_argument('--source-run', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--expected-run-sha', required=True)
    parser.add_argument('--expected-source-sha', required=True)
    parser.add_argument('--owner-confirmed', action='store_true')
    parser.add_argument('--acl-confirmed', action='store_true')
    args = parser.parse_args()
    guard = OfflineGuard(); sys.addaudithook(guard)
    value = execute(args.source_run, args.output, args.expected_run_sha, args.expected_source_sha,
                    owner_confirmed=args.owner_confirmed, acl_confirmed=args.acl_confirmed)
    if guard.violations: value.update(status='BLOCKED', reason='OFFLINE_GUARD_VIOLATION')
    public = {k: value[k] for k in ('status', 'reason', 'checks', 'scope', 'input_label', 'excel_started', 'com_connected',
                                    'approval_id_used', 'gate2_acceptance') if k in value}
    print(json.dumps(public, ensure_ascii=False, indent=2))
    return 0 if value['status'] == 'OFFLINE_E1_REVERIFICATION_PASS' else 2


if __name__ == '__main__':
    raise SystemExit(main())
