"""Fresh tiny3-only one-shot diagnostic; never a Gate 2 acceptance result.

The Windows path remains dormant until an owner supplies a separate,
single-use approval ID.  Linux tests inject the monitor and never import COM.
"""
import argparse
from datetime import datetime
import hashlib
import json
import ntpath
import os
from pathlib import Path
import re
import stat
import subprocess
import sys
import time
import uuid
import zipfile

from openpyxl import load_workbook

from . import E1_DIAGNOSTIC, E1_PROTOCOL, SCHEMA
from . import cases, e1, windows_preflight as pre
from .common import GateError, error_info, read_json, sha, utc


FRESH_SCHEMA = 'G2_FRESH_TINY3_ONE_SHOT_V1'
FRESH_LEDGER = 'lifeplan-excel2024-fresh-tiny3-approvals'
ORIGINAL_SHA = '38c28cbceebc6bf9645b6791819cf00dce790f6f3ed3bbeb1c404bfcfa6e3682'
APPROVAL = re.compile(r'G2-20261017-017-WIN-[0-9]{2}')
MAX_ARCHIVE, MAX_EXPANDED, MAX_ENTRIES = 1024 * 1024, 8 * 1024 * 1024, 64
MAX_ROWS, MAX_COLUMNS, MAX_CELLS = 64, 32, 2048


class Blocked(Exception):
    """Fixed fail-closed reason code; CLI never echoes private arguments."""


def require(value, code):
    if not value:
        raise Blocked(code)


class Parser(argparse.ArgumentParser):
    def error(self, message):
        raise Blocked('ARGUMENTS_INVALID')


def _write_new_json(path, value):
    """Exclusive durable write; an old run or approval is never overwritten."""
    path = Path(path)
    with path.open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, sort_keys=True, indent=2, default=str)
        stream.flush()
        os.fsync(stream.fileno())


def _fixed_ntfs(path, filesystem_check):
    """Require explicit fixed-local NTFS evidence before any Excel launch."""
    try:
        evidence = filesystem_check(Path(path))
    except pre.Blocked as error:
        raise Blocked('LOCAL_FIXED_NTFS_UNVERIFIED') from error
    require(type(evidence) is dict
            and evidence.get('drive_type') == 'FIXED_LOCAL'
            and evidence.get('filesystem') == 'NTFS', 'LOCAL_FIXED_NTFS_UNVERIFIED')
    return evidence


def _new_output(path, *, filesystem_check=pre.local_ntfs):
    raw = Path(path)
    require('..' not in raw.parts, 'OUTPUT_PARENT_TRAVERSAL')
    path = raw.absolute()
    require(not str(path).startswith(('\\\\', '//')), 'OUTPUT_UNC_NOT_ALLOWED')
    require(not path.is_relative_to(pre.REPO), 'OUTPUT_IN_REPOSITORY')
    require(not os.path.lexists(path), 'OUTPUT_MUST_BE_NEW')
    parent = pre.path_check(path.parent, outside=True)
    require(parent.is_dir(), 'OUTPUT_PARENT_REQUIRED')
    return path, _fixed_ntfs(parent, filesystem_check)


def validate_approval_id(value):
    require(type(value) is str and APPROVAL.fullmatch(value) is not None,
            'SEPARATE_WINDOWS_APPROVAL_ID_REQUIRED')


def _regular_json(path, limit=131072, *, filesystem_check=pre.local_ntfs):
    path = pre.path_check(path, outside=True)
    filesystem = _fixed_ntfs(path, filesystem_check)
    info = path.stat()
    require(stat.S_ISREG(info.st_mode) and info.st_nlink == 1 and info.st_size <= limit,
            'PROFILE_REGULAR_FILE_REQUIRED')
    raw = path.read_bytes()
    require(len(raw) <= limit, 'PROFILE_SIZE_LIMIT')
    return path, json.loads(raw.decode('utf-8-sig')), filesystem


def _process_record(record):
    return (type(record) is dict and type(record.get('pid')) is int and record['pid'] > 0
            and type(record.get('ppid')) is int and record['ppid'] >= 0
            and isinstance(record.get('created'), (int, float)) and record.get('path'))


def _same_process(left, right):
    return (_process_record(left) and _process_record(right)
            and left['pid'] == right['pid'] and left['ppid'] == right['ppid']
            and left['created'] == right['created']
            and ntpath.normcase(left['path']) == ntpath.normcase(right['path']))


def validate_manifest(manifest):
    spec = cases.SPEC
    expected = dict(case_id='tiny3', file=spec['file'], sha256=ORIGINAL_SHA,
                    formula_count=3,
                    probes=[[spec['sheet'], cell, value]
                            for cell, value in spec['expected_after'].items()],
                    edits=[[spec['sheet'], cell, value]
                           for cell, value in spec['inputs_after'].items()],
                    challenges=[])
    for key, value in expected.items():
        require(manifest.get(key) == value, 'TINY3_MANIFEST_MISMATCH')
    require(type(manifest.get('run_id')) is str and manifest['run_id'], 'RUN_ID_MISSING')
    return manifest


def _tiny_package(path, *, input_only=False):
    """Bound the tiny workbook before openpyxl iterates any dimensions."""
    path = Path(path)
    require(path.stat().st_size <= MAX_ARCHIVE, 'TINY3_ARCHIVE_LIMIT')
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        names = [entry.filename for entry in entries]
        require(archive.testzip() is None and len(entries) <= MAX_ENTRIES
                and len(set(names)) == len(names)
                and sum(entry.file_size for entry in entries) <= MAX_EXPANDED,
                'TINY3_ZIP_LIMIT_OR_CORRUPTION')
        require(not any(entry.flag_bits & 1 or entry.filename.startswith('/')
                        or '..' in entry.filename.split('/') for entry in entries),
                'TINY3_ZIP_UNSAFE_PATH')
        require(not any('externalLink' in name or 'vbaProject' in name
                        or name == 'xl/connections.xml' or name.startswith('xl/embeddings/')
                        for name in names), 'TINY3_EXTERNAL_CONTENT')
    if input_only:
        # Input fixture must have no synthetic formula caches.  Saved Excel
        # evidence must have caches, so it is checked by openpyxl below.
        require(cases.static_input(path) == 3, 'TINY3_FORMULA_COUNT')


def validate_saved_values(directory, manifest):
    """AND-check all tiny3 formulas and every non-formula value in bounded sheets."""
    directory = Path(directory)
    original = directory / 'inputs' / manifest['file']
    saved = directory / 'results' / manifest['file']
    _tiny_package(original, input_only=True)
    _tiny_package(saved)
    expected_values = dict(cases.SPEC['inputs_after'])
    expected_formulas = dict(cases.SPEC['formulas'])
    book = load_workbook(saved, read_only=True, data_only=False)
    try:
        require(book.sheetnames == [cases.SPEC['sheet']], 'TINY3_UNEXPECTED_SHEET')
        sheet = book[cases.SPEC['sheet']]
        rows, columns = sheet.max_row or 1, sheet.max_column or 1
        require(rows <= MAX_ROWS and columns <= MAX_COLUMNS and rows * columns <= MAX_CELLS,
                'TINY3_CELL_LIMIT')
        values, formulas = {}, {}
        for row in sheet.iter_rows(min_row=1, min_col=1, max_row=rows, max_col=columns):
            for cell in row:
                if cell.data_type == 'f':
                    formulas[cell.coordinate] = cell.value
                elif cell.value is not None:
                    values[cell.coordinate] = cell.value
        require(formulas == expected_formulas, 'TINY3_FORMULA_IDENTITY')
        require(values == expected_values, 'TINY3_NONFORMULA_VALUE_CHANGE')
    finally:
        book.close()
    return dict(all_nonformula_values='matched', formulas='matched', sheets='matched')


def validate_evidence_paths(result, manifest, directory):
    evidence = result.get('e1', {}).get('evidence', [])
    work = ntpath.normcase(str((Path(directory) / 'work' / manifest['file']).resolve()))
    saved = ntpath.normcase(str((Path(directory) / 'results' / manifest['file']).resolve()))
    for item in evidence:
        books = item.get('workbooks', [])
        require(len(books) == 1, 'TINY3_WORKBOOK_INVENTORY')
        book = books[0]
        reopened = item.get('phase') == 'reopened'
        wanted = saved if reopened else work
        require(book.get('name') == manifest['file']
                and ntpath.normcase(book.get('path', '')) == wanted
                and book.get('read_only') is reopened,
                'TINY3_WORKBOOK_PATH_OR_MODE')


def _events(path):
    result = []
    for line in Path(path).read_text(encoding='utf-8').splitlines():
        if line.strip():
            result.append(json.loads(line))
    return result


def validate_progress(directory, result):
    records = _events(Path(directory) / 'progress.jsonl')
    require(records and all(record.get('run_id') == result['run_id']
                            and record.get('case_id') == 'tiny3'
                            and record.get('worker_pid') == result['python_processes']['worker']['pid']
                            for record in records), 'PROGRESS_RUN_OR_PID_MISMATCH')
    names = [record.get('event') for record in records]
    required = ('ownership.verified', 'identity.verified', 'settings.after_call',
                'final.full_rebuild.before_call', 'final.full_rebuild.after_call',
                'final.automatic_read.before_call', 'final.automatic_read.after_call',
                'calculation.wait_done', 'save_as.after_call', 'reopen.after_call',
                'cleanup.quit.after_call', 'cleanup.release.after_call')
    require(all(name in names for name in required), 'PROGRESS_REQUIRED_EVENT_MISSING')
    require(not any(name and (name.endswith('.failed') or name == 'calculation.wait_timeout')
                    for name in names), 'PROGRESS_FAILURE_EVENT')
    observed = {record.get('property'): record.get('value') for record in records
                if record.get('event') == 'cleanup.observed'}
    require(observed.get('Ready') is True and observed.get('Interactive') is True
            and observed.get('DisplayAlerts') is True, 'MODAL_OR_INTERACTIVE_STATE_UNSAFE')
    return dict(events=len(records), calculation_wait_done=names.count('calculation.wait_done'),
                cleanup_observed=observed)


def validate_processes(result):
    proof = result.get('python_processes', {})
    controller, launcher, worker = (proof.get(key, {}) for key in ('controller', 'launcher', 'worker'))
    require(all(_process_record(record) for record in (controller, launcher, worker)),
            'PYTHON_PROCESS_PROOF_MISSING')
    require(launcher['ppid'] == controller['pid'] and _same_process(launcher, worker),
            'PYTHON_PARENT_OR_WORKER_MISMATCH')
    require(launcher['pid'] == result.get('worker_pid')
            and worker.get('sys_executable')
            and ntpath.normcase(worker['sys_executable']) == ntpath.normcase(worker['path']),
            'PYTHON_EXECUTABLE_MISMATCH')
    require(proof.get('worker_exit', {}).get('outcome') in ('PROCESS_EXITED', 'PID_REUSED'),
            'PYTHON_WORKER_EXIT_UNCONFIRMED')
    return 'DIRECT_FRESH_WORKER'


def _approval_marker(profile_path, approval_id, run_id, source_sha256):
    ledger = profile_path.parent / FRESH_LEDGER
    ledger.mkdir(exist_ok=True)
    ledger = pre.path_check(ledger, outside=True)
    require(ledger.is_dir(), 'APPROVAL_LEDGER_INVALID')
    marker = ledger / (approval_id + '.json')
    value = dict(schema=FRESH_SCHEMA, approval_id=approval_id, run_id=run_id,
                 source_sha256=source_sha256, original_sha256=ORIGINAL_SHA,
                 profile_sha256=sha(profile_path), created_utc=utc())
    _write_new_json(marker, value)
    require(marker.stat().st_nlink == 1, 'APPROVAL_MARKER_LINKED')
    return marker, value


def validate_control(request, directory, *, process_identity):
    """Worker-wrapper pre-COM gate, including the actual parent relationship."""
    control = request.get('fresh_tiny3_control', {})
    approval_keys = {'schema', 'approval_id', 'run_id', 'source_sha256', 'original_sha256',
                     'profile_sha256', 'created_utc'}
    control_keys = approval_keys | {'controller', 'profile_path', 'approval_marker',
                                    'approval_marker_sha256'}
    require(set(control) == control_keys and control.get('schema') == FRESH_SCHEMA
            and control.get('run_id') == request.get('run_id')
            and control.get('original_sha256') == ORIGINAL_SHA,
            'FRESH_CONTROL_MISMATCH')
    validate_approval_id(control.get('approval_id'))
    manifest = validate_manifest(read_json(Path(directory) / 'case.json'))
    require(set(request) == {'run_id', 'case_id', 'diagnostic', 'protocol', 'profile',
                             'profile_sha256', 'manifest_sha256', 'source_sha256',
                             'fresh_tiny3_control'}
            and request.get('case_id') == 'tiny3' and request.get('diagnostic') == E1_DIAGNOSTIC
            and request.get('protocol') == E1_PROTOCOL
            and request.get('manifest_sha256') == sha(Path(directory) / 'case.json')
            and request.get('source_sha256') == control.get('source_sha256')
            and request.get('profile_sha256') == control.get('profile_sha256')
            and manifest['run_id'] == request['run_id'], 'FRESH_REQUEST_BINDING')
    marker = pre.path_check(control.get('approval_marker'), outside=True)
    require(marker.parent.name == FRESH_LEDGER and marker.name == control['approval_id'] + '.json'
            and marker.stat().st_nlink == 1 and sha(marker) == control.get('approval_marker_sha256'),
            'APPROVAL_MARKER_MISMATCH')
    approval = read_json(marker)
    require(set(approval) == approval_keys
            and approval == {key: control[key] for key in approval_keys}, 'APPROVAL_MARKER_CONTENT')
    profile_path = pre.path_check(control.get('profile_path'), outside=True)
    require(profile_path.is_file() and profile_path.stat().st_nlink == 1
            and sha(profile_path) == control['profile_sha256']
            and read_json(profile_path) == request.get('profile'), 'PROFILE_FILE_BINDING')
    controller = control.get('controller', {})
    actual_parent = process_identity(os.getppid())
    require(_same_process(controller, actual_parent), 'ACTUAL_PARENT_MISMATCH')
    return process_identity(os.getpid())


def supervise(directory, request, *, popen=subprocess.Popen, process_identity=None,
              clock=time.monotonic, pause=time.sleep, overall_timeout=600):
    """Launch one wrapper, authorize its exact PID, and never terminate it."""
    from . import functions24, run
    if process_identity is None:
        process_identity = functions24.process_identity
    directory = Path(directory)
    _write_new_json(directory / 'request.json', request)
    controller = request['fresh_tiny3_control']['controller']
    require(_same_process(controller, process_identity(os.getpid())), 'CONTROLLER_PROCESS_CHANGED')
    started = clock()
    records, seen, in_flight = [], 0, {}
    outcome = None
    proc = None
    launcher, worker = {}, {}
    try:
        with (directory / 'worker-console.txt').open('x', encoding='utf-8') as console:
            proc = popen([sys.executable, '-m', 'tools.excel2024_validation.fresh_tiny3_worker',
                          '--request', str(directory / 'request.json')], cwd=cases.REPO,
                         stdout=console, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL)
        launcher = process_identity(proc.pid)
        require(launcher['ppid'] == controller['pid'], 'LAUNCHER_PARENT_MISMATCH')
        proof_path = directory / 'python-worker.json'
        while not proof_path.is_file() and proc.poll() is None and clock() - started < 15:
            pause(.05)
        if proof_path.is_file():
            proof = read_json(proof_path)
            worker = proof.get('process', {})
            require(proof.get('request_sha256') == sha(directory / 'request.json')
                    and _same_process(launcher, worker), 'WRAPPER_PROCESS_MISMATCH')
            require(type(worker.get('sys_executable')) is str and worker['sys_executable']
                    and ntpath.normcase(ntpath.normpath(worker['sys_executable']))
                    == ntpath.normcase(ntpath.normpath(worker['path'])),
                    'PYTHON_EXECUTABLE_MISMATCH')
            _write_new_json(directory / 'launch-authorization.json',
                            dict(schema=FRESH_SCHEMA, request_sha256=proof['request_sha256'],
                                 worker_sha256=sha(proof_path)))
        else:
            outcome = dict(status='BLOCKED', reason='WRAPPER_HANDSHAKE_MISSING')
        while outcome is None:
            records = run.events(directory / 'progress.jsonl')
            for event in records[seen:]:
                if (event.get('run_id') != request['run_id'] or event.get('case_id') != 'tiny3'
                        or event.get('worker_pid') != proc.pid):
                    outcome = dict(status='FAIL', reason='PROGRESS_IDENTITY_MISMATCH')
                    break
                name = event.get('event', '')
                if name.endswith('.before_call'):
                    elapsed = max(0, time.time() - datetime.fromisoformat(event['utc']).timestamp())
                    in_flight[name.removesuffix('.before_call')] = clock() + event['timeout_seconds'] - elapsed
                elif name.endswith(('.after_call', '.failed')):
                    in_flight.pop(name.rsplit('.', 1)[0], None)
            seen = len(records)
            if outcome is not None:
                break
            expired = [name for name, deadline in in_flight.items() if clock() >= deadline]
            if expired or clock() - started > overall_timeout:
                outcome = dict(status='TIMEOUT', reason='SYNC_CALL_OR_OVERALL_TIMEOUT', pending_calls=expired)
                break
            if proc.poll() is not None:
                result_path = directory / 'worker.json'
                outcome = read_json(result_path) if result_path.is_file() else dict(
                    status='FAIL', reason='WORKER_RESULT_MISSING', worker_exit=proc.returncode)
                if proc.returncode != 0 and outcome.get('status') == 'SAVED_REOPENED':
                    outcome.update(status='FAIL', reason='WORKER_EXIT_CODE_MISMATCH')
                break
            pause(.1)
    except BaseException as error:
        (directory / 'STOP').touch(exist_ok=True)
        outcome = dict(status=error.status if isinstance(error, GateError) else 'BLOCKED',
                       reason='SUPERVISOR_HANDSHAKE_OR_MONITOR_FAILURE', first_error=error_info(error))
    if outcome.get('status') in ('TIMEOUT', 'INTERRUPTED', 'FAIL', 'BLOCKED'):
        (directory / 'STOP').touch(exist_ok=True)
    still_running = proc is not None and proc.poll() is None
    proof = dict(controller=controller, launcher=launcher, worker=worker)
    if not still_running and worker:
        proof['worker_exit'] = functions24.worker_exit_observation(worker)
    outcome.update(worker_pid=proc.pid if proc else None, worker_still_running=still_running,
                   cleanup_pending=still_running, force_termination=False,
                   python_processes=proof, progress_event_count=len(records))
    if still_running:
        outcome['action'] = '所有者が専用Excelを通常操作で閉じる。既存ExcelやPythonを強制終了しない'
    return outcome


def execute(config, output, expected_source_sha, approval_id, *, confirmed=False,
            environment=None, source_check=None, monitor=supervise, residual=None,
            process_identity=None, run_id_factory=None, filesystem_check=pre.local_ntfs):
    """Prepare and run exactly one fresh tiny3; injectable dependencies are for Linux mocks."""
    from . import functions24, run
    environment = environment or (lambda: pre.windows_environment())
    source_check = source_check or pre.check_source
    residual = residual or run.residual_check
    process_identity = process_identity or functions24.process_identity
    run_id_factory = run_id_factory or (lambda: str(uuid.uuid4()))
    result = dict(schema=FRESH_SCHEMA, status='BLOCKED', reason='NOT_COMPLETED',
                  success=False, exit_code=2, case_id='tiny3', diagnostic_only=True,
                  excel_started=False, com_connected=False, approval_id_used=False,
                  old_e1_reused=False, p1_status='NOT_RUN', annual1_status='NOT_RUN',
                  gate2_acceptance=False, product_acceptance=False,
                  force_termination=False, next_action='STOP_PM_REVIEW')
    root = None
    try:
        require(__debug__, 'ASSERTS_REQUIRED')
        require(confirmed is True, 'OWNER_CONFIRMATION_REQUIRED')
        validate_approval_id(approval_id)
        require(type(expected_source_sha) is str and re.fullmatch('[0-9a-f]{64}', expected_source_sha),
                'SOURCE_PIN_REQUIRED')
        environment()
        source_filesystem = _fixed_ntfs(pre.path_check(cases.REPO), filesystem_check)
        source = source_check(expected_source_sha)
        require(source.get('release_source_sha256') == expected_source_sha, 'SOURCE_SHA_MISMATCH')
        config, profile, profile_filesystem = _regular_json(
            config, filesystem_check=filesystem_check)
        require(profile.get('schema') == SCHEMA, 'PROFILE_SCHEMA_MISMATCH')
        output, output_filesystem = _new_output(output, filesystem_check=filesystem_check)
        result['filesystem'] = dict(source_root=source_filesystem,
                                    profile=profile_filesystem,
                                    output_parent=output_filesystem)
        run_id = run_id_factory()
        require(type(run_id) is str and re.fullmatch('[0-9a-f-]{36}', run_id), 'RUN_ID_INVALID')
        root = output
        root.mkdir(exist_ok=False)
        case_dir = root / 'tiny3'
        case_dir.mkdir()
        marker, approval = _approval_marker(config, approval_id, run_id, source['source_sha256'])
        result['approval_id_used'] = True
        manifest = validate_manifest(cases.prepare('tiny3', case_dir, run_id))
        require(sha(case_dir / 'inputs' / manifest['file']) == ORIGINAL_SHA, 'TINY3_ORIGINAL_SHA')
        controller = process_identity(os.getpid())
        request = dict(run_id=run_id, case_id='tiny3', diagnostic=E1_DIAGNOSTIC,
                       protocol=E1_PROTOCOL, profile=profile,
                       profile_sha256=approval['profile_sha256'],
                       manifest_sha256=sha(case_dir / 'case.json'),
                       source_sha256=source['source_sha256'], fresh_tiny3_control=dict(
                           **approval, controller=controller, profile_path=str(config),
                           approval_marker=str(marker),
                           approval_marker_sha256=sha(marker)))
        runtime = monitor(case_dir, request)
        result.update(excel_started=runtime.get('identity') is not None,
                      com_connected=runtime.get('observed_identity') is not None,
                      runtime_status=runtime.get('status'), worker_still_running=runtime.get('worker_still_running'))
        if runtime.get('status') != 'SAVED_REOPENED':
            status = runtime.get('status')
            raise GateError(status if status in ('BLOCKED', 'TIMEOUT', 'INTERRUPTED', 'FAIL') else 'FAIL',
                            'WORKER_NOT_SAVED_REOPENED')
        require(runtime.get('force_termination') is False and not runtime.get('worker_still_running')
                and not runtime.get('cleanup_pending') and not runtime.get('cleanup_errors')
                and not runtime.get('progress_log_errors') and not runtime.get('first_error'),
                'WORKER_OR_CLEANUP_INCOMPLETE')
        result['python_process_model'] = validate_processes(runtime)
        run.validate_runtime(runtime, case_dir, request, profile)
        e1.validate_evidence(runtime, manifest)
        validate_evidence_paths(runtime, manifest, case_dir)
        result['progress'] = validate_progress(case_dir, runtime)
        verified = cases.verify_saved(case_dir, manifest)
        verified.update(validate_saved_values(case_dir, manifest))
        require(verified['saved_sha256'] == runtime.get('saved_sha256'), 'SAVED_SHA_MISMATCH')
        require(residual(runtime) is False, 'OWNED_EXCEL_PROCESS_REMAINS')
        observation = runtime.get('residual_observation', {})
        require(observation.get('outcome') in ('PROCESS_EXITED', 'PID_REUSED'),
                'OWNED_EXCEL_EXIT_UNCONFIRMED')
        result.update(status='FRESH_TINY3_DIAGNOSTIC_PASS', reason='DIAGNOSTIC_ONLY',
                      success=True, exit_code=0, verification=verified,
                      saved_sha256=verified['saved_sha256'], source=source,
                      diagnostic_recalculation_verified=True)
    except Blocked as error:
        result['reason'] = str(error)
    except GateError as error:
        result.update(status=error.status if error.status in ('BLOCKED', 'TIMEOUT') else 'FAIL',
                      reason='RUNTIME_GATE_FAILURE', first_error=error_info(error), exit_code=2 if error.status == 'BLOCKED' else 1)
    except BaseException:
        # Do not expose profile paths, approval values, exception text, or run UUIDs on the CLI.
        result.update(status='FAIL', reason='FRESH_TINY3_VALIDATION_FAILED', exit_code=1)
    if root is not None:
        try:
            _write_new_json(root / 'run.json', result)
        except BaseException:
            result.update(status='FAIL', reason='REPORT_WRITE_FAILED', success=False, exit_code=1)
    return result


def main(argv=None):
    try:
        parser = Parser(description=__doc__, add_help=False)
        parser.add_argument('--config', type=Path, required=True)
        parser.add_argument('--output', type=Path, required=True)
        parser.add_argument('--expected-source-sha', required=True)
        parser.add_argument('--approval-id', required=True)
        parser.add_argument('--owner-confirmed', action='store_true')
        args = parser.parse_args(argv)
        value = execute(args.config, args.output, args.expected_source_sha, args.approval_id,
                        confirmed=args.owner_confirmed)
    except BaseException:
        value = dict(schema=FRESH_SCHEMA, status='BLOCKED', reason='ARGUMENTS_INVALID',
                     success=False, exit_code=2, gate2_acceptance=False,
                     product_acceptance=False, next_action='STOP_PM_REVIEW')
    public = {key: value.get(key) for key in ('schema', 'status', 'reason', 'success', 'exit_code',
              'case_id', 'runtime_status', 'diagnostic_only', 'diagnostic_recalculation_verified',
              'p1_status', 'annual1_status', 'gate2_acceptance', 'product_acceptance',
              'force_termination', 'worker_still_running', 'next_action')}
    print(json.dumps(public, ensure_ascii=False))
    return value['exit_code']


if __name__ == '__main__':
    raise SystemExit(main())
