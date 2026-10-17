"""Reuse unchanged fixtures/expected values, generating just the approved case."""
from pathlib import Path
from datetime import date
import re
import shutil
import time
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, coordinate_to_tuple
from tools.s1_tiny_diagnostic import SPEC
from tools.s1_d2_diagnostic import EXPECTED as EXPECTED24
from tools.p2a_s1_cases import extra_cases, prepare_regression
from tools.verify_p2a_s1 import formula_digest, verify_entry, saved_calculation_mode
from .common import GateError, sha, write_json, read_json

REPO = Path(__file__).resolve().parents[2]
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

# Only the small functions24 verifier uses these fail-closed resource bounds.
F24_ROWS, F24_COLUMNS, F24_CELLS = 512, 128, 131072
F24_ARCHIVE, F24_EXPANDED, F24_SECONDS = 2 * 1024**2, 16 * 1024**2, 10


def _f24_deadline(started):
    if time.monotonic() - started > F24_SECONDS:
        raise GateError('BLOCKED', 'F24_VERIFIER_TIME_LIMIT')


def _bounded24(path, started):
    """Bound metadata AND actual XML coordinates before existing readers run."""
    if path.stat().st_size > F24_ARCHIVE:
        raise GateError('BLOCKED', 'F24_ARCHIVE_LIMIT')
    with zipfile.ZipFile(path) as z:
        entries = z.infolist()
        names = [e.filename for e in entries]
        if (len(entries) > 64 or len(set(names)) != len(names)
                or sum(e.file_size for e in entries) > F24_EXPANDED
                or any(e.flag_bits & 1 or '..' in e.filename.split('/')
                       or e.filename.startswith('/') for e in entries)
                or any('externalLink' in n or 'vbaProject' in n or n == 'xl/connections.xml' for n in names)):
            raise GateError('BLOCKED', 'F24_ZIP_LIMIT_OR_STRUCTURE')
        sheets, cells, area = 0, 0, 0
        for name in names:
            if not name.endswith('.xml'):
                continue
            raw = z.read(name)  # CRC and declared expansion limits checked by zipfile
            if b'<!DOCTYPE' in raw.upper() or b'<!ENTITY' in raw.upper():
                raise GateError('BLOCKED', 'F24_XML_DECLARATION')
            if not name.startswith('xl/worksheets/'):
                _f24_deadline(started)
                continue
            if not re.fullmatch(r'xl/worksheets/sheet[0-9]+\.xml', name):
                raise GateError('BLOCKED', 'F24_WORKSHEET_PART')
            sheets += 1
            root = ET.fromstring(raw)
            dimension = root.find('m:dimension', NS)
            if dimension is None:
                raise GateError('BLOCKED', 'F24_DIMENSION_MISSING')
            bounds = range_boundaries(dimension.get('ref', ''))
            if bounds is not None and (None in bounds or not 1 <= bounds[0] <= bounds[2] <= F24_COLUMNS
                                       or not 1 <= bounds[1] <= bounds[3] <= F24_ROWS):
                raise GateError('BLOCKED', 'F24_DIMENSION_LIMIT')
            area += bounds[2] * bounds[3]
            if area > F24_CELLS:
                raise GateError('BLOCKED', 'F24_UNION_LIMIT')
            last_row = 0
            for row in root.findall('m:sheetData/m:row', NS):
                r = int(row.get('r', '0'))
                if not last_row < r <= F24_ROWS:
                    raise GateError('BLOCKED', 'F24_ROW_LIMIT_OR_ORDER')
                last_row, last_col = r, 0
                for cell in row.findall('m:c', NS):
                    cr, cc = coordinate_to_tuple(cell.get('r', ''))
                    if cr != r or not last_col < cc <= F24_COLUMNS:
                        raise GateError('BLOCKED', 'F24_CELL_LIMIT_OR_ORDER')
                    if bounds and not (bounds[0] <= cc <= bounds[2] and bounds[1] <= cr <= bounds[3]):
                        raise GateError('BLOCKED', 'F24_DIMENSION_HIDES_CELL')
                    last_col = cc
                    cells += 1
                _f24_deadline(started)
            if sheets > 8 or cells > F24_CELLS:
                raise GateError('BLOCKED', 'F24_SHEET_OR_CELL_LIMIT')
        _f24_deadline(started)


def _nonformula24(before, saved, started):
    """Compare every position in the bounded union, including EmptyCell/None."""
    total = 0
    for name in dict.fromkeys([*before.sheetnames, *saved.sheetnames]):
        left = before[name] if name in before.sheetnames else None
        right = saved[name] if name in saved.sheetnames else None
        # Dimensions have been checked against actual XML cells, not trusted blindly.
        dims = [(s.max_row or 1, s.max_column or 1) if s is not None else (1, 1) for s in (left, right)]
        rows, cols = max(d[0] for d in dims), max(d[1] for d in dims)
        total += rows * cols
        if rows > F24_ROWS or cols > F24_COLUMNS or total > F24_CELLS:
            raise GateError('BLOCKED', 'F24_UNION_LIMIT')
        streams = [s.iter_rows(min_row=1, min_col=1, max_row=rows, max_col=cols)
                   if s is not None else iter([(None,) * cols] * rows) for s in (left, right)]
        # ReadOnlyWorksheet can end at the final serialized row even when a
        # larger max_row is requested. Missing trailing rows are all blank;
        # explicitly compare them rather than truncating zip at the shorter side.
        empty = (None,) * cols
        for row_number in range(1, rows + 1):
            pair = [next(stream, empty) for stream in streams]
            for column_number, (a, b) in enumerate(zip(*pair, strict=True), 1):
                if a is not None and a.data_type == 'f':
                    continue  # Formula identity is independently AND-checked, unchanged.
                av, bv = a.value if a is not None else None, b.value if b is not None else None
                same = (av == bv and isinstance(av, bool) == isinstance(bv, bool)
                        and (not isinstance(av, date) and not isinstance(bv, date) or type(av) is type(bv)))
                if not same or b is not None and b.data_type == 'f':
                    raise ValueError(f'24式の非数式値が変更されています: {name} R{row_number}C{column_number}')
            _f24_deadline(started)


def static_input(path):
    with zipfile.ZipFile(path) as z:
        if z.testzip() is not None:
            raise ValueError('破損ZIP')
        if any('externalLink' in n or 'vbaProject' in n or n == 'xl/connections.xml' for n in z.namelist()):
            raise ValueError('外部リンク/マクロ/接続は禁止')
        count = 0
        for name in z.namelist():
            if name.startswith('xl/worksheets/') and name.endswith('.xml'):
                root = ET.fromstring(z.read(name))
                for cell in root.findall('.//m:c', NS):
                    f = cell.find('m:f', NS)
                    if f is None:
                        continue
                    if f.get('t') in ('array', 'shared') or len(f.text or '') + 1 > 8192:
                        raise ValueError('Gate 2入力は8192文字以下の通常式だけです')
                    v = cell.find('m:v', NS)
                    if v is not None and v.text:
                        raise ValueError('生成入力に数式キャッシュがあります')
                    count += 1
    return count


def prepare(stage, directory, run_id):
    directory = Path(directory)
    inputs = directory / 'inputs'
    inputs.mkdir()
    if stage == 'tiny3':
        original = REPO / 'output/s1_tiny_diagnostic/tiny_calc.xlsx'
        record = read_json(original.parent / 'manifest.json')
        if record['sha256'] != sha(original) or record['spec'] != SPEC:
            raise ValueError('既存3式の正本/定義が一致しません')
        file = original.name
        shutil.copyfile(original, inputs / file)
        edits = [[SPEC['sheet'], address, value] for address, value in SPEC['inputs_after'].items()]
        probes = [[SPEC['sheet'], a, v] for a, v in SPEC['expected_after'].items()]
        challenges = []
        wanted_count = 3
    elif stage == 'functions24':
        original = REPO / 'output/s1_compatibility/01_function_gate.xlsx'
        record = next(r for r in read_json(original.parent / 'manifest.json') if r['file'] == original.name)
        if record['sha256'] != sha(original):
            raise ValueError('既存24式の正本が一致しません')
        file = original.name
        shutil.copyfile(original, inputs / file)
        # Only the operation COPY is temporarily perturbed. Restore B8=100,
        # then compare all 24 UNCHANGED original fixed expectations.
        challenges = [['確認', 'B8', 0]]
        edits = [['確認', 'B8', 100]]
        probes = [['確認', a, v] for a, v in EXPECTED24.items()]
        wanted_count = 24
    elif stage == 'annual1':
        inputs.rmdir()  # existing generator requires a new/empty directory
        case = next(c for c in extra_cases() if c.name == 'slot200_add')
        packet = prepare_regression(inputs, cases=[case])
        record = packet['entries'][0]
        if len(packet['entries']) != 1 or record['expected_cash'] != 1100000 or record['expected_total'] != 1600000:
            raise ValueError('既存1年モデルの固定期待値不一致')
        file = record['file']
        edits = [[e['sheet'], e['cell'], e['value']] for e in record['edits']]
        # Preserve date type from the old manifest; worker writes OLE dates.
        for edit, old in zip(edits, record['edits']):
            if old['type'] == 'date':
                edit[2] = {'date': old['value']}
        probes = [['入力チェック', 'B2', 'OK'], ['月次CF', 'E13', 1100000], ['月次CF', 'F13', 1600000]]
        challenges = []
        wanted_count = None
    else:
        raise ValueError('承認された3段階以外は実行しません')
    count = static_input(inputs / file)
    if wanted_count is not None and count != wanted_count:
        raise ValueError('数式数不一致')
    manifest = dict(run_id=run_id, case_id=stage, file=file, sha256=sha(inputs / file),
                    formula_count=count, probes=probes, edits=edits, challenges=challenges)
    write_json(directory / 'case.json', manifest)
    (directory / 'work').mkdir()
    (directory / 'results').mkdir()
    shutil.copyfile(inputs / file, directory / 'work' / file)
    return manifest


def compare_probes(manifest, value):
    differences = []
    for sheet, cell, expected in manifest['probes']:
        actual = value(sheet, cell)
        if actual != expected or isinstance(actual, bool) != isinstance(expected, bool):
            differences.append(dict(sheet=sheet, cell=cell, expected=expected, actual=actual))
    return differences


def verify_saved(directory, manifest):
    """No writes, no formula evaluation, no synthetic caches."""
    if not __debug__:
        raise ValueError('既存照合器のassertを無効化するpython -Oは使用できません')
    directory = Path(directory)
    original = directory / 'inputs' / manifest['file']
    saved = directory / 'results' / manifest['file']
    started = time.monotonic()
    if manifest['case_id'] == 'functions24':
        _bounded24(original, started)
        _bounded24(saved, started)
    if sha(original) != manifest['sha256']:
        raise ValueError('入力原本が変更されました')
    if formula_digest(original) != formula_digest(saved):
        raise ValueError('数式の消失/変更')
    saved_calculation_mode(saved)
    if manifest['case_id'] == 'annual1':
        entry = read_json(directory / 'inputs' / 'manifest.json')['entries'][0]
        verify_entry(entry, directory / 'inputs', directory / 'results')
    w = load_workbook(saved, read_only=True, data_only=True)
    try:
        differences = compare_probes(manifest, lambda s, c: w[s][c].value)
        if differences:
            raise ValueError('保存キャッシュの固定値不一致: ' + str(differences))
        errors = [(s.title, c.coordinate, c.value) for s in w for row in s for c in row if c.data_type == 'e']
        if errors:
            raise ValueError('Excel数式エラー: ' + str(errors[:20]))
    finally:
        w.close()
    w = load_workbook(saved, read_only=True, data_only=False)
    try:
        expected_inputs = {**SPEC['inputs_before'], **SPEC['inputs_after']} if manifest['case_id'] == 'tiny3' else None
        if expected_inputs and any(w[SPEC['sheet']][c].value != v for c, v in expected_inputs.items()):
            raise ValueError('3式の入力変更が未反映')
        if manifest['case_id'] == 'functions24':
            before = load_workbook(original, read_only=True, data_only=False)
            try:
                _nonformula24(before, w, started)
            finally:
                before.close()
    finally:
        w.close()
    return dict(saved_sha256=sha(saved), numeric_comparison='matched',
                formula_count=manifest['formula_count'], all_months_and_years_checked=manifest['case_id'] == 'annual1')
