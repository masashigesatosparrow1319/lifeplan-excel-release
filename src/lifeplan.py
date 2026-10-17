"""P1: integer-yen monthly cash flow and independently recalculable Excel."""
from dataclasses import dataclass, field, asdict
from datetime import date
from pathlib import Path
import argparse
import hashlib
import json
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


MAX_AMOUNT = 1_000_000_000
MAX_WEIGHT = 1_000


def bounded_integer(value, maximum, label):
    """Match Excel's numeric integer check; 1.0 is integer-valued, 0.1 is not."""
    if type(value) not in (int, float) or not 0 <= value <= maximum or not math.isfinite(value) or value != int(value):
        raise ValueError(f'{label}は0〜{maximum:,}の整数です')
    return int(value)


def months(start, end):
    if start > end:
        raise ValueError('開始年月が終了年月より後です')
    y, m = start
    while (y, m) <= end:
        yield y, m
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)


def allocate(amount, weights):
    """Largest remainder; equal remainders prefer earlier calendar months.

    Allocate the complete annual amount before filtering effective months.
    Integer products and divmod avoid floating-point remainder comparisons.
    """
    amount = bounded_integer(amount, MAX_AMOUNT, '年額')
    if len(weights) != 12:
        raise ValueError('月別比率は12個必要です')
    ws = [bounded_integer(w, MAX_WEIGHT, '月別比率') for w in weights]
    total = sum(ws)
    if total == 0:
        raise ValueError('月別比率の合計が正である必要があります')
    parts = [divmod(amount * w, total) for w in ws]
    base = [q for q, remainder in parts]
    remaining = amount - sum(base)
    eligible = [m for m, w in enumerate(ws) if w > 0]
    priority = sorted(eligible, key=lambda m: (-parts[m][1], m))
    for month in priority[:remaining]:
        base[month] += 1
    return base


@dataclass
class Person:
    id: str
    role: str
    birth: date
    planned: bool = False


@dataclass
class Account:
    id: str
    owner: str
    purpose: str
    balance: int
    restricted: bool = False
    instrument: str = '現預金'


@dataclass
class Rule:
    id: str
    target: str
    kind: str
    annual: int
    start: tuple
    end: tuple
    account: str
    weights: list = field(default_factory=lambda: [1]*12)


@dataclass
class Plan:
    start: tuple
    end: tuple
    persons: list
    accounts: list
    rules: list

    def validate(self):
        def unique(items):
            ids=[x.id for x in items]
            if any(not x for x in ids) or len(ids)!=len(set(ids)):
                raise ValueError('IDが空または重複しています')
        list(months(self.start,self.end))
        for ym in [self.start,self.end]+[v for r in self.rules for v in (r.start,r.end)]:
            date(*ym,1)
        for items in (self.persons,self.accounts,self.rules): unique(items)
        if not 1 <= sum(p.role=='adult' for p in self.persons) <= 2 or sum(p.role=='child' for p in self.persons)>6:
            raise ValueError('大人1〜2人、子ども0〜6人です')
        if any(p.role not in ('adult','child') or (p.role=='adult' and p.planned) for p in self.persons):
            raise ValueError('人物区分が不正です')
        owners={p.id for p in self.persons}|{'household'}
        accounts={a.id:a for a in self.accounts}
        if not accounts: raise ValueError('口座が必要です')
        for a in self.accounts:
            bounded_integer(a.balance, MAX_AMOUNT, '初期現預金残高')
            if a.owner not in owners or a.instrument!='現預金':
                raise ValueError('口座設定が不正です')
            if a.owner in {p.id for p in self.persons if p.role=='child'} and not a.restricted:
                raise ValueError('子ども口座は利用制限が必要です')
        for r in self.rules:
            allocate(r.annual,r.weights)
            if not r.target or r.kind not in ('income','expense') or r.start>r.end or r.account not in accounts:
                raise ValueError('収支ルールが不正です')
            if accounts[r.account].restricted:
                raise ValueError('P1では制限口座への収支計上は対象外です')
        for i,r in enumerate(self.rules):
            for s in self.rules[i+1:]:
                if r.target==s.target and max(r.start,s.start)<=min(r.end,s.end):
                    raise ValueError('同一対象の期間重複: '+r.target)


def calculate(plan):
    plan.validate()
    balances={a.id:bounded_integer(a.balance,MAX_AMOUNT,'初期現預金残高') for a in plan.accounts}
    available={a.id for a in plan.accounts if not a.restricted}
    allocations={r.id:allocate(r.annual,r.weights) for r in plan.rules}
    rows=[]
    for ym in months(plan.start,plan.end):
        income=expense=0
        opening=sum(balances[a] for a in available)
        for r in plan.rules:
            if r.start<=ym<=r.end:
                value=allocations[r.id][ym[1]-1]
                balances[r.account]+=value if r.kind=='income' else -value
                if r.kind=='income': income+=value
                else: expense+=value
        cash=sum(balances[a] for a in available)
        rows.append(dict(month=ym,opening=opening,income=income,expense=expense,cash=cash,total=sum(balances.values()),shortage=max(0,-cash),accounts=dict(balances)))
    return rows


def annualize(rows):
    result={}
    for r in rows:
        y=r['month'][0]
        if y not in result: result[y]=dict(income=0,expense=0,opening=r['opening'])
        result[y]['income']+=r['income']; result[y]['expense']+=r['expense']
        result[y].update(cash=r['cash'],total=r['total'])
    return result


def demo():
    return Plan((2026,1),(2026,12),[Person('adult1','adult',date(1990,1,1))],
        [Account('cash','adult1','生活・予備',1000000),Account('reserve','adult1','教育',500000,True)],
        [Rule('salary','salary','income',3600000,(2026,1),(2026,12),'cash'),Rule('living','living','expense',2400000,(2026,1),(2026,12),'cash')])


def verification_identity(plan):
    payload=json.dumps(asdict(plan),sort_keys=True,ensure_ascii=False,default=str)
    return (hashlib.sha256(payload.encode('utf-8')).hexdigest(),
            hashlib.sha256(Path(__file__).read_text(encoding='utf-8').encode('utf-8')).hexdigest())


def generate(plan,path):
    plan.validate()
    w=Workbook(); w.remove(w.active)
    info=w.create_sheet('説明'); settings=w.create_sheet('設定'); people=w.create_sheet('人物'); accounts=w.create_sheet('口座'); rules=w.create_sheet('収支'); allocation=w.create_sheet('月別配分'); events=w.create_sheet('月次明細'); monthly=w.create_sheet('月次CF'); annual=w.create_sheet('年次CF'); checks=w.create_sheet('入力チェック'); summary=w.create_sheet('不足サマリ')
    helper=w.create_sheet('配分根拠')
    meta=w.create_sheet('_検証メタ')
    input_hash,code_hash=verification_identity(plan)
    meta.append(['入力SHA256',input_hash]); meta.append(['生成コードSHA256',code_hash])
    meta.sheet_state='hidden'
    helper.append(['ルールID']+[f'{m}月切捨額' for m in range(1,13)]+[f'{m}月剰余' for m in range(1,13)]+[f'{m}月順位' for m in range(1,13)]+['残り円数'])
    for line in ['P1 開発版・ダミーデータ。金額は円、給与は手取り。','月末のみの不足判定。月内不足は未判定。自動振替・売却・借入なし。','黄色セルを編集。収支の期間・比率は変更可能。行追加と試算期間変更は再生成が必要。','年額は12か月分。最大剰余法で端数を配分し、同剰余なら1月から優先。比率0の月は常に0円。','制限口座は利用可能現金に含めない。口座単位の負残高も明細で確認する。','入力チェックと不足サマリを確認。入力エラーのある結果は利用しない。','年金・投資・住宅等は未実装。']:
        info.append([line])
    settings.append(['開始年月','終了年月']); settings.append([date(*plan.start,1),date(*plan.end,1)])
    people.append(['ID','区分','生年月日/出生予定日','出生予定'])
    for p in plan.persons: people.append([p.id,p.role,p.birth,p.planned])
    accounts.append(['ID','所有者','用途','保有形態','初期残高','利用制限'])
    for a in plan.accounts: accounts.append([a.id,a.owner,a.purpose,a.instrument,a.balance,int(a.restricted)])
    rules.append(['ID','対象ID','区分','年間手取り/支出','開始月','終了月','口座ID']+[f'{i}月比率' for i in range(1,13)])
    allocation.append(['ルールID']+[f'{i}月円' for i in range(1,13)])
    for i,r in enumerate(plan.rules,2):
        rules.append([r.id,r.target,r.kind,r.annual,date(*r.start,1),date(*r.end,1),r.account]+r.weights)
        bases=[]; remainders=[]; ranks=[]; amounts=[]
        for m in range(12):
            weight=f"'収支'!{get_column_letter(8+m)}{i}"
            quota=f"'収支'!$D{i}*{weight}/SUM('収支'!$H{i}:$S{i})"
            base=f'{get_column_letter(2+m)}{i}'
            remainder=f'{get_column_letter(14+m)}{i}'
            rank=f'{get_column_letter(26+m)}{i}'
            bases.append(f'=IF({weight}=0,0,ROUNDDOWN({quota},0))')
            # P <= 10^12; Q*W <= P. Both products are exactly representable.
            # Compare integer P-Q*W directly: no MOD or fractional key rounding.
            total=f"SUM('収支'!$H{i}:$S{i})"
            remainders.append(f'=IF({weight}=0,0,\'収支\'!$D{i}*{weight}-{base}*{total})')
            # COUNTIF over the prefix deterministically breaks equal-remainder ties.
            ranks.append(f'=IF({weight}=0,13,COUNTIF($N{i}:$Y{i},">"&{remainder})+COUNTIF($N{i}:{remainder},{remainder}))')
            amounts.append(f'=IF({weight}=0,0,\'配分根拠\'!{base}+IF(\'配分根拠\'!{rank}<=\'配分根拠\'!$AL{i},1,0))')
        helper.append([f"='収支'!A{i}"]+bases+remainders+ranks+[f"='収支'!D{i}-SUM(B{i}:M{i})"])
        allocation.append([f"='収支'!A{i}"]+amounts)
    events.append(['年月','口座ID','ルールID','収入','支出'])
    month_list=list(months(plan.start,plan.end))
    for ym in month_list:
        for j,r in enumerate(plan.rules,2):
            i=events.max_row+1
            base=f"IF(AND(A{i}>='収支'!E{j},A{i}<='収支'!F{j}),INDEX('月別配分'!B{j}:M{j},1,MONTH(A{i})),0)"
            events.append([date(*ym,1),f"='収支'!G{j}",f"='収支'!A{j}",f'=IF(\'収支\'!C{j}="income",{base},0)',f'=IF(\'収支\'!C{j}="expense",{base},0)'])
    monthly.append(['年月','期首利用可能現金','収入','支出','月末利用可能現金','現預金総額','月末不足額','警告']+[a.id for a in plan.accounts])
    end=events.max_row
    for i,ym in enumerate(month_list,2):
        dt=date(*ym,1)
        row=[dt,f'=SUMIF(\'口座\'!F2:F{len(plan.accounts)+1},0,\'口座\'!E2:E{len(plan.accounts)+1})' if i==2 else f'=E{i-1}',f'=SUMIF(\'月次明細\'!A2:A{end},A{i},\'月次明細\'!D2:D{end})',f'=SUMIF(\'月次明細\'!A2:A{end},A{i},\'月次明細\'!E2:E{end})',f'=B{i}+C{i}-D{i}',None,f'=MAX(0,-E{i})',f'=IF(G{i}>0,"月末現金不足（月内未判定）","")']
        for j,a in enumerate(plan.accounts,2):
            col=get_column_letter(j+7)
            opening=f"'口座'!E{j}" if i==2 else f'{col}{i-1}'
            row.append(f'={opening}+SUMIFS(\'月次明細\'!D$2:D${end},\'月次明細\'!A$2:A${end},$A{i},\'月次明細\'!B$2:B${end},\'口座\'!A{j})-SUMIFS(\'月次明細\'!E$2:E${end},\'月次明細\'!A$2:A${end},$A{i},\'月次明細\'!B$2:B${end},\'口座\'!A{j})')
        row[5]=f'=SUM(I{i}:{get_column_letter(8+len(plan.accounts))}{i})'
        monthly.append(row)
    annual.append(['年','対象月数','収入','支出','年末利用可能現金','年末現預金総額'])
    for year in sorted({y for y,m in month_list}):
        indices=[i for i,ym in enumerate(month_list,2) if ym[0]==year]; first,last=indices[0],indices[-1]
        annual.append([year,len(indices),f"=SUM('月次CF'!C{first}:C{last})",f"=SUM('月次CF'!D{first}:D{last})",f"='月次CF'!E{last}",f"='月次CF'!F{last}"])
    last_month=monthly.max_row
    summary.append(['項目','結果'])
    summary.append(['初回月末不足年月',f'=IF(COUNTIF(\'月次CF\'!G2:G{last_month},">0")=0,"なし",INDEX(\'月次CF\'!A2:A{last_month},MATCH(1,INDEX((\'月次CF\'!G2:G{last_month}>0)*1,0),0)))'])
    summary['B2'].number_format='yyyy-mm'
    summary.append(['最大月末不足額',f"=MAX('月次CF'!G2:G{last_month})"])
    summary.append(['月内現金不足','未判定（P1対象外）'])
    checks.append(['対象','判定'])
    checks.append(['試算期間（変更時は再生成）',f'=IF(AND(\'設定\'!A2=DATE({plan.start[0]},{plan.start[1]},1),\'設定\'!B2=DATE({plan.end[0]},{plan.end[1]},1)),"OK","要再生成")'])
    last_rule=rules.max_row; last_account=accounts.max_row
    for i,r in enumerate(plan.rules,2):
        ratio=f"'収支'!H{i}:S{i}"
        checks.append([r.id,f'=IFERROR(IF(AND(ISNUMBER(\'収支\'!D{i}),\'収支\'!D{i}>=0,\'収支\'!D{i}<={MAX_AMOUNT},\'収支\'!D{i}=INT(\'収支\'!D{i}),COUNT({ratio})=12,MIN({ratio})>=0,MAX({ratio})<={MAX_WEIGHT},SUMPRODUCT(--({ratio}=INT({ratio})))=12,SUM({ratio})>0,\'収支\'!E{i}<=\'収支\'!F{i},OR(\'収支\'!C{i}="income",\'収支\'!C{i}="expense"),COUNTIFS(\'口座\'!A2:A{last_account},\'収支\'!G{i},\'口座\'!F2:F{last_account},0)=1,SUMPRODUCT((\'収支\'!B2:B{last_rule}=\'収支\'!B{i})*(\'収支\'!E2:E{last_rule}<=\'収支\'!F{i})*(\'収支\'!F2:F{last_rule}>=\'収支\'!E{i}))=1),"OK","入力エラー：金額・整数比率・期間・口座・競合を確認"),"入力エラー：数値・整数比率を確認")'])
    for i,a in enumerate(plan.accounts,2):
        value=f"'口座'!E{i}"
        checks.append([f'口座:{a.id}',f'=IFERROR(IF(AND(ISNUMBER({value}),{value}>=0,{value}<={MAX_AMOUNT},{value}=INT({value})),"OK","入力エラー：初期残高は0〜10億円の整数"),"入力エラー：初期残高を確認")'])
    summary.append(['入力チェック',f'=IF(COUNTIF(\'入力チェック\'!B2:B{checks.max_row},"OK")={checks.max_row-1},"OK","入力エラー／要再生成：結果を利用しないでください")'])
    for sheet,cellrange,kind,f1,f2 in [(rules,f'D2:D{last_rule}','whole','0',str(MAX_AMOUNT)),(rules,f'H2:S{last_rule}','whole','0',str(MAX_WEIGHT)),(accounts,f'E2:E{last_account}','whole','0',str(MAX_AMOUNT)),(rules,f'C2:C{last_rule}','list','"income,expense"',None),(accounts,f'F2:F{last_account}','list','"0,1"',None)]:
        dv=DataValidation(type=kind,operator='between' if kind!='list' else None,formula1=f1,formula2=f2,allow_blank=False)
        dv.errorTitle='入力エラー'; dv.error='許可された値を入力してください'; dv.showErrorMessage=True; dv.errorStyle='stop'
        sheet.add_data_validation(dv); dv.add(cellrange)
    for s in w:
        s.freeze_panes='A2'
        s.auto_filter.ref=s.dimensions
        for c in s[1]: c.font=Font(bold=True)
        for col in s.columns: s.column_dimensions[col[0].column_letter].width=22
        for row in s.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,date): c.number_format='yyyy-mm-dd'
                if s in (people,accounts,rules) and c.data_type!='f': c.fill=PatternFill('solid',fgColor='FFF2CC')
    w.calculation=CalcProperties(calcId=0,fullCalcOnLoad=True,forceFullCalc=True)
    Path(path).parent.mkdir(parents=True,exist_ok=True); w.save(path)
    return w


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--output',default='output/lifeplan_dev.xlsx'); args=parser.parse_args()
    generate(demo(),args.output)
