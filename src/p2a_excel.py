"""Editable P2-A workbook using Excel 2010 functions; P1 code stays untouched."""
from dataclasses import asdict, replace
from datetime import date, datetime
from pathlib import Path
import hashlib
import json
import argparse
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as col
from . import lifeplan as p1
from .p2a import P2Plan, CashEvent, from_p1, calculate, MAX_EVENTS, MAX_RULES, month_number


def identity(plan):
    payload=json.dumps(asdict(plan),sort_keys=True,ensure_ascii=False,default=str)
    files=[Path(__file__),Path(__file__).with_name('p2a.py'),Path(p1.__file__)]
    source='\n'.join(f.read_text(encoding='utf-8') for f in files)
    return hashlib.sha256(payload.encode()).hexdigest(),hashlib.sha256(source.encode()).hexdigest()


def valid_code(ref):
    alphabet='ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789_-'
    return f'AND(ISTEXT({ref}),LEN({ref})>=1,LEN({ref})<=64,SUMPRODUCT(--ISNUMBER(FIND(MID({ref},ROW($A$1:$A$64),1),"{alphabet}")))=64)'


def literal(ref):
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({ref},"~","~~"),"*","~*"),"?","~?")'


def integer(ref,low,high):
    return f'AND(ISNUMBER({ref}),{ref}>={low},{ref}<={high},{ref}=INT({ref}))'


def ym_ok(ref):
    return f'AND(ISNUMBER({ref}),{ref}>=DATE(1900,1,1),{ref}<=DATE(9999,12,1),DAY({ref})=1,{ref}=INT({ref}))'


def add_validation(sheet,range_,type_,f1,f2=None):
    v=DataValidation(type=type_,formula1=f1,formula2=f2,operator='between' if type_=='whole' else None,allow_blank=True)
    v.showErrorMessage=True; v.errorStyle='stop'; v.errorTitle='入力確認'; v.error='許可された値を入力してください'
    sheet.add_data_validation(v); v.add(range_)


def generate(plan,path):
    if Path(path).resolve().parent == Path(__file__).resolve().parents[1]/'references':
        raise ValueError('参照原本の保存先には出力しません')
    calculate(plan)  # reject invalid input, double counting and overflow before export
    with TemporaryDirectory(prefix='lifeplan-p2a-') as temporary:
        # Reuse only the small accepted P1 formula template. Its long-form renderer
        # is deliberately not run for H*R large plans; P2-A builds bounded matrices.
        w=p1.generate(p1.demo(),Path(temporary)/'base.xlsx')
    w['設定']['A2']=date(*plan.start,1); w['設定']['B2']=date(*plan.end,1)
    for name in ('人物','口座','収支'):
        w[name].delete_rows(2,w[name].max_row-1)
    for p in plan.persons: w['人物'].append([p.id,p.role,p.birth,p.planned])
    for a in plan.accounts: w['口座'].append([a.id,a.owner,a.purpose,a.instrument,a.balance,int(a.restricted)])
    for r in plan.rules: w['収支'].append([r.id,r.target,r.kind,r.annual,date(*r.start,1),date(*r.end,1),r.account]+r.weights)
    monthly=w['月次CF']; annual=w['年次CF']
    monthly.delete_rows(1,monthly.max_row); annual.delete_rows(1,annual.max_row)
    monthly.append(['年月','期首利用可能現金','収入','支出','月末利用可能現金','現預金総額','月末不足額','警告']+[a.id for a in plan.accounts])
    month_list=list(p1.months(plan.start,plan.end))
    for i,ym in enumerate(month_list,2):
        monthly.append([date(*ym,1),f'=SUMIF(\'口座\'!F2:F{len(plan.accounts)+1},0,\'口座\'!E2:E{len(plan.accounts)+1})' if i==2 else f'=E{i-1}',0,0,f'=B{i}+C{i}-D{i}',f'=SUM(I{i}:{col(len(plan.accounts)+8)}{i})',f'=MAX(0,-E{i})',''])
    annual.append(['年','対象月数','収入','支出','年末利用可能現金','年末現預金総額'])
    for year in sorted({y for y,m in month_list}):
        indices=[i for i,ym in enumerate(month_list,2) if ym[0]==year]; first,last=indices[0],indices[-1]
        annual.append([year,len(indices),f"=SUM('月次CF'!C{first}:C{last})",f"=SUM('月次CF'!D{first}:D{last})",f"='月次CF'!E{last}",f"='月次CF'!F{last}"])
    last_cf=len(month_list)+1
    w['不足サマリ']['B2']=f'=IF(COUNTIF(\'月次CF\'!G2:G{last_cf},">0")=0,"なし",INDEX(\'月次CF\'!A2:A{last_cf},MATCH(1,INDEX((\'月次CF\'!G2:G{last_cf}>0)*1,0),0)))'
    w['不足サマリ']['B3']=f"=MAX('月次CF'!G2:G{last_cf})"
    for name in ['月次明細','入力チェック']:
        w.remove(w[name])
    rules=w['収支']; events=w.create_sheet('イベント入力'); helper=w.create_sheet('イベント補助')
    used_ids={str(r.id).upper() for r in plan.rules}|{str(e.id).upper() for e in plan.events}
    def fresh_id(prefix):
        n=1
        while f'{prefix}{n:03}' in used_ids: n+=1
        result=f'{prefix}{n:03}'; used_ids.add(result)
        return result
    allocation=w['月別配分']; roots=w['配分根拠']
    root_template=[c.value for c in roots[2]]
    allocation_template=[c.value for c in allocation[2]]
    roots.delete_rows(2,roots.max_row-1); allocation.delete_rows(2,allocation.max_row-1)
    checks=w.create_sheet('入力チェック'); matrix=w.create_sheet('月次計上'); duplicates=w.create_sheet('重複確認')
    prices=w.create_sheet('年価格'); detail=w.create_sheet('選択月明細')
    rules.cell(1,20,'費用項目コード'); rules.cell(1,21,'入力状態')
    keys=plan.keys()
    for j in range(2,52):
        if j-2<len(plan.rules): rules.cell(j,20,keys[plan.rules[j-2].id])
        else:
            for k in range(2,21): rules.cell(j,k).value=None
            rules.cell(j,1,fresh_id('R'))
            for k in range(8,20): rules.cell(j,k,1)
        rules.cell(j,21,f'=IF(COUNTA(B{j}:G{j},T{j})=0,0,IFERROR(IF(AND(ISTEXT(A{j}),LEN(A{j})>0,ISTEXT(B{j}),LEN(TRIM(B{j}))>0,OR(C{j}="income",C{j}="expense"),{integer(f"D{j}",0,10**9)},{ym_ok(f"E{j}")},{ym_ok(f"F{j}")},E{j}<=F{j},COUNT(H{j}:S{j})=12,MIN(H{j}:S{j})>=0,MAX(H{j}:S{j})<=1000,SUMPRODUCT(--(H{j}:S{j}=INT(H{j}:S{j})))=12,SUM(H{j}:S{j})>0,COUNTIFS(\'口座\'!$A$2:$A${len(plan.accounts)+1},{literal(f"G{j}")},\'口座\'!$F$2:$F${len(plan.accounts)+1},0)=1,{valid_code(f"T{j}")},COUNTIFS($A$2:$A$51,{literal(f"A{j}")},$B$2:$B$51,"<>")=1,SUMPRODUCT(--($B$2:$B$51=B{j}),--($E$2:$E$51<=F{j}),--($F$2:$F$51>=E{j}))=1),1,-1),-1))')
        # Keep the accepted P1 allocation formulas, extending with guarded copies.
        from openpyxl.formula.translate import Translator
        for k,value in enumerate(root_template,1):
            roots.cell(j,k,Translator(value,origin=f'{col(k)}2').translate_formula(f'{col(k)}{j}'))
        for k,value in enumerate(allocation_template,1):
            allocation.cell(j,k,Translator(value,origin=f'{col(k)}2').translate_formula(f'{col(k)}{j}'))
        for sheet in (roots,allocation):
            for c in sheet[j]:
                if c.data_type=='f': c.value=f'=IF(\'収支\'!$U{j}=1,{c.value[1:]},0)'
    headers=['ID（自動）','名称','タイプ（空欄=単発）','収支（空欄=支出）','発生/初回年月','最終年月','間隔（月）','1回金額（円）','口座ID','対象者（空欄=世帯）','用途（空欄=その他）','費用項目コード（通常空欄）','追加分確認','物価補正0/1','価格基準年月','年率bp（空欄=0）','夢ID','備考','有効0/1','入力状態']
    events.append(headers)
    helper.append(['有効入力','コード','区分','口座','初回月','最終月','間隔','補正','基準月','年率bp','ID','入力あり'])
    for i in range(2,202):
        events.cell(i,1,fresh_id('E') if i-2>=len(plan.events) else plan.events[i-2].id); events.cell(i,19,1)
        if i-2<len(plan.events):
            e=plan.events[i-2]
            first=e.occurrence_month if e.event_type=='once' else e.first_month
            values=[e.id,e.name,'単発' if e.event_type=='once' else '周期','収入' if e.kind=='income' else '支出',date(*first,1) if first else None,date(*e.last_month,1) if e.last_month else None,e.interval_months,e.amount,e.account_id,e.person_id,e.purpose,e.economic_key or None,'はい' if e.additional_confirmed else None,int(e.inflation_enabled),date(*e.base_month,1) if e.base_month else None,e.rate_bps,e.linked_dream_id or None,e.note or None,int(e.enabled)]
            for k,v in enumerate(values,1): events.cell(i,k,v)
        # Empty rows ignore internal IDs and preset enabled switches.
        raw=f"'イベント入力'!"
        def er(c): return f'{raw}{c}{i}'
        id_=er('A'); first=er('E'); amount=er('H'); kind=er('D'); typ=er('C'); interval=er('G'); last=er('F'); key=er('L'); base=er('O'); rate=er('P')
        settings=f"AND({valid_code(id_)},ISTEXT({er('B')}),LEN(TRIM({er('B')}))>0,OR({typ}=\"\",{typ}=\"単発\",{typ}=\"周期\"),OR({kind}=\"\",{kind}=\"支出\",{kind}=\"収入\"),{ym_ok(first)},{integer(amount,0,10**9)},COUNTIFS('口座'!$A$2:$A${len(plan.accounts)+1},{er('I')},'口座'!$F$2:$F${len(plan.accounts)+1},0)=1,OR({er('J')}=\"\",{er('J')}=\"household\",COUNTIF('人物'!$A$2:$A${len(plan.persons)+1},{er('J')})=1),OR({key}=\"\",{valid_code(key)}),{er('M')}=\"はい\",OR({er('N')}=\"\",{er('N')}=0,{er('N')}=1),IF({typ}=\"周期\",AND(OR({kind}=\"\",{kind}=\"支出\"),{ym_ok(last)},{last}>={first},{integer(interval,1,120000)}),AND({last}=\"\",{interval}=\"\")),IF({er('N')}=1,AND(OR({kind}=\"\",{kind}=\"支出\"),{ym_ok(base)},{first}>={base},OR({rate}=\"\",{integer(rate,-1000,1000)})),TRUE))"
        settings=f'AND({settings},OR({er("K")}="",AND(ISTEXT({er("K")}),LEN(TRIM({er("K")}))>0)))'
        settings=settings.replace(f",{er('I')},'口座'",f",{literal(er('I'))},'口座'")
        settings=settings.replace(f",{er('J')})=1",f",{literal(er('J'))})=1")
        state=f'=IF(COUNTA(\'イベント入力\'!B{i}:R{i})=0,0,IF(NOT(AND(ISNUMBER({er("S")}),OR({er("S")}=0,{er("S")}=1))),-1,IF({er("S")}=0,0,IFERROR(IF({settings},1,-1),-1))))'
        helper.append([state,f'=IF(A{i}=1,UPPER(IF({key}="",{id_},{key})),"")',f'=IF({kind}="収入","income","expense")',f'={er("I")}',f'=IF(A{i}=1,YEAR({first})*12+MONTH({first})-1,0)',f'=IF(A{i}=1,IF({typ}="周期",YEAR({last})*12+MONTH({last})-1,E{i}),0)',f'=IF(A{i}=1,IF({typ}="周期",{interval},1),1)',f'=IF({er("N")}=1,1,0)',f'=IF(AND(A{i}=1,H{i}=1),YEAR({base})*12+MONTH({base})-1,0)',f'=IF({rate}="",0,{rate})',f'={id_}',f'=IF(COUNTA(\'イベント入力\'!B{i}:R{i})=0,0,1)'])
        first_h=f"'イベント補助'!E{i}"; last_h=f"'イベント補助'!F{i}"; interval_h=f"'イベント補助'!G{i}"
        next_due=f'{first_h}+MAX(0,INT(({month_number(plan.start)}-{first_h}+{interval_h}-1)/{interval_h}))*{interval_h}'
        in_range=f'{next_due}<=MIN({last_h},{month_number(plan.end)})'
        events.cell(i,20,f'=IF(\'イベント補助\'!A{i}=-1,"入力エラー",IF(S{i}=0,"無効",IF(\'イベント補助\'!A{i}=0,"空白",IF({in_range},"OK","試算範囲外"))))')
    # Price recurrence is bounded by the Excel calendar; freeze overflow, never clamp it.
    year_count=plan.end[0]-1900
    prices.append(['経過年']+[f'=\'イベント入力\'!A{i}' for i in range(2,202)])
    for n in range(year_count+1):
        row=n+2; prices.cell(row,1,n)
        for s in range(200):
            c=col(s+2); i=s+2
            active=f"AND('イベント補助'!A{i}=1,'イベント補助'!H{i}=1)"
            body=f"'イベント入力'!H{i}" if n==0 else f'IF({c}{row-1}>1000000000,{c}{row-1},INT(({c}{row-1}*(10000+\'イベント補助\'!J{i})+5000)/10000))'
            prices.cell(row,s+2,f'=IF({active},{body},0)')
    # 50 rule + 200 event columns; dates start row 7. All dependencies stay in-row.
    matrix.cell(1,1,'ソースID'); matrix.cell(2,1,'費用項目'); matrix.cell(3,1,'収支'); matrix.cell(4,1,'口座'); matrix.cell(5,1,'入力状態'); matrix.cell(6,1,'年月')
    duplicates.cell(1,1,'費用項目'); duplicates.cell(2,1,'前の同項目列')
    for s in range(250):
        c=col(s+2)
        if s<50:
            i=s+2
            metadata=[f"'収支'!A{i}",f'IF(\'収支\'!U{i}=1,UPPER(\'収支\'!T{i}),"")',f"'収支'!C{i}",f"'収支'!G{i}",f"'収支'!U{i}"]
        else:
            i=s-50+2
            metadata=[f"'イベント補助'!K{i}",f"'イベント補助'!B{i}",f"'イベント補助'!C{i}",f"'イベント補助'!D{i}",f"'イベント補助'!A{i}"]
        for rr,v in enumerate(metadata,1): matrix.cell(rr,s+2,'='+v)
        duplicates.cell(1,s+2,f"='月次計上'!{c}2")
        if s==0: prev='=0'
        else:
            previous=col(s+1)
            prev=f'=IF({c}1="",0,IFERROR(LOOKUP(2,1/($B$1:{previous}$1={c}1),COLUMN($B$1:{previous}$1)-1),0))'
        duplicates.cell(2,s+2,prev)
    for t,ym in enumerate(month_list):
        mr=t+7; dr=t+3; cf=t+2
        matrix.cell(mr,1,date(*ym,1)); duplicates.cell(dr,1,date(*ym,1))
        k=month_number(ym)
        for s in range(250):
            c=col(s+2)
            if s<50:
                i=s+2
                expr=f'IF(AND($A{mr}>=\'収支\'!E{i},$A{mr}<=\'収支\'!F{i}),INDEX(\'月別配分\'!B{i}:M{i},1,MONTH($A{mr})),0)'
            else:
                i=s-50+2; price_col=col(i)
                first=f"'イベント補助'!E{i}"; last=f"'イベント補助'!F{i}"; interval=f"'イベント補助'!G{i}"
                occurs_=f'AND({k}>={first},{k}<={last},{k}-{first}={interval}*INT(({k}-{first})/{interval}))'
                price=f'IF(\'イベント補助\'!H{i}=1,INDEX(\'年価格\'!{price_col}$2:{price_col}${year_count+2},INT(({k}-\'イベント補助\'!I{i})/12)+1),\'イベント入力\'!H{i})'
                expr=f'IF({occurs_},{price},0)'
            matrix.cell(mr,s+2,f'=IF({c}$5=1,{expr},0)')
            indicator=f'IF(\'月次計上\'!{c}{mr}>0,1,0)'
            previous=col(s+1)
            chain=indicator if s==0 else f'IF({c}$2>0,INDEX($B{dr}:{previous}{dr},1,{c}$2)+{indicator},{indicator})'
            duplicates.cell(dr,s+2,'='+chain)
        duplicates.cell(dr,252,f'=IF(MAX(B{dr}:IQ{dr})>1,1,0)')
        duplicates.cell(dr,253,f'=IF(MAX(\'月次計上\'!B{mr}:IQ{mr})>1000000000,1,0)')
        monthly=w['月次CF']
        monthly.cell(cf,3,f'=SUMIF(\'月次計上\'!$B$3:$IQ$3,"income",\'月次計上\'!B{mr}:IQ{mr})')
        monthly.cell(cf,4,f'=SUMIF(\'月次計上\'!$B$3:$IQ$3,"expense",\'月次計上\'!B{mr}:IQ{mr})')
        monthly.cell(cf,8,f'=IF(\'入力チェック\'!B2<>"OK","入力/重複エラー：結果を使用しない",IF(G{cf}>0,"月末不足（月内未判定）",""))')
        for a_index,a in enumerate(plan.accounts):
            ac=col(a_index+9); ar=a_index+2
            opening=f"'口座'!E{ar}" if t==0 else f'{ac}{cf-1}'
            account_criterion=literal(f"'口座'!A{ar}")
            monthly.cell(cf,a_index+9,f'={opening}+SUMIFS(\'月次計上\'!B{mr}:IQ{mr},\'月次計上\'!$B$4:$IQ$4,{account_criterion},\'月次計上\'!$B$3:$IQ$3,"income")-SUMIFS(\'月次計上\'!B{mr}:IQ{mr},\'月次計上\'!$B$4:$IQ$4,{account_criterion},\'月次計上\'!$B$3:$IQ$3,"expense")')
    # Identity/validation is static O(S^2) in headers, duplicate recurrence O(H*S).
    ids_range="'月次計上'!$B$1:$IQ$1"
    checks.append(['項目','判定']); checks.append(['全体',None])
    checks.append(['期間',f'=IF(AND(\'設定\'!A2=DATE({plan.start[0]},{plan.start[1]},1),\'設定\'!B2=DATE({plan.end[0]},{plan.end[1]},1)),"OK","要再生成")'])
    checks.append(['入力行', '=IF(AND(COUNTIF(\'収支\'!U2:U51,-1)=0,COUNTIF(\'イベント補助\'!A2:A201,-1)=0),"OK","入力エラー")'])
    for s in range(250):
        c=col(s+2)
        checks.append([f'ID:{s+1}',f'=IF(\'月次計上\'!{c}5<>1,"OK",IF(COUNTIFS({ids_range},{literal(f"\'月次計上\'!{c}1")},\'月次計上\'!$B$5:$IQ$5,1)=1,"OK","ID重複"))'])
    checks.append(['同月/同費用',f'=IF(SUM(\'重複確認\'!IR3:IR{len(month_list)+2})=0,"OK","二重計上")'])
    checks.append(['補正価格',f'=IF(SUM(\'重複確認\'!IS3:IS{len(month_list)+2})=0,"OK","補正価格上限超過")'])
    for i,a in enumerate(plan.accounts,2):
        owner=f"'口座'!B{i}"; id_=f"'口座'!A{i}"; restriction=f"'口座'!F{i}"
        valid=f'AND({integer(f"\'口座\'!E{i}",0,10**9)},ISTEXT({id_}),LEN({id_})>0,COUNTIF(\'口座\'!A2:A{len(plan.accounts)+1},{literal(id_)})=1,ISNUMBER({restriction}),OR({restriction}=0,{restriction}=1),\'口座\'!D{i}="現預金",OR({owner}="household",COUNTIF(\'人物\'!A2:A{len(plan.persons)+1},{literal(owner)})=1),IF(COUNTIFS(\'人物\'!A2:A{len(plan.persons)+1},{literal(owner)},\'人物\'!B2:B{len(plan.persons)+1},"child")>0,{restriction}=1,TRUE))'
        checks.append([a.id,f'=IFERROR(IF({valid},"OK","口座設定エラー"),"口座設定エラー")'])
    person_last=len(plan.persons)+1
    checks.append(['家族構成',f'=IF(AND(COUNTIF(\'人物\'!B2:B{person_last},"adult")>=1,COUNTIF(\'人物\'!B2:B{person_last},"adult")<=2,COUNTIF(\'人物\'!B2:B{person_last},"child")<=6),"OK","家族構成エラー")'])
    for i,p in enumerate(plan.persons,2):
        pi=f"'人物'!A{i}"; role=f"'人物'!B{i}"; planned=f"'人物'!D{i}"
        valid=f'AND(ISTEXT({pi}),LEN({pi})>0,COUNTIF(\'人物\'!A2:A{person_last},{literal(pi)})=1,OR({role}="adult",{role}="child"),OR(ISLOGICAL({planned}),{integer(planned,0,1)}),IF({role}="adult",{planned}=FALSE,TRUE))'
        checks.append([p.id,f'=IFERROR(IF({valid},"OK","人物設定エラー"),"人物設定エラー")'])
    checks.cell(2,2,f'=IF(COUNTIF(B3:B{checks.max_row},"OK")={checks.max_row-2},"OK","入力/重複エラー")')
    checks.append(['入力済Rule数（エラー含む）','=COUNTIF(\'収支\'!U2:U51,"<>0")'])
    checks.append(['入力済イベント数（無効含む）','=SUM(\'イベント補助\'!L2:L201)'])
    checks.append(['有効正常イベント数','=COUNTIF(\'イベント補助\'!A2:A201,1)'])
    checks.append(['容量','1200か月 / Rule50 / イベント200 / 口座20'])
    summary=w['不足サマリ']; summary['B5']="='入力チェック'!B2"
    detail.append(['選択年月',date(*plan.start,1)]); detail.append(['ID','名称','費用項目','対象者','用途','夢ID','収支','口座','金額'])
    detail.cell(1,3,'全体判定'); detail.cell(1,4,"='入力チェック'!B2")
    for s in range(250):
        c=col(s+2)
        if s<50:
            title=f"'収支'!B{s+2}"; person='"未指定"'; purpose='"未指定"'; dream='""'
        else:
            erow=s-50+2
            title=f"'イベント入力'!B{erow}"; person=f'IF(\'イベント入力\'!J{erow}="","household",\'イベント入力\'!J{erow})'; purpose=f'IF(\'イベント入力\'!K{erow}="","その他",\'イベント入力\'!K{erow})'; dream=f'IF(\'イベント入力\'!Q{erow}="","",\'イベント入力\'!Q{erow})'
        expressions=[f"'月次計上'!{c}1",title,f"'月次計上'!{c}2",person,purpose,dream,f'IF(\'月次計上\'!{c}3="income","収入","支出")',f"'月次計上'!{c}4",f'IFERROR(INDEX(\'月次計上\'!{c}7:{c}{len(month_list)+6},MATCH($B$1,\'月次計上\'!A7:A{len(month_list)+6},0)),"期間外")']
        detail.append([f'=IF(\'月次計上\'!{c}5=1,{v},"")' for v in expressions])
    info=w['説明']; info.delete_rows(1,info.max_row)
    for text in ['ライフプラン開発版：月末の収支・現金と大型支出を確認できます。','イベント入力: 名前・年月・1回金額・口座・追加分確認を入力。空欄タイプは単発、収支は支出。','IDは自動割当てで編集不要。費用項目コードは通常空欄。同じ支払いを紐付ける場合だけ同じコード。','夢IDは希望のグループ、費用項目コードは実際の支払い。別コードの意味上の重複は自動判定できない。','有効0で無効化。有効な空白行は無視。追加分確認は「はい」を選択。','周期: 初回年月・最終年月・正の月間隔。一括支出のみ。年月は月初の日付（例2026/7/1）。','物価: 支出のみ0/1。整数bp（200=2%）、基準月の周年ごとに円四捨五入。','入力チェック全体OKの結果だけを利用。不足は月末判定で、月内の順序は考慮していません。','入力枠は期間別収支50/イベント200、口座最大20、期間最大1200か月。期間/口座の追加は再生成。','ローン・売却・夢の自動CF反映は未対応。']:
        info.append([text])
    add_validation(events,'C2:C201','list','"単発,周期"'); add_validation(events,'D2:D201','list','"収入,支出"')
    add_validation(events,'H2:H201','whole','0','1000000000'); add_validation(events,'G2:G201','whole','1','120000')
    add_validation(rules,'D2:D51','whole','0','1000000000'); add_validation(rules,'H2:S51','whole','0','1000'); add_validation(rules,'C2:C51','list','"income,expense"')
    add_validation(events,'M2:M201','list','"はい"'); add_validation(events,'N2:N201','whole','0','1'); add_validation(events,'S2:S201','whole','0','1'); add_validation(events,'P2:P201','whole','-1000','1000')
    # Excel 2010 data validation cross-sheet references use defined names.
    from openpyxl.workbook.defined_name import DefinedName
    w.defined_names.add(DefinedName('利用口座',attr_text=f"'口座'!$A$2:$A${len(plan.accounts)+1}"))
    references=w.create_sheet('参照一覧'); references.append(['口座ID','対象者ID','費用項目コード','名称'])
    for i,a in enumerate(plan.accounts,2): references.cell(i,1,a.id)
    references.cell(2,2,'household')
    for i,p in enumerate(plan.persons,3): references.cell(i,2,p.id)
    for s in range(250):
        c=col(s+2); references.cell(s+2,3,f"='月次計上'!{c}2")
        references.cell(s+2,4,f"='収支'!B{s+2}" if s<50 else f"='イベント入力'!B{s-50+2}")
    w.defined_names.add(DefinedName('対象者一覧',attr_text=f"'参照一覧'!$B$2:$B${len(plan.persons)+2}"))
    w.defined_names.add(DefinedName('費用項目一覧',attr_text="'参照一覧'!$C$2:$C$251"))
    add_validation(events,'I2:I201','list','=利用口座'); add_validation(events,'J2:J201','list','=対象者一覧')
    add_validation(events,'L2:L201','list','=費用項目一覧')
    events.data_validations.dataValidation[-1].showErrorMessage=False  # new valid codes are also allowed
    add_validation(rules,'T2:T51','list','=費用項目一覧')
    rules.data_validations.dataValidation[-1].showErrorMessage=False
    w.defined_names.add(DefinedName('試算月一覧',attr_text=f"'月次CF'!$A$2:$A${len(month_list)+1}"))
    add_validation(detail,'B1','list','=試算月一覧')
    detail['B1'].number_format='yyyy-mm'
    for s in [events,rules,checks,detail]:
        s.freeze_panes='B2'
        for c in s[1]: c.font=Font(bold=True)
        for k in range(1,s.max_column+1): s.column_dimensions[col(k)].width=22
    for s in (helper,matrix,duplicates,prices): s.sheet_state='hidden'
    for i in range(2,202):
        for k in range(2,20):
            c=events.cell(i,k); c.fill=PatternFill('solid',fgColor='FFF2CC'); c.protection=Protection(locked=False)
        for k in (5,6,15): events.cell(i,k).number_format='yyyy-mm'
    events.protection.sheet=True; events.protection.autoFilter=False
    for i in range(2,52):
        for k in range(2,21):
            rules.cell(i,k).fill=PatternFill('solid',fgColor='FFF2CC')
            rules.cell(i,k).protection=Protection(locked=False)
        for k in (5,6): rules.cell(i,k).number_format='yyyy-mm'
    rules.protection.sheet=True
    h,c=identity(plan); w['_検証メタ']['B1']=h; w['_検証メタ']['B2']=c
    Path(path).parent.mkdir(parents=True,exist_ok=True); w.save(path)
    return w


def read_plan(path):
    """Read actual Excel input cells, excluding blank/disabled slots; no cached formulas."""
    w=load_workbook(path,data_only=False,read_only=True)
    for name,last,end_col in [('イベント入力',201,19),('収支',51,20)]:
        if w[name].max_row>last and any(c.value is not None for row in w[name].iter_rows(min_row=last+1,max_col=end_col) for c in row):
            raise ValueError('入力枠の容量を超えています')
    def ym(v):
        if not isinstance(v,(date,datetime)) or v.day!=1: raise ValueError('年月は月初の日付です')
        if isinstance(v,datetime) and (v.hour or v.minute or v.second or v.microsecond): raise ValueError('年月に時刻は含めません')
        return v.year,v.month
    persons=[]
    for r in w['人物'].iter_rows(min_row=2,values_only=True):
        if not r[0]: continue
        if type(r[3]) not in (bool,int,float) or r[3] not in (0,1): raise ValueError('出生予定は真偽値です')
        persons.append(p1.Person(r[0],r[1],r[2].date() if isinstance(r[2],datetime) else r[2],bool(r[3])))
    accounts=[]
    for r in w['口座'].iter_rows(min_row=2,values_only=True):
        if not r[0]: continue
        if type(r[5]) not in (int,float) or r[5] not in (0,1): raise ValueError('利用制限は0/1です')
        accounts.append(p1.Account(r[0],r[1],r[2],r[4],bool(r[5]),r[3]))
    rules=[]; rule_keys={}
    for r in w['収支'].iter_rows(min_row=2,max_row=51,values_only=True):
        if not any(v is not None for v in (*r[1:7],r[19])): continue
        rules.append(p1.Rule(r[0],r[1],r[2],r[3],ym(r[4]),ym(r[5]),r[6],list(r[7:19])))
        rule_keys[r[0]]=r[19]
    events=[]
    for r in w['イベント入力'].iter_rows(min_row=2,max_row=201,values_only=True):
        if not any(v is not None for v in r[1:18]): continue
        if type(r[18]) not in (int,float) or r[18] not in (0,1): raise ValueError('有効は0/1です')
        if r[18]==0: continue
        try:
            typ={'単発':'once','周期':'periodic',None:'once'}[r[2]]
            kind={'支出':'expense','収入':'income',None:'expense'}[r[3]]
        except KeyError as error: raise ValueError('イベント区分が不正です') from error
        if r[13] is not None and (type(r[13]) not in (int,float) or r[13] not in (0,1)): raise ValueError('補正は0/1です')
        e=CashEvent(r[0],r[1],kind,r[7],r[8],event_type=typ,person_id=r[9] if r[9] is not None else 'household',purpose=r[10] if r[10] is not None else 'その他',economic_key=r[11] if r[11] is not None else '',additional_confirmed=r[12]=='はい',inflation_enabled=r[13]==1,base_month=ym(r[14]) if r[13]==1 and r[14] is not None else None,rate_bps=(r[15] if r[15] is not None else 0) if r[13]==1 else 0,linked_dream_id=r[16] or '',note=r[17] or '')
        if typ=='once':
            e.occurrence_month=ym(r[4]); e.first_month=ym(r[5]) if r[5] is not None else None; e.interval_months=r[6]
        else:
            e.first_month=ym(r[4]); e.last_month=ym(r[5]); e.interval_months=r[6]
        events.append(e)
    plan=P2Plan(ym(w['設定']['A2'].value),ym(w['設定']['B2'].value),persons,accounts,rules,events,rule_keys)
    w.close()
    return plan


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--output',default='output/lifeplan_p2a_dev.xlsx'); parser.add_argument('--years',type=int,default=30)
    args=parser.parse_args()
    if not 1<=args.years<=100: parser.error('--years は1〜100です')
    plan=from_p1(p1.demo(),[CashEvent('E001','旅行','expense',1000000,'cash',occurrence_month=(2026,7),additional_confirmed=True)])
    plan.end=(2026+args.years-1,12)
    plan.rules=[replace(r,end=plan.end) for r in plan.rules]
    generate(plan,args.output)
